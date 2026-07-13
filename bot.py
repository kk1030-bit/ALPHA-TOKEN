from __future__ import annotations

import argparse
from collections import deque
from concurrent.futures import Future, ThreadPoolExecutor
import json
import math
import os
import sys
import threading
import time
import urllib.error
import urllib.parse
import urllib.request
from pathlib import Path
from typing import Any

import openpyxl

from notification_cards import notification_caption, render_notification_card
from phase_signal_demo import classify_phase
from trade_planner import ACTIONABLE_DECISIONS, build_trade_plan, calculate_4h_market_context
from onchain_service import analyze_onchain, format_onchain_report
from orderbook_service import (
    analyze_orderbook_accumulation,
    collect_orderbook_snapshots,
    format_orderbook_signal,
    prune_orderbook_db,
)
from oi_service import (
    ApiError,
    WatchSymbol,
    analyze_symbol,
    format_analysis,
    format_scan,
    fmt_num,
    fmt_pct,
    get_current_oi,
    get_dynamic_watch_symbols,
    get_klines,
    get_oi_history,
    get_mark_price,
    get_oi_snapshots,
    normalize_symbol,
    save_analysis,
    scan_symbols,
)
from wgl_v3 import (
    STRUCTURE_MODEL_VERSION,
    TRIGGER_STATES,
    assess_liquidity,
    component_scores as v3_component_scores,
    live_momentum_score,
    migrate_structure_screen,
    scan_structure_universe,
)


ROOT = Path(__file__).resolve().parent
DATA_ROOT = ROOT
BINANCE_FAPI_BASE = "https://fapi.binance.com"
SUBSCRIBERS_PATH = ROOT / "data" / "subscribers.json"
POSITIONS_PATH = ROOT / "data" / "positions.json"
STRATEGY_POSITIONS_PATH = ROOT / "data" / "strategy_positions.json"
STRATEGY_EVENTS_PATH = ROOT / "data" / "strategy_events.jsonl"
WGL_REPORT_EVENTS_PATH = ROOT / "data" / "wgl_reports"
WGL_SEEN_SYMBOLS_PATH = ROOT / "data" / "wgl_seen_symbols"
WGL_SYMBOL_STATS_PATH = ROOT / "data" / "wgl_symbol_stats.json"
WGL_DAILY_SUMMARIES_PATH = ROOT / "data" / "wgl_daily_summaries"
WGL_DAILY_SUMMARY_STATE_PATH = ROOT / "data" / "wgl_daily_summary_state.json"
WGL_FULL_SCAN_EVENTS_PATH = ROOT / "data" / "wgl_scans"
WGL_LATEST_REPORT_PATH = ROOT / "data" / "latest_wgl_report.json"
WGL_SIGNAL_STATES_PATH = ROOT / "data" / "wgl_signal_states.json"
WGL_SIGNAL_OUTCOMES_PATH = ROOT / "data" / "wgl_signal_outcomes.json"
WGL_TRANSITIONS_PATH = ROOT / "data" / "wgl_transitions"
STRUCTURE_CACHE_PATH = ROOT / "data" / "wgl_structure_cache.json"
DEFAULT_TOKEN_EXCEL_PATH = ROOT / "tokens.xlsx"
WATCH_CACHE: dict[str, Any] = {"expires_at": 0.0, "symbols": []}
RUNTIME_SPIKE_HISTORY: dict[str, deque[dict[str, Any]]] = {}
RUNTIME_POSITION_HISTORY: dict[str, deque[dict[str, Any]]] = {}
RUNTIME_ONCHAIN_OBSERVE_AT: dict[str, float] = {}
RUNTIME_FUNDING_CACHE: dict[str, dict[str, Any]] = {}
RUNTIME_WGL_CONTEXT_CACHE: dict[str, dict[str, Any]] = {}
RUNTIME_STRUCTURE_CACHE: dict[str, dict[str, Any]] = {}
RUNTIME_SPOT_FLOW_CACHE: dict[str, dict[str, Any]] = {}
RUNTIME_TRANSITION_ALERTS: deque[str] = deque()
RUNTIME_TRANSITION_LOCK = threading.Lock()
RUNTIME_ORDERBOOK_LOCK = threading.Lock()
RUNTIME_BINANCE_BACKOFF_UNTIL = 0.0
RUNTIME_SPIKE_CURSOR = 0
BOT_STARTED_AT = time.time()


def load_env(path: Path) -> None:
    if not path.exists():
        return
    for raw in path.read_text(encoding="utf-8").splitlines():
        line = raw.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.lstrip("\ufeff").split("=", 1)
        key = key.strip()
        value = value.strip().strip('"').strip("'")
        os.environ.setdefault(key, value)


def telegram_call(token: str, method: str, params: dict[str, Any] | None = None, *, timeout: int = 60) -> Any:
    url = f"https://api.telegram.org/bot{token}/{method}"
    encoded = urllib.parse.urlencode(params or {}).encode("utf-8")
    request = urllib.request.Request(url, data=encoded, method="POST")
    try:
        with urllib.request.urlopen(request, timeout=timeout) as response:
            payload = json.loads(response.read().decode("utf-8"))
    except urllib.error.HTTPError as exc:
        body = exc.read().decode("utf-8", errors="replace")
        raise RuntimeError(f"Telegram HTTP {exc.code}: {body}") from exc
    if not payload.get("ok"):
        raise RuntimeError(f"Telegram API error: {payload}")
    return payload.get("result")


def telegram_send_photo(
    token: str,
    chat_id: int,
    image_bytes: bytes,
    *,
    caption: str = "",
    timeout: int = 60,
) -> Any:
    boundary = f"----alpha-token-{int(time.time() * 1000)}"
    chunks: list[bytes] = []

    def add_field(name: str, value: object) -> None:
        chunks.extend(
            [
                f"--{boundary}\r\n".encode("ascii"),
                f'Content-Disposition: form-data; name="{name}"\r\n\r\n'.encode("ascii"),
                str(value).encode("utf-8"),
                b"\r\n",
            ]
        )

    add_field("chat_id", chat_id)
    if caption:
        add_field("caption", caption[:1024])
    chunks.extend(
        [
            f"--{boundary}\r\n".encode("ascii"),
            b'Content-Disposition: form-data; name="photo"; filename="alpha-token.png"\r\n',
            b"Content-Type: image/png\r\n\r\n",
            image_bytes,
            b"\r\n",
            f"--{boundary}--\r\n".encode("ascii"),
        ]
    )
    request = urllib.request.Request(
        f"https://api.telegram.org/bot{token}/sendPhoto",
        data=b"".join(chunks),
        method="POST",
        headers={"Content-Type": f"multipart/form-data; boundary={boundary}"},
    )
    try:
        with urllib.request.urlopen(request, timeout=timeout) as response:
            payload = json.loads(response.read().decode("utf-8"))
    except urllib.error.HTTPError as exc:
        body = exc.read().decode("utf-8", errors="replace")
        raise RuntimeError(f"Telegram sendPhoto HTTP {exc.code}: {body}") from exc
    if not payload.get("ok"):
        raise RuntimeError(f"Telegram sendPhoto error: {payload}")
    return payload.get("result")


def initialize_update_offset(token: str) -> int:
    """Start from the newest Telegram update so old /report requests do not backlog after restart."""
    try:
        updates = telegram_call(
            token,
            "getUpdates",
            {"offset": -1, "limit": 1, "timeout": 1, "allowed_updates": json.dumps(["message"])},
            timeout=5,
        )
    except Exception as exc:
        print(f"Telegram offset init failed: {exc}", file=sys.stderr, flush=True)
        return 0
    if not updates:
        return 0
    return max(int(update.get("update_id", 0)) for update in updates) + 1


def allowed_chat(chat_id: int, allowed: set[int]) -> bool:
    return not allowed or chat_id in allowed


def parse_allowed_chat_ids() -> set[int]:
    raw = os.environ.get("ALLOWED_CHAT_IDS", "").strip()
    if not raw:
        return set()
    out = set()
    for chunk in raw.split(","):
        chunk = chunk.strip()
        if chunk:
            out.add(int(chunk))
    return out


def load_subscribers() -> set[int]:
    if not SUBSCRIBERS_PATH.exists():
        return set()
    try:
        data = json.loads(SUBSCRIBERS_PATH.read_text(encoding="utf-8"))
    except Exception:
        return set()
    return {int(item) for item in data.get("chat_ids", [])}


def save_subscribers(chat_ids: set[int]) -> None:
    SUBSCRIBERS_PATH.parent.mkdir(parents=True, exist_ok=True)
    SUBSCRIBERS_PATH.write_text(
        json.dumps({"chat_ids": sorted(chat_ids)}, indent=2),
        encoding="utf-8",
    )


def load_positions() -> list[dict[str, Any]]:
    if not POSITIONS_PATH.exists():
        return []
    try:
        data = json.loads(POSITIONS_PATH.read_text(encoding="utf-8"))
    except Exception:
        return []
    positions = data.get("positions", [])
    if not isinstance(positions, list):
        return []
    out = []
    for item in positions:
        if not isinstance(item, dict):
            continue
        if not item.get("symbol") or not item.get("chat_id") or item.get("entry_price") is None:
            continue
        out.append(item)
    return out


def save_positions(positions: list[dict[str, Any]]) -> None:
    POSITIONS_PATH.parent.mkdir(parents=True, exist_ok=True)
    POSITIONS_PATH.write_text(
        json.dumps({"positions": positions}, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )


def save_position_event(event: dict[str, Any]) -> None:
    path = ROOT / "data" / "positions_closed.jsonl"
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("a", encoding="utf-8") as file:
        file.write(json.dumps(event, ensure_ascii=False, separators=(",", ":")) + "\n")


def load_strategy_positions() -> list[dict[str, Any]]:
    if not STRATEGY_POSITIONS_PATH.exists():
        return []
    try:
        data = json.loads(STRATEGY_POSITIONS_PATH.read_text(encoding="utf-8"))
    except Exception:
        return []
    positions = data.get("positions", [])
    if not isinstance(positions, list):
        return []
    return [item for item in positions if isinstance(item, dict) and item.get("symbol")]


def save_strategy_positions(positions: list[dict[str, Any]]) -> None:
    STRATEGY_POSITIONS_PATH.parent.mkdir(parents=True, exist_ok=True)
    STRATEGY_POSITIONS_PATH.write_text(
        json.dumps({"positions": positions}, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )


def save_strategy_event(event: dict[str, Any]) -> None:
    STRATEGY_EVENTS_PATH.parent.mkdir(parents=True, exist_ok=True)
    with STRATEGY_EVENTS_PATH.open("a", encoding="utf-8") as file:
        file.write(json.dumps(event, ensure_ascii=False, separators=(",", ":")) + "\n")


def load_strategy_events(limit: int = 2000) -> list[dict[str, Any]]:
    if not STRATEGY_EVENTS_PATH.exists():
        return []
    lines = STRATEGY_EVENTS_PATH.read_text(encoding="utf-8").splitlines()
    events = []
    for line in lines[-limit:]:
        try:
            item = json.loads(line)
        except Exception:
            continue
        if isinstance(item, dict):
            events.append(item)
    return events


def token_excel_path() -> Path:
    raw = os.environ.get("TOKEN_EXCEL_PATH", "").strip()
    return Path(raw) if raw else DEFAULT_TOKEN_EXCEL_PATH


def report_interval_seconds() -> int:
    raw = os.environ.get("REPORT_INTERVAL_SECONDS", "3600").strip()
    try:
        return max(60, int(raw))
    except ValueError:
        return 3600


def report_top_n() -> int:
    return env_int("REPORT_TOP_N", 20, 1)


def report_lookback_seconds() -> int:
    return env_int("REPORT_LOOKBACK_SECONDS", 3600, 300)


def watchlist_refresh_seconds() -> int:
    return env_int("WATCHLIST_REFRESH_SECONDS", 600, 60)


def oi_snapshot_workers() -> int:
    return env_int("OI_SNAPSHOT_WORKERS", 16, 1)


def env_int(name: str, default: int, minimum: int) -> int:
    raw = os.environ.get(name, str(default)).strip()
    try:
        return max(minimum, int(raw))
    except ValueError:
        return default


def env_float(name: str, default: float, minimum: float) -> float:
    raw = os.environ.get(name, str(default)).strip()
    try:
        return max(minimum, float(raw))
    except ValueError:
        return default


def env_bool(name: str, default: bool) -> bool:
    raw = os.environ.get(name)
    if raw is None:
        return default
    return raw.strip().lower() in {"1", "true", "yes", "on", "y"}


def notification_cards_enabled() -> bool:
    return env_bool("TELEGRAM_CARD_MODE", True)


def trade_plan_min_confidence() -> int:
    return env_int("TRADE_PLAN_MIN_CONFIDENCE", 70, 50)


def trade_plan_min_risk_reward() -> float:
    return env_float("TRADE_PLAN_MIN_RISK_REWARD", 1.5, 1.0)


def liquidity_reference_notional_usd() -> float:
    return env_float("LIQUIDITY_REFERENCE_NOTIONAL_USD", 5_000.0, 100.0)


def liquidity_min_quote_volume_24h_usd() -> float:
    return env_float("LIQUIDITY_MIN_QUOTE_VOLUME_24H_USD", 5_000_000.0, 0.0)


def liquidity_min_quote_volume_1h_usd() -> float:
    return env_float("LIQUIDITY_MIN_QUOTE_VOLUME_1H_USD", 100_000.0, 0.0)


def liquidity_min_depth_multiple() -> float:
    return env_float("LIQUIDITY_MIN_DEPTH_MULTIPLE", 1.5, 0.1)


def liquidity_max_spread_pct() -> float:
    return env_float("LIQUIDITY_MAX_SPREAD_PCT", 0.20, 0.001)


def liquidity_max_slippage_pct() -> float:
    return env_float("LIQUIDITY_MAX_SLIPPAGE_PCT", 0.30, 0.001)


def quick_liquidity_pass(row: Any) -> bool:
    quote_volume = to_float(getattr(row, "quote_volume_24h_usd", None))
    return bool(
        quote_volume is not None
        and quote_volume >= liquidity_min_quote_volume_24h_usd()
    )


def spike_window_seconds() -> int:
    return env_int("OI_SPIKE_WINDOW_SECONDS", 180, 60)


def spike_check_seconds() -> int:
    return env_int("OI_SPIKE_CHECK_SECONDS", 30, 10)


def spike_batch_size() -> int:
    return env_int("OI_SPIKE_BATCH_SIZE", 220, 20)


def effective_spike_batch_size(universe_size: int) -> int:
    if universe_size <= 0:
        return 0
    minimum_batch = math.ceil(
        universe_size * spike_check_seconds() / max(spike_check_seconds(), spike_window_seconds() / 2)
    )
    return min(universe_size, max(spike_batch_size(), minimum_batch))


def spike_min_change_pct() -> float:
    return env_float("OI_SPIKE_MIN_CHANGE_PCT", 8.0, 0.1)


def spike_min_value_usd() -> float:
    return env_float("OI_SPIKE_MIN_VALUE_USD", 500_000.0, 0.0)


def spike_min_contracts_pct() -> float:
    return env_float("OI_SPIKE_MIN_CONTRACTS_PCT", 3.0, 0.1)


def spike_price_confirm_pct() -> float:
    return env_float("OI_SPIKE_PRICE_CONFIRM_PCT", 0.5, 0.0)


def spike_cooldown_seconds() -> int:
    return env_int("OI_SPIKE_COOLDOWN_SECONDS", 900, 60)


def trend_window_seconds() -> int:
    return env_int("OI_TREND_WINDOW_SECONDS", 3600, 900)


def trend_min_contracts_pct() -> float:
    return env_float("OI_TREND_MIN_CONTRACTS_PCT", 5.0, 0.1)


def trend_min_price_pct() -> float:
    return env_float("OI_TREND_MIN_PRICE_PCT", 0.5, 0.0)


def trend_max_bottom_price_pct() -> float:
    return env_float("OI_TREND_MAX_BOTTOM_PRICE_PCT", 12.0, 1.0)


def momentum_min_contracts_pct() -> float:
    return env_float("OI_MOMENTUM_MIN_CONTRACTS_PCT", 8.0, 0.1)


def momentum_min_price_pct() -> float:
    return env_float("OI_MOMENTUM_MIN_PRICE_PCT", 3.0, 0.1)


def momentum_strong_contracts_pct() -> float:
    return env_float("OI_MOMENTUM_STRONG_CONTRACTS_PCT", 15.0, 0.1)


def momentum_strong_min_price_pct() -> float:
    return env_float("OI_MOMENTUM_STRONG_MIN_PRICE_PCT", 1.0, 0.1)


def trend_max_momentum_price_pct() -> float:
    return env_float("OI_TREND_MAX_MOMENTUM_PRICE_PCT", 20.0, 1.0)


def trend_cooldown_seconds() -> int:
    return env_int("OI_TREND_COOLDOWN_SECONDS", 3600, 300)


def position_check_seconds() -> int:
    return env_int("POSITION_CHECK_SECONDS", 30, 10)


def position_oi_window_seconds() -> int:
    return env_int("POSITION_OI_WINDOW_SECONDS", 3600, 300)


def position_funding_exit_pct() -> float:
    return env_float("POSITION_FUNDING_EXIT_PCT", 0.10, 0.0)


def position_stop_loss_pct() -> float:
    return env_float("POSITION_STOP_LOSS_PCT", 7.0, 0.1)


def position_take_profit_pct() -> float:
    return env_float("POSITION_TAKE_PROFIT_PCT", 10.0, 0.1)


def strategy_report_interval_seconds() -> int:
    return env_int("STRATEGY_REPORT_INTERVAL_SECONDS", 10800, 900)


def strategy_scan_interval_seconds() -> int:
    return env_int("STRATEGY_SCAN_INTERVAL_SECONDS", 300, 60)


def strategy_mode() -> str:
    value = os.environ.get("STRATEGY_MODE", "trade_plan").strip().lower()
    return value if value in {"trade_plan", "legacy"} else "trade_plan"


def trade_plan_paper_margin_usd() -> float:
    return env_float("TRADE_PLAN_PAPER_MARGIN_USD", 500.0, 1.0)


def trade_plan_paper_leverage() -> float:
    return env_float("TRADE_PLAN_PAPER_LEVERAGE", 10.0, 1.0)


def trade_plan_signal_max_age_seconds() -> int:
    return env_int("TRADE_PLAN_SIGNAL_MAX_AGE_SECONDS", 1200, 60)


def trade_plan_max_open_positions() -> int:
    return env_int("TRADE_PLAN_MAX_OPEN_POSITIONS", 5, 1)


def trade_plan_reentry_cooldown_seconds() -> int:
    return env_int("TRADE_PLAN_REENTRY_COOLDOWN_SECONDS", 21600, 0)


def strategy_take_profit_pct() -> float:
    return env_float("STRATEGY_TAKE_PROFIT_PCT", 10.0, 0.1)


def strategy_stop_loss_pct() -> float:
    return env_float("STRATEGY_STOP_LOSS_PCT", 7.0, 0.1)


def strategy_max_open_positions() -> int:
    return env_int("STRATEGY_MAX_OPEN_POSITIONS", 20, 1)


def strategy_signal_cooldown_seconds() -> int:
    return env_int("STRATEGY_SIGNAL_COOLDOWN_SECONDS", 21600, 300)


def strategy_min_history_seconds() -> int:
    return env_int("STRATEGY_MIN_HISTORY_SECONDS", 3600, 300)


def strategy_base_window_seconds() -> int:
    return env_int("STRATEGY_BASE_WINDOW_SECONDS", 14400, 900)


def strategy_max_funding_pct() -> float:
    return env_float("STRATEGY_MAX_FUNDING_PCT", 0.06, 0.0)


def strategy_max_extension_from_low_pct() -> float:
    return env_float("STRATEGY_MAX_EXTENSION_FROM_LOW_PCT", 30.0, 1.0)


def strategy_max_15m_price_pct() -> float:
    return env_float("STRATEGY_MAX_15M_PRICE_PCT", 12.0, 1.0)


def strategy_min_oi_15m_pct() -> float:
    return env_float("STRATEGY_MIN_OI_15M_PCT", 1.2, 0.0)


def strategy_min_oi_1h_pct() -> float:
    return env_float("STRATEGY_MIN_OI_1H_PCT", 3.0, 0.0)


def strategy_min_volume_ratio() -> float:
    return env_float("STRATEGY_MIN_VOLUME_RATIO", 1.2, 0.1)


def strategy_min_daily_candles() -> int:
    return env_int("STRATEGY_MIN_DAILY_CANDLES", 45, 20)


def strategy_daily_lookback_days() -> int:
    return env_int("STRATEGY_DAILY_LOOKBACK_DAYS", 180, 45)


def strategy_max_daily_extension_from_low_pct() -> float:
    return env_float("STRATEGY_MAX_DAILY_EXTENSION_FROM_LOW_PCT", 220.0, 5.0)


def strategy_max_daily_range_position_pct() -> float:
    return env_float("STRATEGY_MAX_DAILY_RANGE_POSITION_PCT", 55.0, 1.0)


def strategy_max_base_band_from_low_pct() -> float:
    return env_float("STRATEGY_MAX_BASE_BAND_FROM_LOW_PCT", 18.0, 1.0)


def strategy_min_base_days() -> int:
    return env_int("STRATEGY_MIN_BASE_DAYS", 12, 1)


def strategy_min_daily_range_multiple() -> float:
    return env_float("STRATEGY_MIN_DAILY_RANGE_MULTIPLE", 3.5, 1.0)


def strategy_max_7d_extension_pct() -> float:
    return env_float("STRATEGY_MAX_7D_EXTENSION_PCT", 65.0, 5.0)


def strategy_max_24h_price_pct() -> float:
    return env_float("STRATEGY_MAX_24H_PRICE_PCT", 18.0, 1.0)


def strategy_max_3d_price_pct() -> float:
    return env_float("STRATEGY_MAX_3D_PRICE_PCT", 40.0, 1.0)


def strategy_max_14d_range_pct() -> float:
    return env_float("STRATEGY_MAX_14D_RANGE_PCT", 95.0, 5.0)


def strategy_min_daily_volume_ratio() -> float:
    return env_float("STRATEGY_MIN_DAILY_VOLUME_RATIO", 1.25, 0.1)


def strategy_min_daily_oi_24h_pct() -> float:
    return env_float("STRATEGY_MIN_DAILY_OI_24H_PCT", 3.5, 0.0)


def strategy_min_daily_oi_2d_pct() -> float:
    return env_float("STRATEGY_MIN_DAILY_OI_2D_PCT", 6.0, 0.0)


def strategy_max_new_positions_per_scan() -> int:
    return env_int("STRATEGY_MAX_NEW_POSITIONS_PER_SCAN", 1, 1)


def strategy_min_global_entry_gap_seconds() -> int:
    return env_int("STRATEGY_MIN_GLOBAL_ENTRY_GAP_SECONDS", 14400, 0)


def strategy_onchain_min_score() -> int:
    return env_int("STRATEGY_ONCHAIN_MIN_SCORE", 0, -100)


def strategy_max_bottom_range_position_pct() -> float:
    return env_float("STRATEGY_MAX_BOTTOM_RANGE_POSITION_PCT", 38.0, 1.0)


def strategy_max_recent_low_extension_pct() -> float:
    return env_float("STRATEGY_MAX_RECENT_LOW_EXTENSION_PCT", 45.0, 1.0)


def strategy_min_drawdown_from_high_pct() -> float:
    return env_float("STRATEGY_MIN_DRAWDOWN_FROM_HIGH_PCT", 55.0, 0.0)


def strategy_max_daily_volume_spike_ratio() -> float:
    return env_float("STRATEGY_MAX_DAILY_VOLUME_SPIKE_RATIO", 5.0, 1.0)


def strategy_max_4h_range_position_pct() -> float:
    return env_float("STRATEGY_MAX_4H_RANGE_POSITION_PCT", 58.0, 1.0)


def strategy_max_4h_24h_price_pct() -> float:
    return env_float("STRATEGY_MAX_4H_24H_PRICE_PCT", 18.0, 1.0)


def strategy_min_4h_24h_price_pct() -> float:
    return env_float("STRATEGY_MIN_4H_24H_PRICE_PCT", -12.0, -100.0)


def strategy_max_4h_3d_price_pct() -> float:
    return env_float("STRATEGY_MAX_4H_3D_PRICE_PCT", 35.0, 1.0)


def strategy_min_4h_3d_price_pct() -> float:
    return env_float("STRATEGY_MIN_4H_3D_PRICE_PCT", -35.0, -100.0)


def strategy_max_launch_range_position_pct() -> float:
    return env_float("STRATEGY_MAX_LAUNCH_RANGE_POSITION_PCT", 68.0, 1.0)


def strategy_max_launch_recent_low_extension_pct() -> float:
    return env_float("STRATEGY_MAX_LAUNCH_RECENT_LOW_EXTENSION_PCT", 160.0, 1.0)


def strategy_min_launch_drawdown_from_high_pct() -> float:
    return env_float("STRATEGY_MIN_LAUNCH_DRAWDOWN_FROM_HIGH_PCT", 25.0, 0.0)


def strategy_min_launch_base_days() -> int:
    return env_int("STRATEGY_MIN_LAUNCH_BASE_DAYS", 6, 1)


def strategy_max_launch_24h_price_pct() -> float:
    return env_float("STRATEGY_MAX_LAUNCH_24H_PRICE_PCT", 45.0, 1.0)


def strategy_max_launch_3d_price_pct() -> float:
    return env_float("STRATEGY_MAX_LAUNCH_3D_PRICE_PCT", 120.0, 1.0)


def strategy_max_launch_7d_price_pct() -> float:
    return env_float("STRATEGY_MAX_LAUNCH_7D_PRICE_PCT", 240.0, 1.0)


def strategy_max_launch_14d_range_pct() -> float:
    return env_float("STRATEGY_MAX_LAUNCH_14D_RANGE_PCT", 220.0, 5.0)


def strategy_min_launch_daily_volume_ratio() -> float:
    return env_float("STRATEGY_MIN_LAUNCH_DAILY_VOLUME_RATIO", 1.15, 0.1)


def strategy_max_launch_daily_volume_ratio() -> float:
    return env_float("STRATEGY_MAX_LAUNCH_DAILY_VOLUME_RATIO", 18.0, 1.0)


def strategy_max_launch_4h_range_position_pct() -> float:
    return env_float("STRATEGY_MAX_LAUNCH_4H_RANGE_POSITION_PCT", 88.0, 1.0)


def strategy_max_launch_4h_24h_price_pct() -> float:
    return env_float("STRATEGY_MAX_LAUNCH_4H_24H_PRICE_PCT", 55.0, 1.0)


def strategy_max_launch_4h_3d_price_pct() -> float:
    return env_float("STRATEGY_MAX_LAUNCH_4H_3D_PRICE_PCT", 140.0, 1.0)


def launch_report_top_n() -> int:
    return env_int("LAUNCH_REPORT_TOP_N", 12, 1)


def launch_min_score() -> int:
    return env_int("LAUNCH_MIN_SCORE", 62, 1)


def onchain_observe_cooldown_seconds() -> int:
    return env_int("ONCHAIN_OBSERVE_COOLDOWN_SECONDS", 21600, 300)


def onchain_report_interval_seconds() -> int:
    return env_int("ONCHAIN_REPORT_INTERVAL_SECONDS", 3600, 900)


def onchain_report_candidate_count() -> int:
    return env_int("ONCHAIN_REPORT_CANDIDATES", 50, 5)


def momentum_deep_quota() -> int:
    return env_int("MOMENTUM_DEEP_QUOTA", 15, 0)


def onchain_deep_candidate_count() -> int:
    return env_int("ONCHAIN_DEEP_CANDIDATES", 10, 1)


def spot_flow_candidate_count() -> int:
    return env_int("SPOT_FLOW_CANDIDATES", 12, 1)


def onchain_report_top_n() -> int:
    return env_int("ONCHAIN_REPORT_TOP_N", 8, 1)


def ravelab_report_top_n() -> int:
    return env_int("RAVELAB_REPORT_TOP_N", 12, 1)


def ravelab_min_score() -> int:
    return env_int("RAVELAB_MIN_SCORE", 60, 1)


def orderbook_enabled() -> bool:
    return env_bool("ORDERBOOK_ENABLED", True)


def orderbook_db_path() -> Path:
    raw = os.environ.get("ORDERBOOK_DB_PATH", str(ROOT / "data" / "orderbook.sqlite")).strip()
    return Path(raw)


def orderbook_collect_interval_seconds() -> int:
    return env_int("ORDERBOOK_COLLECT_INTERVAL_SECONDS", 60, 30)


def orderbook_watch_candidates() -> int:
    return env_int("ORDERBOOK_WATCH_CANDIDATES", 10, 5)


def orderbook_report_seed_candidates() -> int:
    return env_int("ORDERBOOK_REPORT_SEED_CANDIDATES", 10, 1)


def orderbook_workers() -> int:
    return env_int("ORDERBOOK_WORKERS", 12, 1)


def orderbook_lookback_seconds() -> int:
    return env_int("ORDERBOOK_LOOKBACK_SECONDS", 14400, 900)


def orderbook_min_snapshots() -> int:
    return env_int("ORDERBOOK_MIN_SNAPSHOTS", 12, 3)


def orderbook_min_score() -> int:
    return env_int("ORDERBOOK_MIN_SCORE", 50, 1)


def orderbook_report_top_n() -> int:
    return env_int("ORDERBOOK_REPORT_TOP_N", 10, 1)


def composite_report_top_n() -> int:
    return env_int("COMPOSITE_REPORT_TOP_N", 5, 1)


def structure_scan_candidate_count() -> int:
    return env_int("STRUCTURE_SCAN_CANDIDATES", 20, 5)


def structure_cache_seconds() -> int:
    return env_int("STRUCTURE_CACHE_SECONDS", 21600, 900)


def structure_scan_workers() -> int:
    return env_int("STRUCTURE_SCAN_WORKERS", 8, 1)


def structure_refresh_batch_size() -> int:
    return env_int("STRUCTURE_REFRESH_BATCH_SIZE", 60, 10)


def structure_reference_symbols() -> set[str]:
    raw = os.environ.get(
        "STRUCTURE_REFERENCE_SYMBOLS",
        "RAVEUSDT,LABUSDT,MYXUSDT,COAIUSDT,BEATUSDT,BLESSUSDT",
    )
    return {normalize_symbol(value) for value in raw.split(",") if value.strip()}


def wgl_scan_candidate_count() -> int:
    return env_int("WGL_SCAN_CANDIDATES", 18, 5)


def wgl_min_score() -> int:
    return env_int("WGL_MIN_SCORE", 45, 1)


def wgl_pullback_min_score() -> int:
    return env_int("WGL_PULLBACK_MIN_SCORE", 58, 1)


def wgl_pullback_min_24h_price_pct() -> float:
    return env_float("WGL_PULLBACK_MIN_24H_PRICE_PCT", 8.0, 0.0)


def wgl_pullback_max_24h_price_pct() -> float:
    return env_float("WGL_PULLBACK_MAX_24H_PRICE_PCT", 45.0, 1.0)


def wgl_pullback_min_6h_range_position_pct() -> float:
    return env_float("WGL_PULLBACK_MIN_6H_RANGE_POSITION_PCT", 8.0, 0.0)


def wgl_pullback_max_6h_range_position_pct() -> float:
    return env_float("WGL_PULLBACK_MAX_6H_RANGE_POSITION_PCT", 55.0, 1.0)


def wgl_pullback_max_funding_pct() -> float:
    return env_float("WGL_PULLBACK_MAX_FUNDING_PCT", 0.06, 0.0)


def wgl_pullback_min_volume_ratio() -> float:
    return env_float("WGL_PULLBACK_MIN_VOLUME_RATIO", 0.8, 0.0)


def wgl_pullback_max_oi_1h_drop_pct() -> float:
    return env_float("WGL_PULLBACK_MAX_OI_1H_DROP_PCT", -3.0, -100.0)


def wgl_daily_summary_hour() -> int:
    return min(23, env_int("WGL_DAILY_SUMMARY_HOUR", 23, 0))


def wgl_daily_summary_minute() -> int:
    return min(59, env_int("WGL_DAILY_SUMMARY_MINUTE", 59, 0))


def normalize_day_string(value: Any) -> str:
    text = str(value or "").strip()
    if not text:
        return ""
    if "T" in text:
        text = text.split("T", 1)[0]
    if " " in text:
        text = text.split(" ", 1)[0]
    text = text.replace("/", "-")
    if len(text) == 8 and text.isdigit():
        return f"{text[:4]}-{text[4:6]}-{text[6:8]}"
    parts = text.split("-")
    if len(parts) == 3 and all(part.isdigit() for part in parts):
        return f"{int(parts[0]):04d}-{int(parts[1]):02d}-{int(parts[2]):02d}"
    return text


def wgl_stats_start_date() -> str:
    return normalize_day_string(os.environ.get("WGL_STATS_START_DATE", "2026-07-09"))


def orderbook_prune_days() -> int:
    return env_int("ORDERBOOK_PRUNE_DAYS", 7, 1)


def strategy_blocklist() -> set[str]:
    raw = os.environ.get(
        "STRATEGY_SYMBOL_BLOCKLIST",
        "SANTOSUSDT,PSGUSDT,BARUSDT,ATMUSDT,ASRUSDT,LAZIOUSDT,PORTOUSDT,ALPINEUSDT",
    )
    return {item.strip().upper() for item in raw.split(",") if item.strip()}


def watch_mode() -> str:
    mode = os.environ.get("WATCH_MODE", "dynamic").strip().lower()
    if mode not in {"dynamic", "excel", "both"}:
        return "dynamic"
    return mode


def ensure_token_template(path: Path) -> None:
    if path.exists():
        return
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "tokens"
    sheet.append(["symbol"])
    for symbol in ["BTC", "ETH", "SOL", "BNB", "XRP", "DOGE", "HYPE", "ZEC"]:
        sheet.append([symbol])
    workbook.save(path)


def read_symbols_from_excel(path: Path) -> list[str]:
    ensure_token_template(path)
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []

    header = [str(value).strip().lower() if value is not None else "" for value in rows[0]]
    symbol_headers = {"symbol", "symbols", "coin", "coins", "token", "tokens", "幣種", "代幣"}
    col_idx = 0
    for idx, value in enumerate(header):
        if value in symbol_headers:
            col_idx = idx
            break

    start_idx = 1 if any(header) else 0
    symbols: list[str] = []
    for row in rows[start_idx:]:
        if col_idx >= len(row):
            continue
        value = row[col_idx]
        if value is None:
            continue
        symbol = str(value).strip()
        if symbol and not symbol.startswith("#"):
            symbols.append(symbol)
    return symbols


def resolve_watch_symbols(*, force_refresh: bool = False) -> list[WatchSymbol]:
    now = time.time()
    cached = WATCH_CACHE.get("symbols") or []
    if not force_refresh and cached and now < float(WATCH_CACHE.get("expires_at", 0.0)):
        return list(cached)

    mode = watch_mode()
    symbols: list[WatchSymbol] = []
    seen = set()

    if mode in {"dynamic", "both"}:
        for watch in get_dynamic_watch_symbols():
            symbols.append(watch)
            seen.add(watch.symbol)

    if mode in {"excel", "both"}:
        for raw_symbol in read_symbols_from_excel(token_excel_path()):
            try:
                symbol = raw_symbol.strip().upper()
                if not symbol.endswith("USDT"):
                    symbol = f"{symbol}USDT"
            except Exception:
                continue
            if symbol in seen:
                continue
            symbols.append(
                WatchSymbol(
                    symbol=symbol,
                    market_symbol=symbol.removesuffix("USDT"),
                    market_rank=None,
                    marketcap_usd=None,
                    source="excel",
                )
            )
            seen.add(symbol)

    WATCH_CACHE["symbols"] = symbols
    WATCH_CACHE["expires_at"] = now + watchlist_refresh_seconds()
    return list(symbols)


def watch_source_description() -> str:
    mode = watch_mode()
    if mode == "dynamic":
        return "全部 Binance USDT 合約；CryptoBubbles 市值與排名只作參考"
    if mode == "both":
        return "全部 Binance USDT 合約，加上 Excel 補充名單；不設市值條件"
    return f"只使用 Excel：{token_excel_path()}"


def write_report_excel(report_rows: list[Any], source_description: str) -> Path:
    reports_dir = ROOT / "data" / "reports"
    reports_dir.mkdir(parents=True, exist_ok=True)
    path = reports_dir / f"oi_report_{time.strftime('%Y%m%d_%H%M%S')}.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "oi_sorted"
    sheet.append(
        [
            "rank",
            "attention_rank",
            "attention_score",
            "signal",
            "decision",
            "reason",
            "symbol",
            "price",
            "mark_price",
            "open_interest",
            "oi_value_usd",
            "oi_to_marketcap_pct",
            "contracts_change_180s_pct",
            "contracts_change_1h_pct",
            "price_change_180s_pct",
            "price_change_1h_pct",
            "funding_rate_pct",
            "market_rank",
            "marketcap_usd",
            "market_symbol",
            "source",
            "timestamp_utc",
            "watch_source",
        ]
    )
    for idx, item in enumerate(report_rows, 1):
        row = item["row"]
        metrics = item["metrics"]
        sheet.append(
            [
                row.rank,
                idx,
                item["score"],
                item["signal"],
                item["decision"],
                item["reason"],
                row.symbol,
                row.price,
                row.mark_price,
                row.open_interest,
                row.oi_value_usd,
                metrics["oi_to_marketcap_pct"],
                metrics["contracts_180s_pct"],
                metrics["contracts_1h_pct"],
                metrics["price_180s_pct"],
                metrics["price_1h_pct"],
                row.funding_rate_pct,
                row.market_rank,
                row.marketcap_usd,
                row.market_symbol,
                row.source,
                row.timestamp_utc,
                source_description,
            ]
        )
    for col in range(1, sheet.max_column + 1):
        sheet.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 18
    workbook.save(path)
    return path


def history_pct_change(
    history: dict[str, deque[dict[str, Any]]],
    symbol: str,
    current_value: float | None,
    value_key: str,
    now: float,
    lookback_seconds: int,
) -> float | None:
    if current_value is None:
        return None
    samples = history.get(symbol)
    if not samples:
        return None
    baseline = None
    target_ts = now - lookback_seconds
    for sample in samples:
        if float(sample["ts"]) <= target_ts:
            baseline = sample
        else:
            break
    if baseline is None:
        return None
    old_value = baseline.get(value_key)
    if old_value is None or float(old_value) == 0:
        return None
    return (float(current_value) - float(old_value)) / float(old_value) * 100.0


def report_metrics(row: Any, history: dict[str, deque[dict[str, Any]]], now: float) -> dict[str, float | None]:
    oi_to_marketcap_pct = None
    if row.oi_value_usd is not None and row.marketcap_usd:
        oi_to_marketcap_pct = row.oi_value_usd / row.marketcap_usd * 100.0
    return {
        "oi_to_marketcap_pct": oi_to_marketcap_pct,
        "contracts_180s_pct": history_pct_change(history, row.symbol, row.open_interest, "open_interest", now, 180),
        "contracts_1h_pct": history_pct_change(
            history, row.symbol, row.open_interest, "open_interest", now, report_lookback_seconds()
        ),
        "price_180s_pct": history_pct_change(history, row.symbol, row.mark_price, "mark_price", now, 180),
        "price_1h_pct": history_pct_change(history, row.symbol, row.mark_price, "mark_price", now, report_lookback_seconds()),
    }


def remember_report_sample(history: dict[str, deque[dict[str, Any]]], row: Any, now: float) -> None:
    if row.mark_price is None or row.open_interest is None:
        return
    samples = history.setdefault(row.symbol, deque())
    if samples and now - float(samples[-1]["ts"]) < 10:
        return
    samples.append(
        {
            "ts": now,
            "symbol": row.symbol,
            "oi_value_usd": row.oi_value_usd,
            "open_interest": row.open_interest,
            "mark_price": row.mark_price,
            "funding_rate_pct": row.funding_rate_pct,
            "quote_volume_24h_usd": getattr(row, "quote_volume_24h_usd", None),
            "market_rank": row.market_rank,
            "marketcap_usd": row.marketcap_usd,
        }
    )
    max_history_age = max(position_oi_window_seconds(), report_lookback_seconds(), spike_window_seconds()) + 600
    while samples and now - float(samples[0]["ts"]) > max_history_age:
        samples.popleft()


def phase_samples(row: Any, history: dict[str, deque[dict[str, Any]]]) -> list[dict[str, Any]]:
    samples = []
    for sample in history.get(row.symbol, deque()):
        price = sample.get("mark_price", sample.get("price"))
        open_interest = sample.get("open_interest")
        if price is None or open_interest is None:
            continue
        samples.append(
            {
                "ts": sample["ts"],
                "price": price,
                "open_interest": open_interest,
                "funding_rate_pct": sample.get("funding_rate_pct"),
            }
        )
    return sorted(samples, key=lambda item: float(item["ts"]))


def phase_report_decision(row: Any, history: dict[str, deque[dict[str, Any]]]) -> tuple[str, str] | None:
    samples = phase_samples(row, history)
    if len(samples) < 3:
        return None
    if float(samples[-1]["ts"]) - float(samples[0]["ts"]) < strategy_min_history_seconds():
        return None
    result = classify_phase(samples)
    return result.decision, f"{result.phase}｜{result.reason}"


def attention_signal(metrics: dict[str, float | None]) -> str:
    contracts = metrics.get("contracts_180s_pct")
    price = metrics.get("price_180s_pct")
    if contracts is not None and contracts >= spike_min_contracts_pct():
        if price is not None and price >= spike_price_confirm_pct():
            return "多頭建倉"
        if price is not None and price <= -spike_price_confirm_pct():
            return "空頭建倉"
        return "倉位堆積"
    return "觀察"


def report_decision(row: Any, metrics: dict[str, float | None], score: int) -> tuple[str, str]:
    c180 = metrics.get("contracts_180s_pct")
    c1h = metrics.get("contracts_1h_pct")
    p180 = metrics.get("price_180s_pct")
    funding = row.funding_rate_pct

    funding_abs = abs(funding) if funding is not None else None
    funding_hot = funding_abs is not None and funding_abs >= 0.25
    funding_warm = funding_abs is not None and funding_abs >= 0.10
    contracts_confirmed = (c180 is not None and c180 >= spike_min_contracts_pct()) or (c1h is not None and c1h >= 5)
    price_pumped = p180 is not None and p180 >= 1.5
    price_up_without_oi = p180 is not None and p180 >= spike_price_confirm_pct() and (c180 is None or c180 < 1)

    if funding_hot:
        side = "多方" if funding and funding > 0 else "空方"
        return "不要進", f"{side} funding 過熱，容易反抽或插針"
    if price_up_without_oi and funding_warm:
        return "不要進", "價格拉升但合約OI沒跟，疑似誘多"
    if price_pumped and not contracts_confirmed:
        return "不要進", "價格已先拉，OI確認不足，不追"
    if contracts_confirmed and not funding_warm:
        if p180 is not None and p180 <= -spike_price_confirm_pct():
            return "再確認", "OI增加但價格走弱，先等翻強"
        if p180 is not None and p180 <= 1.2:
            return "再確認", "OI增加但需確認不是短線出貨"
        return "再確認", "OI有進場，但價格已動，等回踩確認"

    if score >= 45:
        return "再確認", "分數不低，但方向仍需價格與OI同步"
    return "再確認", "訊號不足，等 OI 或價格確認"


def decision_order(label: str) -> int:
    order = {"埋伏A": 0, "埋伏B": 1, "埋伏": 2, "再確認": 3, "不要進": 4}
    return order.get(label, 9)


def attention_score(row: Any, metrics: dict[str, float | None]) -> int:
    score = 0.0
    c180 = metrics.get("contracts_180s_pct")
    if c180 is not None:
        if c180 >= 8:
            score += 28
        elif c180 >= 3:
            score += 20
        elif c180 >= 1:
            score += 10

    c1h = metrics.get("contracts_1h_pct")
    if c1h is not None:
        if c1h >= 15:
            score += 26
        elif c1h >= 5:
            score += 18
        elif c1h >= 2:
            score += 10

    p180 = metrics.get("price_180s_pct")
    if c180 is not None and c180 > 0 and p180 is not None:
        if abs(p180) >= spike_price_confirm_pct():
            score += 10
        elif abs(p180) >= spike_price_confirm_pct() / 2:
            score += 5

    funding = row.funding_rate_pct
    if funding is not None:
        if abs(funding) <= 0.05:
            score += 8
        elif abs(funding) <= 0.25:
            score += 3
        else:
            score -= 8

    return max(0, min(100, int(round(score))))


def build_oi_report(history: dict[str, deque[dict[str, Any]]] | None = None) -> tuple[str, Path]:
    history = history if history is not None else RUNTIME_SPIKE_HISTORY
    watch_symbols = resolve_watch_symbols()
    if not watch_symbols:
        raise RuntimeError("No watch symbols found")

    all_rows = get_oi_snapshots(watch_symbols, max_workers=oi_snapshot_workers())
    rows = [row for row in all_rows if row.oi_value_usd is not None]
    top_n = report_top_n()
    now = time.time()
    scored_rows = []
    for row in rows:
        remember_report_sample(history, row, now)
        metrics = report_metrics(row, history, now)
        score = attention_score(row, metrics)
        phase_decision = phase_report_decision(row, history)
        if phase_decision:
            decision, reason = phase_decision
        else:
            decision, reason = report_decision(row, metrics, score)
            if decision == "埋伏":
                decision = "再確認"
                reason = f"階段資料不足，舊規則埋伏暫列再確認｜{reason}"
        scored_rows.append(
            {
                "row": row,
                "metrics": metrics,
                "score": score,
                "signal": attention_signal(metrics),
                "decision": decision,
                "reason": reason,
            }
        )
    scored_rows.sort(
        key=lambda item: (
            item["score"],
            item["row"].oi_value_usd or 0,
        ),
        reverse=True,
    )
    top_rows = scored_rows[:top_n]
    display_rows = sorted(top_rows, key=lambda item: (decision_order(item["decision"]), -item["score"]))
    source_description = watch_source_description()
    report_path = write_report_excel(display_rows, source_description)
    unsupported_count = len(all_rows) - len(rows)

    labels = ["埋伏A", "埋伏B", "再確認", "不要進"]
    counts = {label: sum(1 for item in top_rows if item["decision"] == label) for label in labels}
    lines = [
        f"階段 OI 雷達前 {min(top_n, len(rows))}｜埋伏A {counts['埋伏A']}｜埋伏B {counts['埋伏B']}｜再確認 {counts['再確認']}｜不要進 {counts['不要進']}",
        f"監控：{len(rows)}/{len(all_rows)} 可查 OI｜來源：Binance 全部 USDT 合約（無市值門檻）",
        "判讀：埋伏A=早期推升；埋伏B=洗盤收回；再確認=方向未明；不要進=尾端/誘多/過熱。",
    ]
    if unsupported_count:
        lines.append(f"已排除無 OI 資料標的：{unsupported_count}")
    for label in labels:
        group = [item for item in display_rows if item["decision"] == label]
        if not group:
            continue
        lines.append("")
        lines.append(f"【{label}】")
        for idx, item in enumerate(group, 1):
            row = item["row"]
            metrics = item["metrics"]
            rank_text = f"#{row.market_rank}" if row.market_rank else "#n/a"
            lines.append(
                f"{idx}. {row.symbol}｜{item['score']}分｜{item['reason']}｜"
                f"OI/市值 {fmt_pct(metrics['oi_to_marketcap_pct'])}｜"
                f"OI180 {fmt_pct(metrics['contracts_180s_pct'])}｜"
                f"價180 {fmt_pct(metrics['price_180s_pct'])}｜"
                f"F {fmt_pct(row.funding_rate_pct, 4)}｜市值{rank_text}"
            )
    lines.append(f"已存檔：{report_path}")
    return "\n".join(lines), report_path


def save_spike_event(event: dict[str, Any]) -> None:
    path = ROOT / "data" / "spikes" / f"{time.strftime('%Y%m%d')}.jsonl"
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("a", encoding="utf-8") as file:
        file.write(json.dumps(event, ensure_ascii=False) + "\n")


def save_wgl_report_event(candidates: list[dict[str, Any]], report_text: str) -> None:
    WGL_REPORT_EVENTS_PATH.mkdir(parents=True, exist_ok=True)
    path = WGL_REPORT_EVENTS_PATH / f"{time.strftime('%Y%m%d')}.jsonl"
    event = {
        "timestamp_utc": time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime()),
        "timestamp_local": time.strftime("%Y-%m-%dT%H:%M:%S%z"),
        "top_n": len(candidates),
        "report_text": report_text,
        "items": [],
    }
    for idx, item in enumerate(candidates, 1):
        row = item.get("row")
        wgl = item.get("wgl") or {}
        metrics = item.get("metrics") or {}
        event["items"].append(
            {
                "rank": idx,
                "report_rank": item.get("report_rank", idx),
                "symbol": item.get("symbol"),
                "score": item.get("score"),
                "signal_state": item.get("signal_state"),
                "structure_score": item.get("structure_score"),
                "capital_score": item.get("capital_score"),
                "trigger_score": item.get("trigger_score"),
                "quality_score": item.get("quality_score"),
                "risk_score": item.get("risk_score"),
                "liquidity_status": item.get("liquidity_status"),
                "liquidity_score": item.get("liquidity_score"),
                "liquidity_blocked": bool(item.get("liquidity_blocked")),
                "liquidity_reasons": item.get("liquidity_reasons") or [],
                "quote_volume_24h_usd": to_float(item.get("quote_volume_24h_usd")),
                "liquidity_bid_depth_05pct": to_float(item.get("liquidity_bid_depth_05pct")),
                "liquidity_ask_depth_05pct": to_float(item.get("liquidity_ask_depth_05pct")),
                "liquidity_buy_slippage_pct": to_float(item.get("liquidity_buy_slippage_pct")),
                "liquidity_sell_slippage_pct": to_float(item.get("liquidity_sell_slippage_pct")),
                "labels": item.get("labels"),
                "reasons": item.get("reasons") or [],
                "trade_bucket": item.get("trade_bucket"),
                "trade_decision": item.get("trade_decision"),
                "trade_setup": item.get("trade_setup"),
                "trade_reason": item.get("trade_reason"),
                "trade_side": item.get("trade_side"),
                "plan_confidence": item.get("plan_confidence"),
                "entry_low": to_float(item.get("entry_low")),
                "entry_high": to_float(item.get("entry_high")),
                "take_profit_1": to_float(item.get("take_profit_1")),
                "take_profit_2": to_float(item.get("take_profit_2")),
                "stop_loss": to_float(item.get("stop_loss")),
                "risk_reward_1": to_float(item.get("risk_reward_1")),
                "risk_reward_2": to_float(item.get("risk_reward_2")),
                "plan_reason": item.get("plan_reason"),
                "first_seen": item.get("first_seen"),
                "wgl_score": item.get("wgl_score"),
                "wgl_action": item.get("wgl_action") or wgl.get("action"),
                "wgl_stage": wgl.get("stage"),
                "mark_price": to_float(getattr(row, "mark_price", None)),
                "funding_rate_pct": to_float(getattr(row, "funding_rate_pct", None)),
                "open_interest": to_float(getattr(row, "open_interest", None)),
                "oi_value_usd": to_float(getattr(row, "oi_value_usd", None)),
                "market_rank": getattr(row, "market_rank", None),
                "marketcap_usd": to_float(getattr(row, "marketcap_usd", None)),
                "oi_to_marketcap_pct": to_float(metrics.get("oi_to_marketcap_pct")),
                "price_1h_pct": to_float(metrics.get("price_1h_pct")),
                "contracts_1h_pct": to_float(metrics.get("contracts_1h_pct")),
                "wgl_price_24h_pct": to_float(wgl.get("price_24h_pct")),
                "wgl_oi_1h_pct": to_float(wgl.get("oi_1h_pct")),
                "wgl_range_6h_position_pct": to_float(wgl.get("range_6h_position_pct")),
                "wgl_range_24h_position_pct": to_float(wgl.get("range_24h_position_pct")),
                "wgl_drawdown_from_24h_high_pct": to_float(wgl.get("drawdown_from_24h_high_pct")),
                "wgl_strong_pullback": bool(wgl.get("strong_pullback")),
            }
        )
    with path.open("a", encoding="utf-8") as file:
        file.write(json.dumps(event, ensure_ascii=False) + "\n")


def save_wgl_full_scan_event(
    candidates: list[dict[str, Any]],
    selected_symbols: set[str],
    deep_candidates: list[dict[str, Any]] | None = None,
    component_candidates: list[dict[str, Any]] | None = None,
) -> None:
    WGL_FULL_SCAN_EVENTS_PATH.mkdir(parents=True, exist_ok=True)
    path = WGL_FULL_SCAN_EVENTS_PATH / f"{time.strftime('%Y%m%d')}.jsonl"
    event = {
        "schema": "wgl-full-universe-v3",
        "timestamp_utc": time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime()),
        "timestamp_local": time.strftime("%Y-%m-%dT%H:%M:%S%z"),
        "universe_size": len(candidates),
        "selected_count": len(selected_symbols),
        "items": [],
    }
    deep_by_symbol = {
        item["row"].symbol: item
        for item in (deep_candidates or [])
    }
    component_by_symbol = {
        str(item.get("symbol") or "").upper(): item
        for item in (component_candidates or [])
    }
    for item in candidates:
        deep = deep_by_symbol.get(item["row"].symbol) or item
        row = deep["row"]
        screen = item.get("structure_screen") or {}
        metrics = deep.get("metrics") or {}
        component = component_by_symbol.get(str(row.symbol).upper()) or {}
        event["items"].append(
            {
                "symbol": row.symbol,
                "selected_for_deep_scan": row.symbol in selected_symbols,
                "structure_score": screen.get("score"),
                "structure_eligible": bool(screen.get("eligible")),
                "structure_state": screen.get("state_hint"),
                "rejection_reason": screen.get("reject_reason") or "",
                "prefilter_score": deep.get("prefilter_score") or item.get("prefilter_score"),
                "live_momentum_score": deep.get("live_momentum_score") or item.get("live_momentum_score"),
                "momentum_rank_score": deep.get("momentum_rank_score"),
                "selection_lane": deep.get("selection_lane"),
                "mark_price": to_float(getattr(row, "mark_price", None)),
                "funding_rate_pct": to_float(getattr(row, "funding_rate_pct", None)),
                "oi_value_usd": to_float(getattr(row, "oi_value_usd", None)),
                "contracts_180s_pct": to_float(metrics.get("contracts_180s_pct")),
                "contracts_1h_pct": to_float(metrics.get("contracts_1h_pct")),
                "price_180s_pct": to_float(metrics.get("price_180s_pct")),
                "price_1h_pct": to_float(metrics.get("price_1h_pct")),
                "basis_pct": to_float(metrics.get("basis_pct")),
                "spot_taker_imbalance": to_float(metrics.get("spot_taker_imbalance")),
                "spot_taker_notional": to_float(metrics.get("spot_taker_notional")),
                "signal_state": component.get("signal_state"),
                "overall_score": component.get("score"),
                "capital_score": component.get("capital_score"),
                "trigger_score": component.get("trigger_score"),
                "quality_score": component.get("quality_score"),
                "risk_score": component.get("risk_score"),
                "liquidity_status": component.get("liquidity_status"),
                "liquidity_score": component.get("liquidity_score"),
                "liquidity_blocked": bool(component.get("liquidity_blocked")),
                "liquidity_reasons": component.get("liquidity_reasons") or [],
                "quote_volume_24h_usd": to_float(getattr(row, "quote_volume_24h_usd", None)),
                "liquidity_bid_depth_05pct": to_float(component.get("liquidity_bid_depth_05pct")),
                "liquidity_ask_depth_05pct": to_float(component.get("liquidity_ask_depth_05pct")),
                "liquidity_buy_slippage_pct": to_float(component.get("liquidity_buy_slippage_pct")),
                "liquidity_sell_slippage_pct": to_float(component.get("liquidity_sell_slippage_pct")),
                "short_squeeze": bool(component.get("short_squeeze")),
                "trade_decision": component.get("trade_decision"),
                "trade_side": component.get("trade_side"),
                "plan_confidence": component.get("plan_confidence"),
                "entry_low": to_float(component.get("entry_low")),
                "entry_high": to_float(component.get("entry_high")),
                "take_profit_1": to_float(component.get("take_profit_1")),
                "take_profit_2": to_float(component.get("take_profit_2")),
                "stop_loss": to_float(component.get("stop_loss")),
                "risk_reward_1": to_float(component.get("risk_reward_1")),
                "risk_reward_2": to_float(component.get("risk_reward_2")),
                "plan_reason": component.get("plan_reason"),
                "market_rank_reference": getattr(row, "market_rank", None),
                "marketcap_reference_usd": to_float(getattr(row, "marketcap_usd", None)),
                "data_points": screen.get("data_points"),
                "range_position_pct": screen.get("range_position_pct"),
                "recent_range_position_pct": screen.get("recent_range_position_pct"),
                "drawdown_from_high_pct": screen.get("drawdown_from_high_pct"),
                "recent_low_extension_pct": screen.get("recent_low_extension_pct"),
                "base_days": screen.get("base_days"),
                "compression_ratio": screen.get("compression_ratio"),
                "volume_ratio_7d": screen.get("volume_ratio_7d"),
                "prior_test_pump_pct": screen.get("prior_test_pump_pct"),
            }
        )
    with path.open("a", encoding="utf-8") as file:
        file.write(json.dumps(event, ensure_ascii=False, separators=(",", ":")) + "\n")


def load_structure_cache() -> dict[str, dict[str, Any]]:
    if not STRUCTURE_CACHE_PATH.exists():
        return {}
    try:
        payload = json.loads(STRUCTURE_CACHE_PATH.read_text(encoding="utf-8"))
    except Exception:
        return {}
    entries = payload.get("entries") if isinstance(payload, dict) else None
    if not isinstance(entries, dict):
        return {}
    migrated: dict[str, dict[str, Any]] = {}
    for symbol, value in entries.items():
        screen = migrate_structure_screen(value.get("screen") or {})
        if screen is None or "日線讀取失敗" in str(screen.get("reject_reason") or ""):
            continue
        migrated[symbol] = {"cached_at": value.get("cached_at", 0.0), "screen": screen}
    return migrated


def save_structure_cache() -> None:
    cutoff = time.time() - max(structure_cache_seconds() * 4, 86400)
    entries = {
        symbol: value
        for symbol, value in RUNTIME_STRUCTURE_CACHE.items()
        if (
            isinstance(value, dict)
            and float(value.get("cached_at", 0.0)) >= cutoff
            and (value.get("screen") or {}).get("model_version") == STRUCTURE_MODEL_VERSION
        )
    }
    payload = {
        "schema": "wgl-structure-cache-v3",
        "updated_utc": time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime()),
        "entries": entries,
    }
    STRUCTURE_CACHE_PATH.parent.mkdir(parents=True, exist_ok=True)
    temp_path = STRUCTURE_CACHE_PATH.with_suffix(".tmp")
    temp_path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
    temp_path.replace(STRUCTURE_CACHE_PATH)


def save_latest_wgl_report(report_text: str, candidates: list[dict[str, Any]], universe_size: int) -> None:
    WGL_LATEST_REPORT_PATH.parent.mkdir(parents=True, exist_ok=True)
    payload = {
        "schema": "wgl-latest-report-v3",
        "generated_at": time.time(),
        "generated_utc": time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime()),
        "generated_local": time.strftime("%Y-%m-%d %H:%M"),
        "universe_size": universe_size,
        "symbols": [item.get("symbol") for item in candidates],
        "items": [
            {
                "symbol": item.get("symbol"),
                "report_rank": item.get("report_rank"),
                "signal_state": item.get("signal_state"),
                "score": item.get("score"),
                **compact_trade_plan(item),
            }
            for item in candidates
        ],
        "report_text": report_text,
    }
    WGL_LATEST_REPORT_PATH.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def load_cached_wgl_report() -> str:
    if not WGL_LATEST_REPORT_PATH.exists():
        return "資金雷達正在建立第一份全市場報告，完成後 /report 會立即顯示快取結果。"
    try:
        payload = json.loads(WGL_LATEST_REPORT_PATH.read_text(encoding="utf-8"))
    except Exception:
        return "最近報告快取損壞，系統將在下一輪掃描重建。"
    report = str(payload.get("report_text") or "目前沒有候選。")
    generated_at = to_float(payload.get("generated_at"))
    age_minutes = max(0, int((time.time() - generated_at) / 60)) if generated_at else None
    age_text = f"{age_minutes} 分鐘前" if age_minutes is not None else "時間未知"
    return f"{report}\n\n資料更新：{payload.get('generated_local', '-')}（{age_text}）"


def initial_report_due_at(now: float | None = None) -> float:
    current = time.time() if now is None else now
    if not WGL_LATEST_REPORT_PATH.exists():
        return current
    try:
        payload = json.loads(WGL_LATEST_REPORT_PATH.read_text(encoding="utf-8"))
        generated_at = to_float(payload.get("generated_at"))
    except (OSError, ValueError, TypeError):
        return current
    if generated_at is None or generated_at > current + 300:
        return current
    return max(current, generated_at + report_interval_seconds())


def load_wgl_signal_states() -> dict[str, dict[str, Any]]:
    if not WGL_SIGNAL_STATES_PATH.exists():
        return {}
    try:
        payload = json.loads(WGL_SIGNAL_STATES_PATH.read_text(encoding="utf-8"))
    except Exception:
        return {}
    symbols = payload.get("symbols") if isinstance(payload, dict) else None
    return symbols if isinstance(symbols, dict) else {}


def format_trade_plan_alert(symbol: str, plan: dict[str, Any], *, source: str) -> str:
    decision = str(plan.get("trade_decision") or "不交易")
    if decision not in ACTIONABLE_DECISIONS:
        return ""
    return "\n".join(
        [
            f"交易計畫｜{symbol}",
            f"方向：{decision}｜信心：{int(to_float(plan.get('plan_confidence')) or 0)}/100",
            f"進場區($)：{format_plan_price(plan.get('entry_low'))} - {format_plan_price(plan.get('entry_high'))}",
            (
                f"TP1($)：{format_plan_price(plan.get('take_profit_1'))}｜"
                f"RR {float(plan.get('risk_reward_1') or 0):.2f}"
            ),
            (
                f"TP2($)：{format_plan_price(plan.get('take_profit_2'))}｜"
                f"RR {float(plan.get('risk_reward_2') or 0):.2f}"
            ),
            (
                f"SL($)：{format_plan_price(plan.get('stop_loss'))}｜"
                f"風險 {float(plan.get('stop_distance_pct') or 0):.2f}%"
            ),
            f"理由：{plan.get('plan_reason') or '-'}",
            f"倉位管理：{plan.get('plan_management') or '-'}",
            f"失效條件：{plan.get('plan_invalidation') or '-'}",
            f"來源：{source}",
        ]
    )


def compact_trade_plan(plan: dict[str, Any]) -> dict[str, Any]:
    numeric_fields = (
        "plan_confidence",
        "entry_low",
        "entry_high",
        "entry_mid",
        "take_profit_1",
        "take_profit_2",
        "stop_loss",
        "risk_reward_1",
        "risk_reward_2",
        "stop_distance_pct",
        "take_profit_1_pct",
        "take_profit_2_pct",
        "long_score",
        "short_score",
    )
    compact = {
        "trade_side": str(plan.get("trade_side") or "NONE"),
        "trade_decision": str(plan.get("trade_decision") or "不交易"),
        "plan_reason": str(plan.get("plan_reason") or ""),
        "plan_management": str(plan.get("plan_management") or ""),
        "plan_invalidation": str(plan.get("plan_invalidation") or ""),
    }
    compact.update({field: to_float(plan.get(field)) for field in numeric_fields})
    return compact


def update_wgl_signal_states(candidates: list[dict[str, Any]]) -> None:
    states = load_wgl_signal_states()
    transitions: list[dict[str, Any]] = []
    now_utc = time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime())
    now_local = time.strftime("%Y-%m-%d %H:%M")
    for item in candidates:
        symbol = str(item.get("symbol") or "").upper()
        state = str(item.get("signal_state") or "結構未成熟")
        if not symbol:
            continue
        previous = states.get(symbol) or {}
        previous_state = str(previous.get("state") or "")
        previous_decision = str(previous.get("trade_decision") or "不交易")
        record = {
            "state": state,
            "score": item.get("score"),
            "structure_score": item.get("structure_score"),
            "capital_score": item.get("capital_score"),
            "trigger_score": item.get("trigger_score"),
            "quality_score": item.get("quality_score"),
            "risk_score": item.get("risk_score"),
            "liquidity_status": item.get("liquidity_status"),
            "liquidity_score": item.get("liquidity_score"),
            "liquidity_ready": bool(item.get("liquidity_ready")),
            "liquidity_blocked": bool(item.get("liquidity_blocked")),
            "trade_decision": item.get("trade_decision") or "不交易",
            "trade_side": item.get("trade_side") or "NONE",
            "plan_confidence": item.get("plan_confidence"),
            "entry_low": item.get("entry_low"),
            "entry_high": item.get("entry_high"),
            "take_profit_1": item.get("take_profit_1"),
            "take_profit_2": item.get("take_profit_2"),
            "stop_loss": item.get("stop_loss"),
            "risk_reward_1": item.get("risk_reward_1"),
            "risk_reward_2": item.get("risk_reward_2"),
            "stop_distance_pct": item.get("stop_distance_pct"),
            "plan_reason": item.get("plan_reason"),
            "plan_management": item.get("plan_management"),
            "plan_invalidation": item.get("plan_invalidation"),
            "updated_utc": now_utc,
            "updated_local": now_local,
        }
        states[symbol] = record
        decision = str(record["trade_decision"])
        if previous_state != state or previous_decision != decision:
            transition = {
                "timestamp_utc": now_utc,
                "timestamp_local": now_local,
                "symbol": symbol,
                "from_state": previous_state or "未追蹤",
                "to_state": state,
                **record,
            }
            transitions.append(transition)
            if decision in ACTIONABLE_DECISIONS and previous_decision != decision:
                alert = format_trade_plan_alert(symbol, record, source=f"狀態 {previous_state or '未追蹤'} → {state}")
                with RUNTIME_TRANSITION_LOCK:
                    RUNTIME_TRANSITION_ALERTS.append(alert)

    payload = {
        "schema": "wgl-signal-state-v3",
        "updated_utc": now_utc,
        "updated_local": now_local,
        "symbols": states,
    }
    WGL_SIGNAL_STATES_PATH.parent.mkdir(parents=True, exist_ok=True)
    WGL_SIGNAL_STATES_PATH.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    if transitions:
        WGL_TRANSITIONS_PATH.mkdir(parents=True, exist_ok=True)
        path = WGL_TRANSITIONS_PATH / f"{time.strftime('%Y%m%d')}.jsonl"
        with path.open("a", encoding="utf-8") as file:
            for transition in transitions:
                file.write(json.dumps(transition, ensure_ascii=False, separators=(",", ":")) + "\n")


def consume_wgl_transition_alerts() -> list[str]:
    with RUNTIME_TRANSITION_LOCK:
        alerts = list(RUNTIME_TRANSITION_ALERTS)
        RUNTIME_TRANSITION_ALERTS.clear()
    return alerts


def update_wgl_signal_outcomes(candidates: list[dict[str, Any]], current_rows: list[dict[str, Any]]) -> None:
    payload: dict[str, Any] = {"schema": "wgl-signal-outcomes-v3", "signals": []}
    if WGL_SIGNAL_OUTCOMES_PATH.exists():
        try:
            loaded = json.loads(WGL_SIGNAL_OUTCOMES_PATH.read_text(encoding="utf-8"))
            if isinstance(loaded, dict) and isinstance(loaded.get("signals"), list):
                payload = loaded
        except Exception:
            pass
    signals = [item for item in payload.get("signals", []) if isinstance(item, dict)]
    price_by_symbol = {
        str(item["row"].symbol).upper(): to_float(getattr(item["row"], "mark_price", None))
        for item in current_rows
    }
    now = time.time()
    now_utc = time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime())

    for signal in signals:
        entry = to_float(signal.get("entry_price"))
        price = price_by_symbol.get(str(signal.get("symbol") or "").upper())
        if entry is None or price is None or entry <= 0:
            continue
        change = pct_change(price, entry)
        signal["last_price"] = price
        signal["last_return_pct"] = change
        signal["last_updated_utc"] = now_utc
        previous_mfe = to_float(signal.get("snapshot_mfe_pct"))
        previous_mae = to_float(signal.get("snapshot_mae_pct"))
        signal["snapshot_mfe_pct"] = max(previous_mfe if previous_mfe is not None else 0.0, change or 0.0)
        signal["snapshot_mae_pct"] = min(previous_mae if previous_mae is not None else 0.0, change or 0.0)
        if not signal.get("first_snapshot_hit"):
            if change is not None and change >= strategy_take_profit_pct():
                signal["first_snapshot_hit"] = "TP"
                signal["first_snapshot_hit_utc"] = now_utc
            elif change is not None and change <= -strategy_stop_loss_pct():
                signal["first_snapshot_hit"] = "SL"
                signal["first_snapshot_hit_utc"] = now_utc

    day_key = local_day_key()
    existing_keys = {
        (str(item.get("symbol") or "").upper(), str(item.get("signal_state") or ""), str(item.get("entry_date") or ""))
        for item in signals
    }
    for item in candidates:
        state = str(item.get("signal_state") or "")
        symbol = str(item.get("symbol") or "").upper()
        if state not in TRIGGER_STATES or (symbol, state, day_key) in existing_keys:
            continue
        row = item.get("row")
        entry = to_float(getattr(row, "mark_price", None))
        if not symbol or entry is None:
            continue
        signals.append(
            {
                "id": f"{compact_day_key()}-{symbol}-{state}",
                "symbol": symbol,
                "signal_state": state,
                "entry_date": day_key,
                "entry_ts": now,
                "entry_utc": now_utc,
                "entry_price": entry,
                "last_price": entry,
                "last_return_pct": 0.0,
                "snapshot_mfe_pct": 0.0,
                "snapshot_mae_pct": 0.0,
                "first_snapshot_hit": None,
                "score": item.get("score"),
                "structure_score": item.get("structure_score"),
                "capital_score": item.get("capital_score"),
                "trigger_score": item.get("trigger_score"),
                "quality_score": item.get("quality_score"),
                "risk_score": item.get("risk_score"),
            }
        )
        existing_keys.add((symbol, state, day_key))

    payload["updated_utc"] = now_utc
    payload["signals"] = signals[-5000:]
    WGL_SIGNAL_OUTCOMES_PATH.parent.mkdir(parents=True, exist_ok=True)
    WGL_SIGNAL_OUTCOMES_PATH.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def evaluate_signal_path_1m(signal: dict[str, Any], end_ts: float) -> dict[str, Any] | None:
    symbol = str(signal.get("symbol") or "").upper()
    entry = to_float(signal.get("entry_price"))
    start_ts = to_float(signal.get("entry_ts"))
    if not symbol or entry is None or entry <= 0 or start_ts is None or end_ts <= start_ts:
        return None
    start_ms = int(start_ts * 1000)
    end_ms = int(end_ts * 1000)
    rows: list[list[Any]] = []
    cursor = start_ms
    while cursor <= end_ms and len(rows) < 3000:
        payload = binance_market_json(
            "/fapi/v1/klines",
            {
                "symbol": symbol,
                "interval": "1m",
                "startTime": cursor,
                "endTime": end_ms,
                "limit": 1500,
            },
            timeout=20,
        )
        if not isinstance(payload, list) or not payload:
            break
        rows.extend(payload)
        next_cursor = int(payload[-1][0]) + 60_000
        if next_cursor <= cursor:
            break
        cursor = next_cursor
    if not rows:
        return None

    tp_pct = strategy_take_profit_pct()
    sl_pct = strategy_stop_loss_pct()
    first_hit = None
    first_hit_ts = None
    highs = []
    lows = []
    for row in rows:
        high = to_float(row[2])
        low = to_float(row[3])
        if high is None or low is None:
            continue
        highs.append(high)
        lows.append(low)
        up = pct_change(high, entry)
        down = pct_change(low, entry)
        if first_hit is None and up is not None and down is not None:
            if up >= tp_pct and down <= -sl_pct:
                first_hit = "同分鐘不確定"
                first_hit_ts = int(row[0]) / 1000.0
            elif up >= tp_pct:
                first_hit = "TP"
                first_hit_ts = int(row[0]) / 1000.0
            elif down <= -sl_pct:
                first_hit = "SL"
                first_hit_ts = int(row[0]) / 1000.0
    if not highs or not lows:
        return None
    final_price = to_float(rows[-1][4])
    return {
        "path_mfe_pct": pct_change(max(highs), entry),
        "path_mae_pct": pct_change(min(lows), entry),
        "path_last_return_pct": pct_change(final_price, entry),
        "path_first_hit": first_hit,
        "path_first_hit_utc": (
            time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime(first_hit_ts)) if first_hit_ts else None
        ),
        "path_evaluated_until_utc": time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime(end_ts)),
        "path_candle_count": len(rows),
    }


def wgl_seen_symbols_path() -> Path:
    return WGL_SYMBOL_STATS_PATH


def local_day_key(ts: float | None = None) -> str:
    return time.strftime("%Y-%m-%d", time.localtime(ts or time.time()))


def compact_day_key(day_key: str | None = None) -> str:
    return (day_key or local_day_key()).replace("-", "")


def migrate_legacy_wgl_seen_symbols() -> dict[str, dict[str, Any]]:
    symbols: dict[str, dict[str, Any]] = {}
    start_date = wgl_stats_start_date()
    if not WGL_SEEN_SYMBOLS_PATH.exists():
        return symbols
    for path in sorted(WGL_SEEN_SYMBOLS_PATH.glob("*.json")):
        try:
            payload = json.loads(path.read_text(encoding="utf-8"))
        except Exception:
            continue
        day_key = str(payload.get("date") or "")
        if not day_key and len(path.stem) == 8:
            day_key = f"{path.stem[:4]}-{path.stem[4:6]}-{path.stem[6:8]}"
        day_key = normalize_day_string(day_key)
        if start_date and day_key and day_key < start_date:
            continue
        rows = payload.get("symbols", {})
        if not isinstance(rows, dict):
            continue
        for symbol, row in rows.items():
            if not isinstance(row, dict):
                continue
            symbol_key = str(symbol).upper()
            first_seen_local = row.get("first_seen_local") or row.get("first_seen_utc") or day_key
            record = symbols.setdefault(
                symbol_key,
                {
                    "first_seen_utc": row.get("first_seen_utc"),
                    "first_seen_local": first_seen_local,
                    "first_seen_date": day_key,
                    "first_price": row.get("first_price"),
                    "first_direction": row.get("first_direction"),
                    "total_push_count": 0,
                    "days": {},
                },
            )
            record["total_push_count"] = int(to_float(record.get("total_push_count")) or 0) + int(
                to_float(row.get("push_count")) or 1
            )
            days = record.setdefault("days", {})
            days[day_key] = {
                "date": day_key,
                "first_seen_local": first_seen_local,
                "last_seen_local": row.get("last_seen_local") or first_seen_local,
                "first_price": row.get("first_price"),
                "last_price": row.get("last_price"),
                "push_count": int(to_float(row.get("push_count")) or 1),
                "best_score": row.get("score") or row.get("last_score"),
                "last_score": row.get("last_score") or row.get("score"),
                "last_trade_bucket": row.get("last_trade_bucket") or row.get("trade_bucket"),
                "last_trade_decision": row.get("last_trade_decision") or row.get("trade_decision"),
                "appearances": [],
            }
    return symbols


def local_minute_to_utc(value: Any) -> str | None:
    text = str(value or "").strip()
    try:
        parsed = time.strptime(text[:16], "%Y-%m-%d %H:%M")
    except ValueError:
        return None
    return time.strftime("%Y-%m-%dT%H:%M:00Z", time.gmtime(time.mktime(parsed)))


def normalize_wgl_symbol_stats(symbols: dict[str, dict[str, Any]]) -> dict[str, dict[str, Any]]:
    start_date = wgl_stats_start_date()
    if not start_date:
        return symbols
    normalized: dict[str, dict[str, Any]] = {}
    for symbol, raw_record in symbols.items():
        if not isinstance(raw_record, dict):
            continue
        days_raw = raw_record.get("days") if isinstance(raw_record.get("days"), dict) else {}
        days: dict[str, dict[str, Any]] = {}
        for raw_day_key, raw_day in days_raw.items():
            day_key = normalize_day_string(raw_day_key or (raw_day or {}).get("date"))
            if not day_key or day_key < start_date or not isinstance(raw_day, dict):
                continue
            day = dict(raw_day)
            day["date"] = day_key
            if day.get("first_seen_local"):
                day["first_seen_utc"] = local_minute_to_utc(day.get("first_seen_local")) or day.get("first_seen_utc")
            if day.get("last_seen_local"):
                day["last_seen_utc"] = local_minute_to_utc(day.get("last_seen_local")) or day.get("last_seen_utc")
            days[day_key] = day
        if not days:
            continue
        first_day_key = min(days)
        last_day_key = max(days)
        first_day = days[first_day_key]
        last_day = days[last_day_key]
        record = dict(raw_record)
        record["days"] = days
        record["first_seen_date"] = first_day_key
        record["first_seen_utc"] = first_day.get("first_seen_utc") or first_day.get("last_seen_utc")
        record["first_seen_local"] = first_day.get("first_seen_local") or first_day.get("last_seen_local") or first_day_key
        record["first_price"] = to_float(first_day.get("first_price")) or to_float(first_day.get("last_price"))
        record["first_direction"] = raw_record.get("first_direction") or "做多"
        record["last_seen_date"] = last_day_key
        record["last_seen_utc"] = last_day.get("last_seen_utc") or last_day.get("first_seen_utc")
        record["last_seen_local"] = last_day.get("last_seen_local") or last_day.get("first_seen_local")
        record["last_price"] = to_float(last_day.get("last_price")) or to_float(last_day.get("first_price"))
        record["last_score"] = last_day.get("last_score") or last_day.get("best_score")
        record["last_trade_bucket"] = last_day.get("last_trade_bucket")
        record["last_trade_decision"] = last_day.get("last_trade_decision")
        record["last_trade_setup"] = last_day.get("last_trade_setup")
        record["total_push_count"] = sum(
            int(to_float(day.get("push_count")) or len(day.get("appearances") or []) or 0)
            for day in days.values()
        )
        normalized[str(symbol).upper()] = record
    return normalized


def load_wgl_seen_symbols() -> dict[str, dict[str, Any]]:
    path = wgl_seen_symbols_path()
    if path.exists():
        try:
            data = json.loads(path.read_text(encoding="utf-8"))
        except Exception:
            return {}
    else:
        migrated = migrate_legacy_wgl_seen_symbols()
        if migrated:
            save_wgl_seen_symbols(migrated)
        return normalize_wgl_symbol_stats(migrated)
    if not isinstance(data, dict):
        return {}
    symbols = data.get("symbols", {})
    return normalize_wgl_symbol_stats(symbols) if isinstance(symbols, dict) else {}


def save_wgl_seen_symbols(seen: dict[str, dict[str, Any]]) -> None:
    path = wgl_seen_symbols_path()
    path.parent.mkdir(parents=True, exist_ok=True)
    payload = {
        "schema": "wgl-symbol-stats-v2",
        "updated_utc": time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime()),
        "updated_local": time.strftime("%Y-%m-%d %H:%M"),
        "symbols": seen,
    }
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def mark_wgl_report_seen(candidates: list[dict[str, Any]], seen: dict[str, dict[str, Any]]) -> None:
    now_utc = time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime())
    now_local = time.strftime("%Y-%m-%d %H:%M")
    day_key = local_day_key()
    changed = False
    for item in candidates:
        symbol = str(item.get("symbol") or "").upper()
        if not symbol:
            continue
        row = item.get("row")
        mark_price = to_float(getattr(row, "mark_price", None)) or to_float(getattr(row, "price", None))
        direction = wgl_card_mode(item)
        score = int(to_float(item.get("score")) or 0)
        record = seen.setdefault(
            symbol,
            {
                "first_seen_utc": now_utc,
                "first_seen_local": now_local,
                "first_seen_date": day_key,
                "first_price": mark_price,
                "first_direction": direction,
                "total_push_count": 0,
                "days": {},
            },
        )
        if not record.get("first_seen_date"):
            record["first_seen_date"] = day_key
        if record.get("first_price") is None:
            record["first_price"] = mark_price
        if not record.get("first_direction"):
            record["first_direction"] = direction

        record["total_push_count"] = int(
            to_float(record.get("total_push_count")) or to_float(record.get("push_count")) or 0
        ) + 1
        record["last_seen_utc"] = now_utc
        record["last_seen_local"] = now_local
        record["last_seen_date"] = day_key
        record["last_price"] = mark_price
        record["last_score"] = score
        record["last_trade_bucket"] = item.get("trade_bucket")
        record["last_trade_decision"] = item.get("trade_decision")
        record["last_trade_setup"] = item.get("trade_setup")

        days = record.setdefault("days", {})
        day = days.setdefault(
            day_key,
            {
                "date": day_key,
                "first_seen_utc": now_utc,
                "first_seen_local": now_local,
                "first_price": mark_price,
                "push_count": 0,
                "appearances": [],
            },
        )
        day["push_count"] = int(to_float(day.get("push_count")) or 0) + 1
        day["last_seen_utc"] = now_utc
        day["last_seen_local"] = now_local
        day["last_price"] = mark_price
        day["last_score"] = score
        day["last_trade_bucket"] = item.get("trade_bucket")
        day["last_trade_decision"] = item.get("trade_decision")
        day["last_trade_setup"] = item.get("trade_setup")
        day["best_score"] = max(int(to_float(day.get("best_score")) or score), score)
        day["worst_score"] = min(int(to_float(day.get("worst_score")) or score), score)
        if mark_price is not None:
            day["max_price"] = max(to_float(day.get("max_price")) or mark_price, mark_price)
            day["min_price"] = min(to_float(day.get("min_price")) or mark_price, mark_price)

        appearances = day.setdefault("appearances", [])
        if isinstance(appearances, list):
            appearances.append(
                {
                    "time_utc": now_utc,
                    "time_local": now_local,
                    "report_rank": item.get("report_rank"),
                    "symbol": symbol,
                    "score": score,
                    "signal_state": item.get("signal_state"),
                    "structure_score": item.get("structure_score"),
                    "capital_score": item.get("capital_score"),
                    "trigger_score": item.get("trigger_score"),
                    "quality_score": item.get("quality_score"),
                    "risk_score": item.get("risk_score"),
                    "liquidity_status": item.get("liquidity_status"),
                    "liquidity_score": item.get("liquidity_score"),
                    "liquidity_blocked": bool(item.get("liquidity_blocked")),
                    "quote_volume_24h_usd": to_float(item.get("quote_volume_24h_usd")),
                    "wgl_score": item.get("wgl_score"),
                    "trade_bucket": item.get("trade_bucket"),
                    "trade_decision": item.get("trade_decision"),
                    "trade_setup": item.get("trade_setup"),
                    "trade_reason": item.get("trade_reason"),
                    "trade_side": item.get("trade_side"),
                    "plan_confidence": item.get("plan_confidence"),
                    "entry_low": to_float(item.get("entry_low")),
                    "entry_high": to_float(item.get("entry_high")),
                    "take_profit_1": to_float(item.get("take_profit_1")),
                    "take_profit_2": to_float(item.get("take_profit_2")),
                    "stop_loss": to_float(item.get("stop_loss")),
                    "risk_reward_1": to_float(item.get("risk_reward_1")),
                    "risk_reward_2": to_float(item.get("risk_reward_2")),
                    "plan_reason": item.get("plan_reason"),
                    "price": mark_price,
                    "funding_rate_pct": to_float(getattr(row, "funding_rate_pct", None)),
                    "oi_value_usd": to_float(getattr(row, "oi_value_usd", None)),
                    "marketcap_usd": to_float(getattr(row, "marketcap_usd", None)),
                }
            )

        changed = True
    if changed:
        save_wgl_seen_symbols(seen)


def classify_spike_regime(price_pct: float | None, contracts_pct: float | None, confirm_pct: float) -> str:
    if contracts_pct is None:
        return "未知 OI 堆積"
    if price_pct is None:
        return "OI 堆積，價格未知"
    if price_pct >= confirm_pct and contracts_pct > 0:
        return "多頭建倉"
    if price_pct <= -confirm_pct and contracts_pct > 0:
        return "空頭建倉"
    if abs(price_pct) < confirm_pct and contracts_pct > 0:
        return "OI 堆積，方向未確認"
    return "混合訊號"


def classify_spike_grade(
    value_pct: float,
    contracts_pct: float | None,
    price_pct: float | None,
    min_value_pct: float,
    min_contracts_pct: float,
    confirm_pct: float,
) -> str:
    if contracts_pct is None:
        return "C"
    price_confirmed = price_pct is not None and abs(price_pct) >= confirm_pct
    if value_pct >= min_value_pct * 2 and contracts_pct >= min_contracts_pct * 2 and price_confirmed:
        return "A"
    if value_pct >= min_value_pct and contracts_pct >= min_contracts_pct:
        return "B"
    return "C"


def classify_oi_trend_signal(
    *,
    contracts_1h_pct: float | None,
    price_1h_pct: float | None,
    oi_value_usd: float | None,
    funding_rate_pct: float | None,
    structure: dict[str, Any] | None,
    quote_volume_24h_usd: float | None = None,
    min_quote_volume_24h_usd: float | None = None,
) -> dict[str, str] | None:
    if min_quote_volume_24h_usd is not None and (
        quote_volume_24h_usd is None
        or quote_volume_24h_usd < min_quote_volume_24h_usd
    ):
        return None
    if (
        contracts_1h_pct is None
        or price_1h_pct is None
        or oi_value_usd is None
        or oi_value_usd < spike_min_value_usd()
        or contracts_1h_pct <= 0
        or price_1h_pct <= 0
    ):
        return None
    if funding_rate_pct is not None and funding_rate_pct >= 0.10:
        return None

    screen = structure or {}
    structure_score = float(screen.get("score") or 0)
    base_days = int(screen.get("base_days") or 0)
    recent_extension = to_float(screen.get("recent_low_extension_pct"))
    bottom_structure = bool(
        structure_score >= 60
        and base_days >= 10
        and recent_extension is not None
        and recent_extension <= 80
    )
    bottom_trigger = bool(
        bottom_structure
        and contracts_1h_pct >= trend_min_contracts_pct()
        and trend_min_price_pct() <= price_1h_pct <= trend_max_bottom_price_pct()
    )
    momentum_trigger = bool(
        price_1h_pct <= trend_max_momentum_price_pct()
        and (
            (
                contracts_1h_pct >= momentum_min_contracts_pct()
                and price_1h_pct >= momentum_min_price_pct()
            )
            or (
                contracts_1h_pct >= momentum_strong_contracts_pct()
                and price_1h_pct >= momentum_strong_min_price_pct()
            )
        )
    )
    if not bottom_trigger and not momentum_trigger:
        return None

    short_squeeze = funding_rate_pct is not None and funding_rate_pct <= -0.10
    if bottom_trigger:
        signal_type = "軋空點火" if short_squeeze else "底部點火"
        action = "再確認：等待 5-15 分鐘回踩守住，價/OI未轉弱再考慮"
        lane = "bottom"
    else:
        signal_type = "軋空延續" if short_squeeze else "強勢延續"
        action = "只列入動能觀察：不是底部埋伏，禁止直接追價"
        lane = "momentum"
    return {"signal_type": signal_type, "action": action, "lane": lane}


def to_float(value: Any) -> float | None:
    if value is None:
        return None
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    if number != number:
        return None
    return number


def pct_change(new: float | None, old: float | None) -> float | None:
    if new is None or old is None or old == 0:
        return None
    return (new - old) / old * 100.0


def side_label(side: str) -> str:
    return "空" if side.upper() == "SHORT" else "多"


def position_pnl_pct(position: dict[str, Any], mark_price: float | None) -> float | None:
    entry_price = to_float(position.get("entry_price"))
    if entry_price is None or mark_price is None or entry_price <= 0:
        return None
    if str(position.get("side", "LONG")).upper() == "SHORT":
        return (entry_price - mark_price) / entry_price * 100.0
    return (mark_price - entry_price) / entry_price * 100.0


def position_history_key(position: dict[str, Any]) -> str:
    return str(position.get("symbol", "")).upper()


def current_position_sample(symbol: str) -> dict[str, Any]:
    mark = get_mark_price(symbol)
    current_oi = get_current_oi(symbol)
    mark_price = to_float(mark.get("markPrice"))
    funding_rate = to_float(mark.get("lastFundingRate"))
    open_interest = to_float(current_oi.get("openInterest")) if isinstance(current_oi, dict) else None
    return {
        "ts": time.time(),
        "symbol": symbol,
        "mark_price": mark_price,
        "open_interest": open_interest,
        "oi_value_usd": open_interest * mark_price if open_interest is not None and mark_price is not None else None,
        "funding_rate_pct": funding_rate * 100.0 if funding_rate is not None else None,
    }


def append_position_sample(
    history: dict[str, deque[dict[str, Any]]],
    symbol: str,
    sample: dict[str, Any],
) -> deque[dict[str, Any]]:
    samples = history.setdefault(symbol, deque())
    samples.append(sample)
    max_age = position_oi_window_seconds() + max(300, position_check_seconds() * 4)
    while samples and time.time() - float(samples[0]["ts"]) > max_age:
        samples.popleft()
    return samples


def baseline_sample(samples: deque[dict[str, Any]], window_seconds: int) -> dict[str, Any] | None:
    target_ts = time.time() - window_seconds
    baseline = None
    for old in samples:
        if float(old["ts"]) <= target_ts:
            baseline = old
        else:
            break
    return baseline


def market_context(symbol: str) -> WatchSymbol | None:
    try:
        for watch in resolve_watch_symbols():
            if watch.symbol == symbol:
                return watch
    except Exception:
        return None
    return None


def position_metrics(
    sample: dict[str, Any],
    samples: deque[dict[str, Any]],
    watch: WatchSymbol | None,
) -> dict[str, float | None]:
    baseline_180 = baseline_sample(samples, spike_window_seconds())
    baseline_1h = baseline_sample(samples, position_oi_window_seconds())
    oi_to_mcap = None
    if watch and watch.marketcap_usd and sample.get("oi_value_usd") is not None and watch.marketcap_usd > 0:
        oi_to_mcap = float(sample["oi_value_usd"]) / watch.marketcap_usd * 100.0
    return {
        "oi_to_marketcap_pct": oi_to_mcap,
        "contracts_180s_pct": pct_change(
            to_float(sample.get("open_interest")),
            to_float(baseline_180.get("open_interest")) if baseline_180 else None,
        ),
        "contracts_1h_pct": pct_change(
            to_float(sample.get("open_interest")),
            to_float(baseline_1h.get("open_interest")) if baseline_1h else None,
        ),
        "price_180s_pct": pct_change(
            to_float(sample.get("mark_price")),
            to_float(baseline_180.get("mark_price")) if baseline_180 else None,
        ),
        "price_1h_pct": pct_change(
            to_float(sample.get("mark_price")),
            to_float(baseline_1h.get("mark_price")) if baseline_1h else None,
        ),
    }


def format_position(position: dict[str, Any]) -> str:
    symbol = str(position.get("symbol", "")).upper()
    side = side_label(str(position.get("side", "LONG")))
    entry_price = to_float(position.get("entry_price"))
    last_price = to_float(position.get("last_price"))
    pnl = to_float(position.get("last_pnl_pct"))
    partial = "｜已半停利" if position.get("partial_taken") else ""
    return (
        f"{symbol}｜{side}｜進場 {fmt_num(entry_price, 6)}"
        f"｜現價 {fmt_num(last_price, 6)}｜PnL {fmt_pct(pnl, 2)}{partial}"
    )


def format_positions(chat_id: int) -> str:
    positions = [p for p in load_positions() if int(p.get("chat_id", 0)) == int(chat_id)]
    if not positions:
        return "目前沒有登記中的進場倉位。"
    lines = ["目前盯盤倉位："]
    lines.extend(format_position(position) for position in positions)
    lines.append("")
    lines.append(f"監控間隔：{position_check_seconds()} 秒")
    lines.append(f"停損：-{position_stop_loss_pct():.2f}%｜半停利：+{position_take_profit_pct():.2f}%")
    lines.append(f"Funding 出場：>{position_funding_exit_pct():.4f}% 且已獲利")
    return "\n".join(lines)


def close_position_alert(
    position: dict[str, Any],
    sample: dict[str, Any],
    metrics: dict[str, float | None],
    action: str,
    reason: str,
    pnl_pct: float | None,
    decision: str,
    decision_reason: str,
) -> str:
    symbol = str(position.get("symbol", "")).upper()
    return (
        f"倉位警報｜{symbol}\n"
        f"動作：{action}\n"
        f"原因：{reason}\n"
        f"方向：{side_label(str(position.get('side', 'LONG')))}｜進場 {fmt_num(to_float(position.get('entry_price')), 6)}"
        f"｜現價 {fmt_num(to_float(sample.get('mark_price')), 6)}｜PnL {fmt_pct(pnl_pct, 2)}\n"
        f"Funding：{fmt_pct(to_float(sample.get('funding_rate_pct')), 4)}"
        f"｜OI1h：{fmt_pct(metrics.get('contracts_1h_pct'), 2)}"
        f"｜OI180：{fmt_pct(metrics.get('contracts_180s_pct'), 2)}\n"
        f"目前判定：{decision}｜{decision_reason}"
    )


def position_event(
    position: dict[str, Any],
    sample: dict[str, Any],
    metrics: dict[str, float | None],
    action: str,
    reason: str,
    pnl_pct: float | None,
    decision: str,
) -> dict[str, Any]:
    return {
        "timestamp": time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime()),
        "chat_id": position.get("chat_id"),
        "symbol": position.get("symbol"),
        "side": position.get("side", "LONG"),
        "entry_price": position.get("entry_price"),
        "mark_price": sample.get("mark_price"),
        "pnl_pct": pnl_pct,
        "action": action,
        "reason": reason,
        "decision": decision,
        "funding_rate_pct": sample.get("funding_rate_pct"),
        "open_interest": sample.get("open_interest"),
        "oi_value_usd": sample.get("oi_value_usd"),
        "contracts_1h_pct": metrics.get("contracts_1h_pct"),
        "contracts_180s_pct": metrics.get("contracts_180s_pct"),
        "price_1h_pct": metrics.get("price_1h_pct"),
        "price_180s_pct": metrics.get("price_180s_pct"),
        "partial_taken": bool(position.get("partial_taken")),
    }


def collect_position_alerts(history: dict[str, deque[dict[str, Any]]]) -> list[tuple[int, str]]:
    positions = load_positions()
    if not positions:
        return []

    alerts: list[tuple[int, str]] = []
    remaining: list[dict[str, Any]] = []
    changed = False
    sample_cache: dict[str, dict[str, Any]] = {}
    context_cache: dict[str, WatchSymbol | None] = {}

    for position in positions:
        symbol = str(position.get("symbol", "")).upper()
        chat_id = int(position.get("chat_id", 0))
        if not symbol or not chat_id:
            changed = True
            continue

        try:
            if symbol not in sample_cache:
                sample_cache[symbol] = current_position_sample(symbol)
            sample = sample_cache[symbol]
        except Exception as exc:
            position["last_error"] = str(exc)
            remaining.append(position)
            changed = True
            continue

        samples = append_position_sample(history, symbol, sample)
        watch = context_cache.setdefault(symbol, market_context(symbol))
        metrics = position_metrics(sample, samples, watch)
        rank = watch.market_rank if watch else None
        row = type("PositionRow", (), {"funding_rate_pct": sample.get("funding_rate_pct"), "market_rank": rank})()
        score = attention_score(row, metrics)
        decision, decision_reason = report_decision(row, metrics, score)
        pnl_pct = position_pnl_pct(position, to_float(sample.get("mark_price")))

        position["last_price"] = sample.get("mark_price")
        position["last_pnl_pct"] = pnl_pct
        position["last_funding_rate_pct"] = sample.get("funding_rate_pct")
        position["last_open_interest"] = sample.get("open_interest")
        position["last_checked_at"] = sample.get("ts")
        position["last_decision"] = decision

        side = str(position.get("side", "LONG")).upper()
        funding = to_float(sample.get("funding_rate_pct"))
        oi_1h_pct = metrics.get("contracts_1h_pct")
        action = None
        reason = None
        remove_position = False

        if decision == "不要進":
            action = "立即出場"
            reason = f"倉位判定轉為不要進：{decision_reason}"
            remove_position = True
        elif pnl_pct is not None and position.get("partial_taken") and pnl_pct <= 0:
            action = "剩餘半倉出場"
            reason = "已半停利，價格回到開倉價，保本出剩餘倉"
            remove_position = True
        elif pnl_pct is not None and pnl_pct <= -position_stop_loss_pct():
            action = "停損出場"
            reason = f"虧損達 -{position_stop_loss_pct():.2f}%"
            remove_position = True
        elif funding is not None and pnl_pct is not None and pnl_pct > 0 and side == "LONG" and funding > position_funding_exit_pct():
            action = "立即出場"
            reason = f"Funding > {position_funding_exit_pct():.4f}% 且價格已漲"
            remove_position = True
        elif funding is not None and pnl_pct is not None and pnl_pct > 0 and side == "SHORT" and funding < -position_funding_exit_pct():
            action = "立即出場"
            reason = f"空單 funding < -{position_funding_exit_pct():.4f}% 且已獲利"
            remove_position = True
        elif oi_1h_pct is not None and oi_1h_pct < 0:
            action = "立即出場"
            reason = "OI 1小時轉負"
            remove_position = True
        elif pnl_pct is not None and not position.get("partial_taken") and pnl_pct >= position_take_profit_pct():
            action = "先停利一半"
            reason = f"獲利達 +{position_take_profit_pct():.2f}%，剩餘倉止損拉到開倉價"
            position["partial_taken"] = True
            position["break_even_stop"] = True
            changed = True

        if action and reason:
            message = close_position_alert(position, sample, metrics, action, reason, pnl_pct, decision, decision_reason)
            alerts.append((chat_id, message))
            save_position_event(position_event(position, sample, metrics, action, reason, pnl_pct, decision))
            if remove_position:
                changed = True
                continue

        remaining.append(position)

    if changed:
        save_positions(remaining)
    return alerts


def latest_history_sample(samples: deque[dict[str, Any]]) -> dict[str, Any] | None:
    if not samples:
        return None
    sample = samples[-1]
    if time.time() - float(sample.get("ts", 0)) > max(300, spike_check_seconds() * 4):
        return None
    return sample


TRADE_PLAN_FINAL_ACTIONS = {"PLAN_TP2", "PLAN_SL", "PLAN_BE"}


def is_trade_plan_position(position: dict[str, Any]) -> bool:
    return str(position.get("strategy_model") or "") == "trade_plan_v1"


def directional_return_pct(side: str, exit_price: float | None, entry_price: float | None) -> float | None:
    if entry_price is None or exit_price is None or entry_price <= 0:
        return None
    raw = (exit_price - entry_price) / entry_price * 100.0
    return -raw if str(side).upper() == "SHORT" else raw


def strategy_pnl_pct(position: dict[str, Any], mark_price: float | None) -> float | None:
    entry_price = to_float(position.get("entry_price"))
    return directional_return_pct(str(position.get("side") or "LONG"), mark_price, entry_price)


def strategy_trade_id(symbol: str) -> str:
    return f"{symbol}-{int(time.time())}"


def strategy_event_timestamp(event_ts: float) -> str:
    return time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime(event_ts))


def trade_plan_strategy_event(
    position: dict[str, Any],
    action: str,
    *,
    event_ts: float,
    fill_price: float,
    pnl_pct: float,
    note: str = "",
) -> dict[str, Any]:
    leverage = to_float(position.get("leverage")) or trade_plan_paper_leverage()
    margin_usd = to_float(position.get("margin_usd")) or trade_plan_paper_margin_usd()
    leveraged_pnl_pct = pnl_pct * leverage
    return {
        "timestamp": strategy_event_timestamp(event_ts),
        "ts": event_ts,
        "action": action,
        "id": position.get("id"),
        "strategy_model": "trade_plan_v1",
        "symbol": position.get("symbol"),
        "side": position.get("side"),
        "entry_phase": position.get("entry_phase"),
        "entry_reason": position.get("entry_reason"),
        "entry_price": position.get("entry_price"),
        "entry_low": position.get("entry_low"),
        "entry_high": position.get("entry_high"),
        "take_profit_1": position.get("take_profit_1"),
        "take_profit_2": position.get("take_profit_2"),
        "initial_stop_loss": position.get("initial_stop_loss"),
        "stop_loss": position.get("stop_loss"),
        "mark_price": fill_price,
        "fill_price": fill_price,
        "pnl_pct": pnl_pct,
        "leverage": leverage,
        "leveraged_pnl_pct": leveraged_pnl_pct,
        "margin_usd": margin_usd,
        "pnl_usd": margin_usd * leveraged_pnl_pct / 100.0,
        "remaining_fraction": position.get("remaining_fraction"),
        "tp1_hit": bool(position.get("tp1_hit")),
        "note": note,
    }


def trade_plan_level_hit(side: str, *, high: float, low: float, level: float, target: bool) -> bool:
    if str(side).upper() == "SHORT":
        return low <= level if target else high >= level
    return high >= level if target else low <= level


def evaluate_trade_plan_candles(
    position: dict[str, Any],
    klines: list[list[Any]],
) -> tuple[list[dict[str, Any]], bool]:
    side = str(position.get("side") or "LONG").upper()
    entry = to_float(position.get("entry_price"))
    tp1 = to_float(position.get("take_profit_1"))
    tp2 = to_float(position.get("take_profit_2"))
    stop = to_float(position.get("stop_loss"))
    if None in {entry, tp1, tp2, stop}:
        return [], False

    events: list[dict[str, Any]] = []
    tp1_hit = bool(position.get("tp1_hit"))
    realized = to_float(position.get("realized_return_pct")) or 0.0
    closed = False
    for row in sorted(klines, key=lambda item: int(item[0])):
        if len(row) < 5:
            continue
        high = to_float(row[2])
        low = to_float(row[3])
        if high is None or low is None:
            continue
        event_ts = (to_float(row[6]) or (to_float(row[0]) or 0) + 59_999) / 1000.0
        position["last_evaluated_ms"] = int(row[0]) + 60_000

        if not tp1_hit:
            stop_hit = trade_plan_level_hit(side, high=high, low=low, level=stop, target=False)
            tp1_now = trade_plan_level_hit(side, high=high, low=low, level=tp1, target=True)
            if stop_hit:
                pnl = directional_return_pct(side, stop, entry) or 0.0
                note = "同一分鐘亦觸及TP1，依保守原則先計SL" if tp1_now else "初始止損"
                position["remaining_fraction"] = 0.0
                events.append(
                    trade_plan_strategy_event(
                        position,
                        "PLAN_SL",
                        event_ts=event_ts,
                        fill_price=stop,
                        pnl_pct=pnl,
                        note=note,
                    )
                )
                closed = True
                break
            if tp1_now:
                tp1_hit = True
                position["tp1_hit"] = True
                position["remaining_fraction"] = 0.5
                realized = 0.5 * (directional_return_pct(side, tp1, entry) or 0.0)
                position["realized_return_pct"] = realized
                position["stop_loss"] = entry
                stop = entry
                events.append(
                    trade_plan_strategy_event(
                        position,
                        "PLAN_TP1",
                        event_ts=event_ts,
                        fill_price=tp1,
                        pnl_pct=realized,
                        note="停利一半，剩餘止損移到進場價",
                    )
                )
                if trade_plan_level_hit(side, high=high, low=low, level=tp2, target=True):
                    total = realized + 0.5 * (directional_return_pct(side, tp2, entry) or 0.0)
                    position["remaining_fraction"] = 0.0
                    events.append(
                        trade_plan_strategy_event(
                            position,
                            "PLAN_TP2",
                            event_ts=event_ts,
                            fill_price=tp2,
                            pnl_pct=total,
                            note="同一分鐘依序穿越TP1與TP2",
                        )
                    )
                    closed = True
                    break
                continue

        tp2_hit = trade_plan_level_hit(side, high=high, low=low, level=tp2, target=True)
        breakeven_hit = trade_plan_level_hit(side, high=high, low=low, level=stop, target=False)
        if tp2_hit and breakeven_hit:
            position["remaining_fraction"] = 0.0
            events.append(
                trade_plan_strategy_event(
                    position,
                    "PLAN_BE",
                    event_ts=event_ts,
                    fill_price=stop,
                    pnl_pct=realized,
                    note="同一分鐘觸及TP2與成本止損，依保守原則計成本出場",
                )
            )
            closed = True
            break
        if tp2_hit:
            total = realized + 0.5 * (directional_return_pct(side, tp2, entry) or 0.0)
            position["remaining_fraction"] = 0.0
            events.append(
                trade_plan_strategy_event(
                    position,
                    "PLAN_TP2",
                    event_ts=event_ts,
                    fill_price=tp2,
                    pnl_pct=total,
                    note="TP2出清剩餘半倉",
                )
            )
            closed = True
            break
        if breakeven_hit:
            position["remaining_fraction"] = 0.0
            events.append(
                trade_plan_strategy_event(
                    position,
                    "PLAN_BE",
                    event_ts=event_ts,
                    fill_price=stop,
                    pnl_pct=realized,
                    note="TP1後剩餘半倉於成本出場",
                )
            )
            closed = True
            break

    return events, closed


def trade_plan_klines_since(position: dict[str, Any], *, now: float | None = None) -> list[list[Any]]:
    current = time.time() if now is None else now
    start_ms = int(to_float(position.get("last_evaluated_ms")) or 0)
    if start_ms <= 0:
        entry_ts = to_float(position.get("entry_ts")) or current
        start_ms = (int(entry_ts * 1000) // 60_000 + 1) * 60_000
    end_ms = int(current * 1000)
    output: list[list[Any]] = []
    for _ in range(20):
        if start_ms >= end_ms:
            break
        rows = binance_market_json(
            "/fapi/v1/klines",
            {
                "symbol": position.get("symbol"),
                "interval": "1m",
                "startTime": start_ms,
                "endTime": end_ms,
                "limit": 1500,
            },
        )
        if not isinstance(rows, list) or not rows:
            break
        closed_rows = [row for row in rows if len(row) > 6 and int(row[6]) <= end_ms]
        output.extend(closed_rows)
        next_start = int(rows[-1][0]) + 60_000
        if next_start <= start_ms:
            break
        start_ms = next_start
        if len(rows) < 1500:
            break
    return output


def latest_trade_plan_entries(
    positions: list[dict[str, Any]],
    events: list[dict[str, Any]],
    *,
    now: float | None = None,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[str]]:
    current = time.time() if now is None else now
    if not WGL_LATEST_REPORT_PATH.exists():
        return positions, [], []
    try:
        payload = json.loads(WGL_LATEST_REPORT_PATH.read_text(encoding="utf-8"))
    except Exception:
        return positions, [], []
    generated_at = to_float(payload.get("generated_at"))
    items = payload.get("items")
    if (
        generated_at is None
        or generated_at > current + 60
        or current - generated_at > trade_plan_signal_max_age_seconds()
        or not isinstance(items, list)
    ):
        return positions, [], []

    existing_ids = {str(item.get("id")) for item in events + positions if item.get("id")}
    open_symbols = {
        str(item.get("symbol") or "").upper()
        for item in positions
        if is_trade_plan_position(item)
    }
    latest_final_by_symbol: dict[str, float] = {}
    for event in events:
        if event.get("action") not in TRADE_PLAN_FINAL_ACTIONS:
            continue
        symbol = str(event.get("symbol") or "").upper()
        latest_final_by_symbol[symbol] = max(
            latest_final_by_symbol.get(symbol, 0.0),
            to_float(event.get("ts")) or 0.0,
        )

    capacity = max(
        0,
        trade_plan_max_open_positions()
        - sum(1 for item in positions if is_trade_plan_position(item)),
    )
    new_events: list[dict[str, Any]] = []
    alerts: list[str] = []
    for item in items:
        if capacity <= 0 or not isinstance(item, dict):
            break
        decision = str(item.get("trade_decision") or "不交易")
        side = str(item.get("trade_side") or "NONE").upper()
        symbol = str(item.get("symbol") or "").upper()
        trade_id = f"PLAN-{symbol}-{int(generated_at)}"
        entry = to_float(item.get("entry_mid"))
        entry_low = to_float(item.get("entry_low"))
        entry_high = to_float(item.get("entry_high"))
        tp1 = to_float(item.get("take_profit_1"))
        tp2 = to_float(item.get("take_profit_2"))
        stop = to_float(item.get("stop_loss"))
        if (
            decision not in ACTIONABLE_DECISIONS
            or side not in {"LONG", "SHORT"}
            or not symbol
            or trade_id in existing_ids
            or symbol in open_symbols
            or None in {entry, entry_low, entry_high, tp1, tp2, stop}
        ):
            continue
        last_final = latest_final_by_symbol.get(symbol, 0.0)
        if last_final and generated_at - last_final < trade_plan_reentry_cooldown_seconds():
            continue
        levels_valid = (
            stop < entry < tp1 < tp2
            if side == "LONG"
            else tp2 < tp1 < entry < stop
        )
        if not levels_valid:
            continue
        position = {
            "id": trade_id,
            "strategy_model": "trade_plan_v1",
            "symbol": symbol,
            "side": side,
            "entry_ts": generated_at,
            "entry_time": strategy_event_timestamp(generated_at),
            "entry_price": entry,
            "entry_low": entry_low,
            "entry_high": entry_high,
            "entry_phase": f"{decision}計畫",
            "entry_reason": item.get("plan_reason"),
            "plan_confidence": item.get("plan_confidence"),
            "take_profit_1": tp1,
            "take_profit_2": tp2,
            "initial_stop_loss": stop,
            "stop_loss": stop,
            "risk_reward_1": item.get("risk_reward_1"),
            "risk_reward_2": item.get("risk_reward_2"),
            "margin_usd": trade_plan_paper_margin_usd(),
            "leverage": trade_plan_paper_leverage(),
            "remaining_fraction": 1.0,
            "tp1_hit": False,
            "realized_return_pct": 0.0,
            "last_price": entry,
            "last_pnl_pct": 0.0,
            "source_report_local": payload.get("generated_local"),
        }
        positions.append(position)
        open_symbols.add(symbol)
        existing_ids.add(trade_id)
        event = trade_plan_strategy_event(
            position,
            "PLAN_OPEN",
            event_ts=generated_at,
            fill_price=entry,
            pnl_pct=0.0,
            note="按每小時報告訊號價模擬成交",
        )
        new_events.append(event)
        alerts.append(
            f"策略開單｜{symbol}\n"
            f"方向：{decision}｜進場：{fmt_num(entry, 6)}\n"
            f"TP1：{fmt_num(tp1, 6)}｜TP2：{fmt_num(tp2, 6)}｜SL：{fmt_num(stop, 6)}\n"
            f"模擬：{fmt_num(position['margin_usd'])}U × {fmt_num(position['leverage'])}倍｜"
            f"信心 {int(to_float(position.get('plan_confidence')) or 0)}/100"
        )
        capacity -= 1
    return positions, new_events, alerts


def strategy_phase_result(symbol: str, history: dict[str, deque[dict[str, Any]]]) -> Any | None:
    row = type("StrategyRow", (), {"symbol": symbol})()
    samples = phase_samples(row, history)
    if len(samples) < 3:
        return None
    if float(samples[-1]["ts"]) - float(samples[0]["ts"]) < strategy_min_history_seconds():
        return None
    return classify_phase(samples)


def list_median(values: list[float]) -> float | None:
    clean = sorted(value for value in values if value is not None)
    if not clean:
        return None
    mid = len(clean) // 2
    if len(clean) % 2:
        return clean[mid]
    return (clean[mid - 1] + clean[mid]) / 2.0


def kline_float(row: list[Any], idx: int) -> float | None:
    if idx >= len(row):
        return None
    return to_float(row[idx])


def closed_klines(rows: list[list[Any]]) -> list[list[Any]]:
    now_ms = int(time.time() * 1000)
    return [row for row in rows if len(row) > 6 and int(row[6]) <= now_ms]


def daily_average(values: list[float | None]) -> float | None:
    clean = [value for value in values if value is not None]
    if not clean:
        return None
    return sum(clean) / len(clean)


def daily_strategy_setup(
    symbol: str,
    phase_result: Any,
    sample: dict[str, Any],
    watch: WatchSymbol | None,
) -> tuple[bool, str, int]:
    if symbol in strategy_blocklist():
        return False, "symbol blocklist: not a long daily setup target", 0

    marketcap = watch.marketcap_usd if watch else None
    oi_value = to_float(sample.get("oi_value_usd"))
    oi_to_mcap = oi_value / marketcap * 100.0 if oi_value is not None and marketcap and marketcap > 0 else None

    funding = to_float(sample.get("funding_rate_pct"))
    if funding is None or abs(funding) > strategy_max_funding_pct():
        return False, f"funding too hot: {fmt_pct(funding, 4)}", 0

    daily_klines = closed_klines(get_klines(symbol, interval="1d", limit=strategy_daily_lookback_days()))
    if len(daily_klines) < strategy_min_daily_candles():
        return False, f"not enough daily candles: {len(daily_klines)}", 0

    closes = [kline_float(row, 4) for row in daily_klines]
    highs = [kline_float(row, 2) for row in daily_klines]
    lows = [kline_float(row, 3) for row in daily_klines]
    quote_volumes = [kline_float(row, 7) for row in daily_klines]
    valid_closes = [value for value in closes if value is not None]
    valid_highs = [value for value in highs if value is not None]
    valid_lows = [value for value in lows if value is not None]
    if not valid_closes or not valid_highs or not valid_lows or closes[-1] is None:
        return False, "bad daily kline data", 0

    last_price = closes[-1]
    low_lookback = min(valid_lows)
    high_lookback = max(valid_highs)
    range_size = high_lookback - low_lookback
    if low_lookback <= 0 or high_lookback <= low_lookback:
        return False, "bad daily range", 0
    range_multiple = high_lookback / low_lookback
    if range_multiple < strategy_min_daily_range_multiple():
        return False, f"daily range too small: {range_multiple:.2f}x", 0

    extension_from_low = pct_change(last_price, low_lookback)
    range_position = (last_price - low_lookback) / range_size * 100.0
    recent_lows_60 = [value for value in lows[-60:] if value is not None]
    recent_low_60 = min(recent_lows_60) if recent_lows_60 else low_lookback
    recent_extension_from_low = pct_change(last_price, recent_low_60)
    drawdown_from_high = (high_lookback - last_price) / high_lookback * 100.0
    price_24h_pct = pct_change(last_price, closes[-2] if len(closes) >= 2 else None)
    price_3d_pct = pct_change(last_price, closes[-4] if len(closes) >= 4 else None)
    price_7d_pct = pct_change(last_price, closes[-8] if len(closes) >= 8 else None)
    high_14d = max(value for value in highs[-14:] if value is not None)
    low_14d = min(value for value in lows[-14:] if value is not None)
    range_14d_pct = pct_change(high_14d, low_14d)

    if extension_from_low is None or extension_from_low > strategy_max_daily_extension_from_low_pct():
        return False, f"daily too far from base low: {fmt_pct(extension_from_low)}", 0
    bottom_range_limit = min(strategy_max_daily_range_position_pct(), strategy_max_bottom_range_position_pct())
    if drawdown_from_high < strategy_min_drawdown_from_high_pct():
        return False, f"not deep enough from prior high: {fmt_pct(drawdown_from_high)}", 0
    if range_position > bottom_range_limit:
        return False, f"daily range position too high: {fmt_pct(range_position)}", 0
    if recent_extension_from_low is None or recent_extension_from_low > strategy_max_recent_low_extension_pct():
        return False, f"too far from recent 60d low: {fmt_pct(recent_extension_from_low)}", 0
    if price_24h_pct is None or price_24h_pct > strategy_max_24h_price_pct() or price_24h_pct < -18:
        return False, f"daily candle too stretched/weak: {fmt_pct(price_24h_pct)}", 0
    if price_3d_pct is not None and price_3d_pct > strategy_max_3d_price_pct():
        return False, f"3d price already stretched: {fmt_pct(price_3d_pct)}", 0
    if price_7d_pct is not None and price_7d_pct > strategy_max_7d_extension_pct():
        return False, f"7d price already stretched: {fmt_pct(price_7d_pct)}", 0
    if range_14d_pct is not None and range_14d_pct > strategy_max_14d_range_pct():
        return False, f"14d range too wide: {fmt_pct(range_14d_pct)}", 0

    base_range_position_limit = strategy_max_base_band_from_low_pct()
    base_days = sum(
        1
        for close in closes[-30:]
        if close is not None and ((close - low_lookback) / range_size * 100.0) <= base_range_position_limit
    )
    if base_days < strategy_min_base_days():
        return False, f"not enough daily base days: {base_days}", 0

    latest_volume = daily_average(quote_volumes[-3:])
    baseline_volume = list_median([value for value in quote_volumes[-33:-3] if value is not None])
    if baseline_volume is None:
        baseline_volume = list_median([value for value in quote_volumes[:-3] if value is not None])
    volume_ratio = latest_volume / baseline_volume if latest_volume and baseline_volume and baseline_volume > 0 else None
    if volume_ratio is None or volume_ratio < strategy_min_daily_volume_ratio():
        return False, f"daily volume not expanding: x{volume_ratio:.2f}" if volume_ratio else "missing daily volume baseline", 0
    if volume_ratio > strategy_max_daily_volume_spike_ratio():
        return False, f"daily volume already blow-off: x{volume_ratio:.2f}", 0

    oi_history = get_oi_history(symbol, period="1d", limit=30)
    if len(oi_history) < 3:
        return False, "not enough daily OI history", 0
    oi_values = [
        to_float(row.get("sumOpenInterestValue")) or to_float(row.get("sumOpenInterest"))
        for row in oi_history
    ]
    if oi_values[-1] is None:
        return False, "bad daily OI history", 0
    oi_24h_pct = pct_change(oi_values[-1], oi_values[-2] if len(oi_values) >= 2 else None)
    oi_2d_pct = pct_change(oi_values[-1], oi_values[-3] if len(oi_values) >= 3 else None)
    if (oi_24h_pct is None or oi_24h_pct < strategy_min_daily_oi_24h_pct()) and (
        oi_2d_pct is None or oi_2d_pct < strategy_min_daily_oi_2d_pct()
    ):
        return False, f"daily OI not building: 1d {fmt_pct(oi_24h_pct)}, 2d {fmt_pct(oi_2d_pct)}", 0

    four_h = closed_klines(get_klines(symbol, interval="4h", limit=180))
    if len(four_h) < 60:
        return False, f"not enough 4h candles: {len(four_h)}", 0
    h4_closes = [kline_float(row, 4) for row in four_h]
    h4_highs = [kline_float(row, 2) for row in four_h]
    h4_lows = [kline_float(row, 3) for row in four_h]
    h4_volumes = [kline_float(row, 7) for row in four_h]
    if h4_closes[-1] is None:
        return False, "bad 4h kline data", 0
    h4_last = h4_closes[-1]
    h4_low_90 = min(value for value in h4_lows[-90:] if value is not None)
    h4_high_90 = max(value for value in h4_highs[-90:] if value is not None)
    h4_low_30 = min(value for value in h4_lows[-30:] if value is not None)
    h4_low_prev = min(value for value in h4_lows[-90:-30] if value is not None)
    h4_high_30 = max(value for value in h4_highs[-30:] if value is not None)
    h4_range_90 = pct_change(h4_high_90, h4_low_90)
    h4_range_30 = pct_change(h4_high_30, h4_low_30)
    h4_position = (h4_last - h4_low_90) / (h4_high_90 - h4_low_90) * 100.0 if h4_high_90 > h4_low_90 else None
    h4_p24 = pct_change(h4_last, h4_closes[-7] if len(h4_closes) >= 7 else None)
    h4_p3d = pct_change(h4_last, h4_closes[-19] if len(h4_closes) >= 19 else None)
    h4_volume = daily_average(h4_volumes[-6:])
    h4_volume_base = list_median([value for value in h4_volumes[-78:-6] if value is not None])
    h4_volume_ratio = h4_volume / h4_volume_base if h4_volume and h4_volume_base and h4_volume_base > 0 else None
    h4_higher_low = h4_low_30 >= h4_low_prev * 0.92 if h4_low_prev else False
    h4_reclaim = h4_last >= (list_median([value for value in h4_closes[-30:] if value is not None]) or h4_last)
    h4_compressed = h4_range_30 is not None and h4_range_90 is not None and h4_range_30 <= h4_range_90 * 0.75

    if h4_position is None or h4_position > strategy_max_4h_range_position_pct():
        return False, f"4h position too high for bottom: {fmt_pct(h4_position)}", 0
    if h4_p24 is not None and h4_p24 > strategy_max_4h_24h_price_pct():
        return False, f"4h 24h price already stretched: {fmt_pct(h4_p24)}", 0
    if h4_p24 is not None and h4_p24 < strategy_min_4h_24h_price_pct():
        return False, f"4h 24h still falling: {fmt_pct(h4_p24)}", 0
    if h4_p3d is not None and h4_p3d > strategy_max_4h_3d_price_pct():
        return False, f"4h 3d price already stretched: {fmt_pct(h4_p3d)}", 0
    if h4_p3d is not None and h4_p3d < strategy_min_4h_3d_price_pct():
        return False, f"4h 3d still falling: {fmt_pct(h4_p3d)}", 0
    if not (h4_higher_low or h4_reclaim or h4_compressed):
        return False, "4h bottom structure not formed", 0

    score = int(getattr(phase_result, "confidence", 50))
    if range_position <= bottom_range_limit:
        score += 12
    if drawdown_from_high >= strategy_min_drawdown_from_high_pct() + 15:
        score += 10
    if recent_extension_from_low is not None and recent_extension_from_low <= strategy_max_recent_low_extension_pct() * 0.6:
        score += 8
    if base_days >= strategy_min_base_days() + 5:
        score += 8
    if oi_24h_pct is not None and oi_24h_pct >= strategy_min_daily_oi_24h_pct() * 2:
        score += 8
    if volume_ratio >= strategy_min_daily_volume_ratio() * 1.8:
        score += 7
    if h4_position is not None and h4_position <= strategy_max_4h_range_position_pct() * 0.75:
        score += 8
    if h4_higher_low:
        score += 8
    if h4_reclaim:
        score += 6
    if h4_compressed:
        score += 6
    reason = (
        f"日線長底部 | 無市值限制 | "
        f"歷史區間 {range_multiple:.1f}x | 高點回撤{fmt_pct(drawdown_from_high)} | "
        f"低點+{fmt_pct(extension_from_low)} | 近低+{fmt_pct(recent_extension_from_low)} | 區間{fmt_pct(range_position)} | "
        f"P1d {fmt_pct(price_24h_pct)} | P3d {fmt_pct(price_3d_pct)} | "
        f"base {base_days}d | OI1d {fmt_pct(oi_24h_pct)} | OI2d {fmt_pct(oi_2d_pct)} | "
        f"Vol x{volume_ratio:.2f} | 4H位階{fmt_pct(h4_position)} | 4H24 {fmt_pct(h4_p24)} | 4H3d {fmt_pct(h4_p3d)}"
    )
    return True, reason, min(score, 100)


def launch_structure_setup(
    symbol: str,
    marketcap: float | None,
    funding: float | None,
    oi_to_mcap: float | None,
    market_rank: int | None = None,
) -> dict[str, Any] | None:
    symbol = str(symbol).upper()
    if symbol in strategy_blocklist():
        return None
    if funding is None or abs(funding) > strategy_max_funding_pct():
        return None

    daily = closed_klines(get_klines(symbol, interval="1d", limit=strategy_daily_lookback_days()))
    if len(daily) < strategy_min_daily_candles():
        return None
    closes = [kline_float(item, 4) for item in daily]
    highs = [kline_float(item, 2) for item in daily]
    lows = [kline_float(item, 3) for item in daily]
    volumes = [kline_float(item, 7) for item in daily]
    if closes[-1] is None:
        return None
    valid_highs = [value for value in highs if value is not None]
    valid_lows = [value for value in lows if value is not None]
    if not valid_highs or not valid_lows:
        return None

    low = min(valid_lows)
    high = max(valid_highs)
    if low <= 0 or high <= low:
        return None
    last = closes[-1]
    range_size = high - low
    range_multiple = high / low
    range_position = (last - low) / range_size * 100.0
    extension_from_low = pct_change(last, low)
    recent_lows_60 = [value for value in lows[-60:] if value is not None]
    recent_low_60 = min(recent_lows_60) if recent_lows_60 else low
    recent_extension_from_low = pct_change(last, recent_low_60)
    drawdown_from_high = (high - last) / high * 100.0
    p1d = pct_change(last, closes[-2] if len(closes) >= 2 else None)
    p3d = pct_change(last, closes[-4] if len(closes) >= 4 else None)
    p7d = pct_change(last, closes[-8] if len(closes) >= 8 else None)
    high_14d = max(value for value in highs[-14:] if value is not None)
    low_14d = min(value for value in lows[-14:] if value is not None)
    range_14d_pct = pct_change(high_14d, low_14d)
    launch_base_band = max(strategy_max_base_band_from_low_pct(), 28.0)
    base_days = sum(
        1
        for close in closes[-45:]
        if close is not None and ((close - low) / range_size * 100.0) <= launch_base_band
    )
    daily_volume = avg_clean(volumes[-3:])
    daily_volume_base = list_median([value for value in volumes[-63:-3] if value is not None])
    daily_volume_ratio = daily_volume / daily_volume_base if daily_volume and daily_volume_base and daily_volume_base > 0 else None

    if range_multiple < max(2.0, strategy_min_daily_range_multiple() * 0.65):
        return None
    if extension_from_low is None or extension_from_low > max(strategy_max_daily_extension_from_low_pct(), 220.0):
        return None
    if drawdown_from_high < strategy_min_launch_drawdown_from_high_pct():
        return None
    if range_position > strategy_max_launch_range_position_pct():
        return None
    if recent_extension_from_low is None or recent_extension_from_low > strategy_max_launch_recent_low_extension_pct():
        return None
    if p1d is None or p1d < -18 or p1d > strategy_max_launch_24h_price_pct():
        return None
    if p3d is not None and p3d > strategy_max_launch_3d_price_pct():
        return None
    if p7d is not None and p7d > strategy_max_launch_7d_price_pct():
        return None
    if range_14d_pct is not None and range_14d_pct > strategy_max_launch_14d_range_pct():
        return None
    if base_days < strategy_min_launch_base_days():
        return None
    if daily_volume_ratio is None or daily_volume_ratio < strategy_min_launch_daily_volume_ratio():
        return None
    if daily_volume_ratio > strategy_max_launch_daily_volume_ratio():
        return None

    oi_history = get_oi_history(symbol, period="1d", limit=30)
    if len(oi_history) < 3:
        return None
    oi_values = [
        to_float(row.get("sumOpenInterestValue")) or to_float(row.get("sumOpenInterest"))
        for row in oi_history
    ]
    if oi_values[-1] is None:
        return None
    oi_24h_pct = pct_change(oi_values[-1], oi_values[-2] if len(oi_values) >= 2 else None)
    oi_2d_pct = pct_change(oi_values[-1], oi_values[-3] if len(oi_values) >= 3 else None)
    min_oi_24h = strategy_min_daily_oi_24h_pct() * 0.6
    min_oi_2d = strategy_min_daily_oi_2d_pct() * 0.8
    if (oi_24h_pct is None or oi_24h_pct < min_oi_24h) and (oi_2d_pct is None or oi_2d_pct < min_oi_2d):
        return None

    four_h = closed_klines(get_klines(symbol, interval="4h", limit=180))
    if len(four_h) < 60:
        return None
    h4_closes = [kline_float(row, 4) for row in four_h]
    h4_highs = [kline_float(row, 2) for row in four_h]
    h4_lows = [kline_float(row, 3) for row in four_h]
    h4_volumes = [kline_float(row, 7) for row in four_h]
    if h4_closes[-1] is None:
        return None
    h4_last = h4_closes[-1]
    h4_low_90 = min(value for value in h4_lows[-90:] if value is not None)
    h4_high_90 = max(value for value in h4_highs[-90:] if value is not None)
    h4_low_30 = min(value for value in h4_lows[-30:] if value is not None)
    h4_high_30 = max(value for value in h4_highs[-30:] if value is not None)
    h4_low_prev = min(value for value in h4_lows[-90:-30] if value is not None)
    h4_range_90 = pct_change(h4_high_90, h4_low_90)
    h4_range_30 = pct_change(h4_high_30, h4_low_30)
    h4_position = (h4_last - h4_low_90) / (h4_high_90 - h4_low_90) * 100.0 if h4_high_90 > h4_low_90 else None
    h4_p24 = pct_change(h4_last, h4_closes[-7] if len(h4_closes) >= 7 else None)
    h4_p3d = pct_change(h4_last, h4_closes[-19] if len(h4_closes) >= 19 else None)
    h4_volume = avg_clean(h4_volumes[-6:])
    h4_volume_base = list_median([value for value in h4_volumes[-78:-6] if value is not None])
    h4_volume_ratio = h4_volume / h4_volume_base if h4_volume and h4_volume_base and h4_volume_base > 0 else None
    h4_higher_low = h4_low_30 >= h4_low_prev * 0.92 if h4_low_prev else False
    h4_reclaim = h4_last >= (list_median([value for value in h4_closes[-30:] if value is not None]) or h4_last)
    h4_compressed = h4_range_30 is not None and h4_range_90 is not None and h4_range_30 <= h4_range_90 * 0.75

    if h4_position is None or h4_position > strategy_max_launch_4h_range_position_pct():
        return None
    if h4_p24 is not None and (h4_p24 < -12 or h4_p24 > strategy_max_launch_4h_24h_price_pct()):
        return None
    if h4_p3d is not None and (h4_p3d < -35 or h4_p3d > strategy_max_launch_4h_3d_price_pct()):
        return None
    ignition = (
        (p1d is not None and p1d >= 4)
        or (h4_p24 is not None and h4_p24 >= 4)
        or (daily_volume_ratio is not None and daily_volume_ratio >= 1.5)
        or (h4_volume_ratio is not None and h4_volume_ratio >= 1.3)
    )
    if not ignition:
        return None
    if not (h4_higher_low or h4_reclaim or h4_compressed or (h4_p24 is not None and h4_p24 >= 8)):
        return None

    score = 0
    score += 14 if range_multiple >= strategy_min_daily_range_multiple() else 9
    score += 12 if drawdown_from_high >= strategy_min_launch_drawdown_from_high_pct() + 20 else 6
    score += 14 if range_position <= 50 else 7
    score += 12 if recent_extension_from_low <= 80 else 5
    score += 12 if base_days >= strategy_min_launch_base_days() * 2 else 6
    if p1d is not None and 4 <= p1d <= 30:
        score += 10
    elif p1d is not None and 0 <= p1d <= strategy_max_launch_24h_price_pct():
        score += 6
    if p3d is not None and 8 <= p3d <= 80:
        score += 10
    elif p3d is not None and 0 <= p3d <= strategy_max_launch_3d_price_pct():
        score += 4
    if daily_volume_ratio is not None and 1.5 <= daily_volume_ratio <= 8:
        score += 10
    else:
        score += 5
    if oi_24h_pct is not None and oi_24h_pct >= strategy_min_daily_oi_24h_pct() * 2:
        score += 12
    elif oi_2d_pct is not None and oi_2d_pct >= strategy_min_daily_oi_2d_pct() * 2:
        score += 10
    else:
        score += 5
    score += 12 if h4_position <= 70 else 5
    if h4_higher_low:
        score += 8
    if h4_reclaim:
        score += 8
    if h4_compressed:
        score += 6
    if h4_p24 is not None and 4 <= h4_p24 <= 35:
        score += 10
    elif h4_p24 is not None and 0 <= h4_p24 <= strategy_max_launch_4h_24h_price_pct():
        score += 5
    if h4_volume_ratio is not None and h4_volume_ratio >= 1.3:
        score += 8
    if abs(funding) <= 0.02:
        score += 4

    if score < launch_min_score():
        return None

    flags = ["起漲確認"]
    if h4_higher_low:
        flags.append("4H低點抬高")
    if h4_reclaim:
        flags.append("4H回收均衡")
    if h4_compressed:
        flags.append("4H壓縮")
    if daily_volume_ratio is not None and daily_volume_ratio >= 1.5:
        flags.append("日量放大")
    if oi_24h_pct is not None and oi_24h_pct >= strategy_min_daily_oi_24h_pct():
        flags.append("OI增倉")

    reason = (
        f"起漲確認 | 日線區間{fmt_pct(range_position)} | 高點回撤{fmt_pct(drawdown_from_high)} | "
        f"近低+{fmt_pct(recent_extension_from_low)} | P1d {fmt_pct(p1d)} | P3d {fmt_pct(p3d)} | "
        f"base {base_days}d | OI1d {fmt_pct(oi_24h_pct)} | OI2d {fmt_pct(oi_2d_pct)} | "
        f"Vol x{daily_volume_ratio:.2f} | 4H位階{fmt_pct(h4_position)} | 4H24 {fmt_pct(h4_p24)}"
    )
    if h4_volume_ratio is not None:
        reason += f" | 4H量x{h4_volume_ratio:.2f}"

    return {
        "symbol": symbol,
        "score": min(score, 100),
        "reason": reason,
        "flags": "、".join(flags),
        "setup_type": "起漲確認",
        "oi_to_mcap": oi_to_mcap,
        "funding": funding,
        "market_rank": market_rank,
        "marketcap": marketcap,
        "extension_from_low": extension_from_low,
        "recent_extension_from_low": recent_extension_from_low,
        "drawdown_from_high": drawdown_from_high,
        "p1d": p1d,
        "p3d": p3d,
        "p7d": p7d,
    }


def launch_strategy_setup(
    symbol: str,
    phase_result: Any,
    sample: dict[str, Any],
    watch: WatchSymbol | None,
) -> tuple[bool, str, int]:
    marketcap = watch.marketcap_usd if watch else None
    oi_value = to_float(sample.get("oi_value_usd"))
    oi_to_mcap = oi_value / marketcap * 100.0 if oi_value is not None and marketcap and marketcap > 0 else None
    funding = to_float(sample.get("funding_rate_pct"))
    setup = launch_structure_setup(
        symbol,
        marketcap,
        funding,
        oi_to_mcap,
        watch.market_rank if watch else None,
    )
    if setup is None:
        return False, "not a bottom-or-launch setup", 0
    score = min(100, int(setup["score"]) + max(0, int(getattr(phase_result, "confidence", 50)) - 50) // 2)
    return True, str(setup["reason"]), score


def precision_strategy_setup(
    symbol: str,
    phase_result: Any,
    sample: dict[str, Any],
    watch: WatchSymbol | None,
) -> tuple[bool, str, int]:
    passed, reason, score = daily_strategy_setup(symbol, phase_result, sample, watch)
    if passed:
        return passed, reason, score
    return launch_strategy_setup(symbol, phase_result, sample, watch)


def strategy_event_base(position: dict[str, Any], sample: dict[str, Any], action: str, pnl_pct: float | None) -> dict[str, Any]:
    return {
        "timestamp": time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime()),
        "ts": time.time(),
        "action": action,
        "id": position.get("id"),
        "symbol": position.get("symbol"),
        "entry_phase": position.get("entry_phase"),
        "entry_reason": position.get("entry_reason"),
        "entry_price": position.get("entry_price"),
        "mark_price": sample.get("mark_price"),
        "pnl_pct": pnl_pct,
        "take_profit_pct": position.get("take_profit_pct"),
        "stop_loss_pct": position.get("stop_loss_pct"),
    }


def format_strategy_position(position: dict[str, Any]) -> str:
    symbol = str(position.get("symbol", "")).upper()
    entry = to_float(position.get("entry_price"))
    last_price = to_float(position.get("last_price"))
    pnl = to_float(position.get("last_pnl_pct"))
    if is_trade_plan_position(position):
        side = "做空" if str(position.get("side")).upper() == "SHORT" else "做多"
        leverage = to_float(position.get("leverage")) or trade_plan_paper_leverage()
        remaining = to_float(position.get("remaining_fraction")) or 0.0
        return (
            f"{symbol}｜{side}｜進 {fmt_num(entry, 6)}｜現 {fmt_num(last_price, 6)}｜"
            f"浮動 {fmt_pct(pnl, 2)} / {fmt_pct((pnl or 0.0) * leverage, 2)}({fmt_num(leverage)}倍)｜"
            f"TP1 {fmt_num(to_float(position.get('take_profit_1')), 6)}｜"
            f"TP2 {fmt_num(to_float(position.get('take_profit_2')), 6)}｜"
            f"SL {fmt_num(to_float(position.get('stop_loss')), 6)}｜剩餘 {remaining * 100:.0f}%"
        )
    tp_pct = to_float(position.get("take_profit_pct")) or strategy_take_profit_pct()
    sl_pct = to_float(position.get("stop_loss_pct")) or strategy_stop_loss_pct()
    tp_price = entry * (1 + tp_pct / 100.0) if entry is not None else None
    sl_price = entry * (1 - sl_pct / 100.0) if entry is not None else None
    phase = position.get("entry_phase", "n/a")
    return (
        f"{symbol}｜{phase}｜進 {fmt_num(entry, 6)}｜現 {fmt_num(last_price, 6)}｜PnL {fmt_pct(pnl, 2)}"
        f"｜TP {fmt_num(tp_price, 6)}｜SL {fmt_num(sl_price, 6)}"
    )


def collect_strategy_alerts(
    history: dict[str, deque[dict[str, Any]]],
    last_signal_at: dict[str, float],
) -> list[str]:
    positions = load_strategy_positions()
    alerts: list[str] = []
    remaining: list[dict[str, Any]] = []
    changed = False
    open_symbols = {str(item.get("symbol", "")).upper() for item in positions}

    for position in positions:
        symbol = str(position.get("symbol", "")).upper()
        if is_trade_plan_position(position):
            try:
                klines = trade_plan_klines_since(position)
                plan_events, closed = evaluate_trade_plan_candles(position, klines)
            except Exception as exc:
                print(f"Trade-plan tracking error for {symbol}: {exc}", file=sys.stderr, flush=True)
                remaining.append(position)
                continue
            if klines:
                last_close = to_float(klines[-1][4])
                position["last_price"] = last_close
                position["last_pnl_pct"] = strategy_pnl_pct(position, last_close)
                position["last_checked_at"] = time.time()
                changed = True
            for event in plan_events:
                save_strategy_event(event)
                action = str(event.get("action") or "")
                label = {
                    "PLAN_TP1": "TP1 半倉停利",
                    "PLAN_TP2": "TP2 全部出清",
                    "PLAN_SL": "SL 停損",
                    "PLAN_BE": "成本保護出場",
                }.get(action, action)
                alerts.append(
                    f"策略{label}｜{symbol}\n"
                    f"成交：{fmt_num(to_float(event.get('fill_price')), 6)}｜"
                    f"策略報酬 {fmt_pct(to_float(event.get('pnl_pct')), 2)}｜"
                    f"{fmt_num(to_float(event.get('leverage')))}倍損益 "
                    f"{fmt_pct(to_float(event.get('leveraged_pnl_pct')), 2)}｜"
                    f"{fmt_num(to_float(event.get('pnl_usd')))}U\n"
                    f"{event.get('note') or ''}"
                )
            if plan_events:
                changed = True
            if closed:
                last_signal_at[symbol] = time.time()
                continue
            remaining.append(position)
            continue

        sample = latest_history_sample(history.get(symbol, deque()))
        if sample is None:
            remaining.append(position)
            continue

        mark_price = to_float(sample.get("mark_price"))
        pnl_pct = strategy_pnl_pct(position, mark_price)
        position["last_price"] = mark_price
        position["last_pnl_pct"] = pnl_pct
        position["last_checked_at"] = sample.get("ts")

        tp_pct = to_float(position.get("take_profit_pct")) or strategy_take_profit_pct()
        sl_pct = to_float(position.get("stop_loss_pct")) or strategy_stop_loss_pct()
        close_action = None
        if pnl_pct is not None and pnl_pct >= tp_pct:
            close_action = "TP"
        elif pnl_pct is not None and pnl_pct <= -sl_pct:
            close_action = "SL"

        if close_action:
            event = strategy_event_base(position, sample, close_action, pnl_pct)
            save_strategy_event(event)
            last_signal_at[symbol] = time.time()
            changed = True
            alerts.append(
                f"策略{close_action}｜{symbol}\n"
                f"進場：{fmt_num(to_float(position.get('entry_price')), 6)}｜出場：{fmt_num(mark_price, 6)}｜PnL {fmt_pct(pnl_pct, 2)}\n"
                f"來源：{position.get('entry_phase')}｜{position.get('entry_reason')}"
            )
            continue

        remaining.append(position)

    positions = remaining
    open_symbols = {str(item.get("symbol", "")).upper() for item in positions}
    if strategy_mode() == "trade_plan":
        positions, plan_open_events, plan_alerts = latest_trade_plan_entries(
            positions,
            load_strategy_events(),
        )
        for event in plan_open_events:
            save_strategy_event(event)
        alerts.extend(plan_alerts)
        if changed or plan_open_events:
            save_strategy_positions(positions)
        return alerts

    capacity = max(0, strategy_max_open_positions() - len(positions))
    if capacity <= 0:
        if changed:
            save_strategy_positions(positions)
        return alerts

    gap_seconds = strategy_min_global_entry_gap_seconds()
    if gap_seconds > 0:
        last_open_ts = 0.0
        for event in load_strategy_events():
            if event.get("action") == "OPEN":
                last_open_ts = max(last_open_ts, to_float(event.get("ts")) or 0.0)
        if last_open_ts and time.time() - last_open_ts < gap_seconds:
            if changed:
                save_strategy_positions(positions)
            return alerts

    try:
        watch_by_symbol = {watch.symbol: watch for watch in resolve_watch_symbols()}
    except Exception:
        watch_by_symbol = {}

    candidates = []
    for symbol, samples in history.items():
        symbol = str(symbol).upper()
        if symbol in open_symbols:
            continue
        if time.time() - last_signal_at.get(symbol, 0.0) < strategy_signal_cooldown_seconds():
            continue
        sample = latest_history_sample(samples)
        if sample is None:
            continue
        result = strategy_phase_result(symbol, history)
        if result is None or result.decision == "不要進":
            continue
        mark_price = to_float(sample.get("mark_price"))
        if mark_price is None:
            continue
        watch = watch_by_symbol.get(symbol)
        try:
            passed, setup_reason, setup_score = precision_strategy_setup(
                symbol,
                result,
                sample,
                watch,
            )
        except Exception as exc:
            print(f"Strategy precision filter error for {symbol}: {exc}", file=sys.stderr, flush=True)
            continue
        if not passed:
            continue
        try:
            onchain_signal = analyze_onchain(
                symbol,
                market_symbol=watch.market_symbol if watch else symbol.removesuffix("USDT"),
                provider_id=watch.provider_id if watch else None,
            )
        except Exception as exc:
            print(f"Strategy on-chain filter error for {symbol}: {exc}", file=sys.stderr, flush=True)
            continue
        if onchain_signal.identity_verified and onchain_signal.score < strategy_onchain_min_score():
            now = time.time()
            if now - RUNTIME_ONCHAIN_OBSERVE_AT.get(symbol, 0.0) >= onchain_observe_cooldown_seconds():
                reason_text = "\n".join(f"- {reason}" for reason in onchain_signal.reasons[:4])
                if not reason_text:
                    reason_text = "- 鏈上資料不足或中性偏弱。"
                alerts.append(
                    f"鏈上觀察｜{symbol}｜{onchain_signal.verdict} {onchain_signal.score:+d}\n"
                    "日線條件通過，但鏈上分數未達開倉門檻，暫不開單。\n"
                    f"{reason_text}"
                )
                RUNTIME_ONCHAIN_OBSERVE_AT[symbol] = now
            continue
        if onchain_signal.identity_verified:
            setup_score += max(0, min(onchain_signal.score, 10))
            setup_reason = f"{setup_reason} | 鏈上已驗證 {onchain_signal.verdict} {onchain_signal.score:+d}"
        else:
            setup_reason = f"{setup_reason} | 鏈上地址未驗證，不計分"
        if orderbook_enabled():
            try:
                seed_orderbook_symbol(symbol)
                orderbook_signal = analyze_orderbook_accumulation(
                    symbol,
                    orderbook_db_path(),
                    lookback_seconds=orderbook_lookback_seconds(),
                    min_snapshots=orderbook_min_snapshots(),
                )
            except Exception as exc:
                print(f"Strategy orderbook filter error for {symbol}: {exc}", file=sys.stderr, flush=True)
                orderbook_signal = None
            if orderbook_signal is None:
                continue
            liquidity_row = type(
                "StrategyLiquidityRow",
                (),
                {"quote_volume_24h_usd": sample.get("quote_volume_24h_usd")},
            )()
            liquidity = assess_liquidity(
                row=liquidity_row,
                wgl={},
                book=orderbook_signal,
                reference_notional_usd=liquidity_reference_notional_usd(),
                min_quote_volume_24h_usd=liquidity_min_quote_volume_24h_usd(),
                min_quote_volume_1h_usd=liquidity_min_quote_volume_1h_usd(),
                min_depth_multiple=liquidity_min_depth_multiple(),
                max_spread_pct=liquidity_max_spread_pct(),
                max_slippage_pct=liquidity_max_slippage_pct(),
            )
            if not liquidity["liquidity_ready"]:
                continue
            setup_reason = (
                f"{setup_reason} | 流動性{liquidity['liquidity_score']}分 "
                f"24H ${fmt_num(liquidity['quote_volume_24h_usd'])}"
            )
            if orderbook_signal.snapshot_count >= orderbook_min_snapshots():
                if orderbook_signal.verdict == "偏弱/派發" and orderbook_signal.score <= 20:
                    continue
                if orderbook_signal.score >= orderbook_min_score():
                    setup_score += min(12, max(4, orderbook_signal.score // 10))
                    setup_reason = f"{setup_reason} | 訂單簿{orderbook_signal.verdict} {orderbook_signal.score}分"
                else:
                    setup_reason = f"{setup_reason} | 訂單簿{orderbook_signal.verdict} {orderbook_signal.score}分"
        else:
            continue
        candidates.append((setup_score, symbol, result, sample, mark_price, setup_reason))

    for _, symbol, result, sample, mark_price, setup_reason in sorted(
        candidates, key=lambda item: (item[0], item[1]), reverse=True
    )[
        : min(capacity, strategy_max_new_positions_per_scan())
    ]:
        entry_phase = "日線起漲" if "起漲確認" in setup_reason else "日線埋伏"
        position = {
            "id": strategy_trade_id(symbol),
            "symbol": symbol,
            "entry_ts": time.time(),
            "entry_time": time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime()),
            "entry_price": mark_price,
            "entry_phase": entry_phase,
            "entry_reason": setup_reason,
            "take_profit_pct": strategy_take_profit_pct(),
            "stop_loss_pct": strategy_stop_loss_pct(),
            "last_price": mark_price,
            "last_pnl_pct": 0.0,
            "last_checked_at": sample.get("ts"),
        }
        positions.append(position)
        last_signal_at[symbol] = time.time()
        changed = True
        save_strategy_event(strategy_event_base(position, sample, "OPEN", 0.0))
        alerts.append(
            f"策略開單｜{symbol}｜{entry_phase}\n"
            f"進場：{fmt_num(mark_price, 6)}｜TP +{strategy_take_profit_pct():.2f}%｜SL -{strategy_stop_loss_pct():.2f}%\n"
            f"{setup_reason}"
        )

    if changed:
        save_strategy_positions(positions)
    return alerts


def build_strategy_report(history: dict[str, deque[dict[str, Any]]] | None = None) -> str:
    history = history or RUNTIME_SPIKE_HISTORY
    positions = load_strategy_positions()
    for position in positions:
        symbol = str(position.get("symbol", "")).upper()
        sample = latest_history_sample(history.get(symbol, deque()))
        if sample is None:
            continue
        mark_price = to_float(sample.get("mark_price"))
        position["last_price"] = mark_price
        position["last_pnl_pct"] = strategy_pnl_pct(position, mark_price)
        position["last_checked_at"] = sample.get("ts")
    if positions:
        save_strategy_positions(positions)

    events = load_strategy_events()
    plan_positions = [item for item in positions if is_trade_plan_position(item)]
    legacy_positions = [item for item in positions if not is_trade_plan_position(item)]
    plan_events = [item for item in events if item.get("strategy_model") == "trade_plan_v1"]
    closed = [item for item in plan_events if item.get("action") in TRADE_PLAN_FINAL_ACTIONS]
    wins = [item for item in closed if (to_float(item.get("pnl_pct")) or 0.0) > 0]
    losses = [item for item in closed if (to_float(item.get("pnl_pct")) or 0.0) <= 0]
    returns = [to_float(item.get("pnl_pct")) for item in closed]
    returns = [value for value in returns if value is not None]
    avg = sum(returns) / len(returns) if returns else None
    total_usd = sum(to_float(item.get("pnl_usd")) or 0.0 for item in closed)

    lines = [
        "策略回報｜方向型 TP1 / TP2 / SL",
        (
            f"主策略：只採用做多/做空完整計畫｜每筆 {fmt_num(trade_plan_paper_margin_usd())}U × "
            f"{fmt_num(trade_plan_paper_leverage())}倍｜TP1半倉後移動SL到成本"
        ),
        (
            f"新模型持有 {len(plan_positions)}｜已平倉 {len(closed)}｜勝 {len(wins)}｜敗 {len(losses)}｜"
            f"平均 {fmt_pct(avg, 2)}｜累計損益 ${fmt_num(total_usd)}"
        ),
    ]
    if plan_positions:
        lines.extend(["", "【新模型持有中】"])
        for index, position in enumerate(
            sorted(plan_positions, key=lambda item: to_float(item.get("last_pnl_pct")) or 0, reverse=True)[:20],
            1,
        ):
            lines.append(f"{index}. {format_strategy_position(position)}")
    recent = closed[-10:]
    if recent:
        lines.extend(["", "【新模型最近平倉】"])
        labels = {"PLAN_TP2": "TP2", "PLAN_SL": "SL", "PLAN_BE": "成本保護"}
        for index, event in enumerate(reversed(recent), 1):
            lines.append(
                f"{index}. {event.get('symbol')}｜{labels.get(str(event.get('action')), event.get('action'))}｜"
                f"策略 {fmt_pct(to_float(event.get('pnl_pct')), 2)}｜"
                f"槓桿後 {fmt_pct(to_float(event.get('leveraged_pnl_pct')), 2)}｜"
                f"{fmt_num(to_float(event.get('pnl_usd')))}U"
            )
    if not plan_positions and not recent:
        lines.append("目前尚無新模型交易；沒有完整 TP/SL 計畫就不開單。")
    lines.append("")
    lines.append(f"舊策略持有 {len(legacy_positions)} 檔（只追蹤出場，已停止新增）")
    for position in legacy_positions[:10]:
        lines.append(f"- {format_strategy_position(position)}")
    return "\n".join(lines)


def ranked_oi_candidates(
    history: dict[str, deque[dict[str, Any]]],
    *,
    limit: int,
) -> list[dict[str, Any]]:
    watch_symbols = resolve_watch_symbols()
    rows = [
        row
        for row in get_oi_snapshots(watch_symbols, max_workers=oi_snapshot_workers())
        if row.oi_value_usd is not None
    ]
    now = time.time()
    out: list[dict[str, Any]] = []
    for row in rows:
        remember_report_sample(history, row, now)
        metrics = report_metrics(row, history, now)
        metrics["basis_pct"] = pct_change(
            to_float(getattr(row, "mark_price", None)),
            to_float(getattr(row, "price", None)),
        )
        score = attention_score(row, metrics)
        out.append({"row": row, "metrics": metrics, "score": score})
    out.sort(
        key=lambda item: (
            item["score"],
            item["row"].oi_value_usd or 0,
        ),
        reverse=True,
    )
    return out if limit <= 0 else out[:limit]


def active_history_symbols(
    history: dict[str, deque[dict[str, Any]]],
    *,
    limit: int = 20,
) -> list[str]:
    active: list[tuple[float, str]] = []
    for symbol, samples in list(history.items()):
        rows = list(samples)
        if len(rows) < 2:
            continue
        latest = rows[-1]
        oldest = rows[0]
        latest_oi = to_float(latest.get("open_interest"))
        oldest_oi = to_float(oldest.get("open_interest"))
        latest_price = to_float(latest.get("mark_price"))
        oldest_price = to_float(oldest.get("mark_price"))
        oi_change = abs(pct_change(latest_oi, oldest_oi) or 0.0)
        price_change = abs(pct_change(latest_price, oldest_price) or 0.0)
        activity = oi_change * 2.0 + price_change
        if oi_change >= 1.0 or price_change >= 2.0:
            active.append((activity, str(symbol).upper()))
    active.sort(reverse=True)
    return [symbol for _, symbol in active[:limit]]


def select_deep_candidates(
    candidates: list[dict[str, Any]],
    active_symbols: set[str],
    *,
    limit: int,
    momentum_quota: int,
) -> list[dict[str, Any]]:
    if limit <= 0:
        return []
    structure_ranked = sorted(
        candidates,
        key=lambda item: (
            item.get("prefilter_score") or 0,
            (item.get("structure_screen") or {}).get("score") or 0,
            item.get("live_momentum_score") or 0,
        ),
        reverse=True,
    )
    momentum_ranked = sorted(
        (item for item in candidates if item["row"].symbol in active_symbols),
        key=lambda item: (
            item.get("momentum_rank_score") or 0,
            item.get("live_momentum_score") or 0,
            item.get("prefilter_score") or 0,
        ),
        reverse=True,
    )
    selected: list[dict[str, Any]] = []
    selected_symbols: set[str] = set()
    for item in momentum_ranked[: min(limit, max(0, momentum_quota))]:
        selected.append(item)
        selected_symbols.add(item["row"].symbol)
    for item in structure_ranked:
        if item["row"].symbol in selected_symbols:
            continue
        selected.append(item)
        selected_symbols.add(item["row"].symbol)
        if len(selected) >= limit:
            break
    selected.sort(
        key=lambda item: (
            max(
                item.get("prefilter_score") or 0,
                (item.get("momentum_rank_score") or 0)
                if item.get("selection_lane") == "momentum"
                else 0,
            ),
            item.get("prefilter_score") or 0,
            item.get("momentum_rank_score") or 0,
        ),
        reverse=True,
    )
    return selected[:limit]


def structure_first_oi_candidates(
    history: dict[str, deque[dict[str, Any]]],
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    if not RUNTIME_STRUCTURE_CACHE:
        RUNTIME_STRUCTURE_CACHE.update(load_structure_cache())
    watches = resolve_watch_symbols()
    shells = [{"row": watch, "metrics": {}, "score": 0} for watch in watches]
    saved_states = load_wgl_signal_states()
    active_symbols = active_history_symbols(history, limit=max(20, momentum_deep_quota() * 2))
    active_symbol_set = set(active_symbols)
    priority_symbols = structure_reference_symbols()
    priority_symbols.update(active_symbol_set)
    priority_symbols.update(
        symbol
        for symbol, state in saved_states.items()
        if str(state.get("state") or "") in {"資金預備", *TRIGGER_STATES}
    )
    structured_universe = scan_structure_universe(
        shells,
        get_klines,
        RUNTIME_STRUCTURE_CACHE,
        cache_seconds=structure_cache_seconds(),
        workers=structure_scan_workers(),
        lookback_days=strategy_daily_lookback_days(),
        min_candles=20,
        max_refresh=structure_refresh_batch_size(),
        priority_symbols=priority_symbols,
    )
    save_structure_cache()

    watch_by_symbol = {watch.symbol: watch for watch in watches}
    selected_watches: list[WatchSymbol] = []
    seen: set[str] = set()
    preselect_limit = max(onchain_report_candidate_count(), structure_scan_candidate_count())
    for item in structured_universe[:preselect_limit]:
        watch = item["row"]
        if watch.symbol not in seen:
            selected_watches.append(watch)
            seen.add(watch.symbol)
    for symbol in active_symbols:
        watch = watch_by_symbol.get(symbol)
        if watch and symbol not in seen:
            selected_watches.append(watch)
            seen.add(symbol)
    for symbol, state in saved_states.items():
        if str(state.get("state") or "") not in {"資金預備", *TRIGGER_STATES}:
            continue
        watch = watch_by_symbol.get(symbol)
        if watch and symbol not in seen:
            selected_watches.append(watch)
            seen.add(symbol)

    rows = [
        row
        for row in get_oi_snapshots(selected_watches, max_workers=oi_snapshot_workers())
        if row.oi_value_usd is not None
    ]
    screen_by_symbol = {
        item["row"].symbol: item.get("structure_screen") or {}
        for item in structured_universe
    }
    now = time.time()
    deep: list[dict[str, Any]] = []
    for row in rows:
        remember_report_sample(history, row, now)
        metrics = report_metrics(row, history, now)
        metrics["basis_pct"] = pct_change(
            to_float(getattr(row, "mark_price", None)),
            to_float(getattr(row, "price", None)),
        )
        momentum = live_momentum_score(row, metrics)
        structure = screen_by_symbol.get(row.symbol) or {}
        effective_structure = float(structure.get("score") or 0)
        if not structure.get("eligible"):
            effective_structure = min(effective_structure, 40.0)
        oi_1h = max(0.0, to_float(metrics.get("contracts_1h_pct")) or 0.0)
        price_1h = max(0.0, to_float(metrics.get("price_1h_pct")) or 0.0)
        momentum_rank_score = min(
            100.0,
            max(
                float(momentum),
                oi_1h * 4.0 + price_1h * 3.0 + (20.0 if oi_1h > 0 and price_1h > 0 else 0.0),
            ),
        )
        deep.append(
            {
                "row": row,
                "metrics": metrics,
                "score": attention_score(row, metrics),
                "structure_screen": structure,
                "live_momentum_score": momentum,
                "momentum_rank_score": round(momentum_rank_score, 2),
                "selection_lane": "momentum" if row.symbol in active_symbol_set else "structure",
                "prefilter_score": round(effective_structure * 0.85 + momentum * 0.15, 2),
            }
        )

    deep_limit = onchain_report_candidate_count()
    selected = select_deep_candidates(
        deep,
        active_symbol_set,
        limit=deep_limit,
        momentum_quota=momentum_deep_quota(),
    )
    return structured_universe, selected


def avg_clean(values: list[float | None]) -> float | None:
    clean = [value for value in values if value is not None]
    if not clean:
        return None
    return sum(clean) / len(clean)


def extract_binance_ban_until(body: str) -> float | None:
    marker = "banned until "
    if marker not in body:
        return None
    tail = body.split(marker, 1)[1]
    digits = []
    for char in tail:
        if char.isdigit():
            digits.append(char)
        else:
            break
    if not digits:
        return None
    try:
        return int("".join(digits)) / 1000.0
    except ValueError:
        return None


def raise_if_binance_backoff() -> None:
    if time.time() < RUNTIME_BINANCE_BACKOFF_UNTIL:
        until = time.strftime("%Y-%m-%d %H:%M:%S", time.localtime(RUNTIME_BINANCE_BACKOFF_UNTIL))
        raise ApiError(f"Binance rate limit cooldown until {until}")


def binance_market_json(path: str, params: dict[str, Any], *, timeout: int = 10) -> Any:
    global RUNTIME_BINANCE_BACKOFF_UNTIL
    raise_if_binance_backoff()
    query = urllib.parse.urlencode(params)
    url = f"{BINANCE_FAPI_BASE}{path}?{query}"
    request = urllib.request.Request(
        url,
        headers={
            "User-Agent": "oi-phase-wgl/1.0",
            "Accept": "application/json",
        },
    )
    try:
        with urllib.request.urlopen(request, timeout=timeout) as response:
            return json.loads(response.read().decode("utf-8"))
    except urllib.error.HTTPError as exc:
        body = exc.read().decode("utf-8", errors="replace")
        if exc.code in {418, 429}:
            until = extract_binance_ban_until(body)
            if until is None:
                until = time.time() + (600 if exc.code == 418 else 90)
            RUNTIME_BINANCE_BACKOFF_UNTIL = max(RUNTIME_BINANCE_BACKOFF_UNTIL, until)
        raise ApiError(f"Binance HTTP {exc.code}: {body}") from exc
    except urllib.error.URLError as exc:
        raise ApiError(f"Binance network error: {exc}") from exc


def get_spot_taker_flow(market_symbol: str) -> dict[str, Any] | None:
    spot_symbol = f"{str(market_symbol).upper()}USDT"
    now = time.time()
    cached = RUNTIME_SPOT_FLOW_CACHE.get(spot_symbol)
    if cached and now - float(cached.get("ts", 0.0)) < 300:
        return dict(cached.get("flow") or {}) or None
    query = urllib.parse.urlencode({"symbol": spot_symbol, "limit": 500})
    request = urllib.request.Request(
        f"https://api.binance.com/api/v3/aggTrades?{query}",
        headers={"User-Agent": "oi-phase-spot-flow/1.0", "Accept": "application/json"},
    )
    try:
        with urllib.request.urlopen(request, timeout=10) as response:
            trades = json.loads(response.read().decode("utf-8"))
    except Exception:
        RUNTIME_SPOT_FLOW_CACHE[spot_symbol] = {"ts": now, "flow": {}}
        return None
    buy_notional = 0.0
    sell_notional = 0.0
    for trade in trades if isinstance(trades, list) else []:
        price = to_float(trade.get("p"))
        quantity = to_float(trade.get("q"))
        if price is None or quantity is None:
            continue
        notional = price * quantity
        if bool(trade.get("m")):
            sell_notional += notional
        else:
            buy_notional += notional
    total = buy_notional + sell_notional
    if total <= 0:
        return None
    flow = {
        "symbol": spot_symbol,
        "taker_buy_notional": buy_notional,
        "taker_sell_notional": sell_notional,
        "taker_imbalance": (buy_notional - sell_notional) / total,
        "sample_trades": len(trades),
    }
    RUNTIME_SPOT_FLOW_CACHE[spot_symbol] = {"ts": now, "flow": flow}
    return flow


def get_funding_history_pct(symbol: str, *, limit: int = 8) -> list[float]:
    symbol = symbol.strip().upper()
    key = f"{symbol}:{limit}"
    now = time.time()
    cached = RUNTIME_FUNDING_CACHE.get(key)
    if cached and now - float(cached.get("ts", 0)) < 600:
        return list(cached.get("rates") or [])

    data = binance_market_json("/fapi/v1/fundingRate", {"symbol": symbol, "limit": limit})
    if not isinstance(data, list):
        raise ApiError(f"Unexpected funding response for {symbol}: {data}")
    rows = sorted(data, key=lambda row: int(row.get("fundingTime", 0)))
    rates = [
        rate * 100.0
        for rate in (to_float(row.get("fundingRate")) for row in rows)
        if rate is not None
    ]
    RUNTIME_FUNDING_CACHE[key] = {"ts": now, "rates": rates}
    return rates


def pct_from_tail(values: list[float | None], back: int) -> float | None:
    clean = [value for value in values if value is not None]
    if len(clean) <= back:
        return None
    return pct_change(clean[-1], clean[-1 - back])


def wgl_recent_context(symbol: str) -> dict[str, Any]:
    symbol = symbol.strip().upper()
    now = time.time()
    cached = RUNTIME_WGL_CONTEXT_CACHE.get(symbol)
    if cached and now - float(cached.get("ts", 0)) < 600:
        return dict(cached.get("context") or {})

    context: dict[str, Any] = {
        "price_1h_pct": None,
        "price_6h_pct": None,
        "price_24h_pct": None,
        "oi_1h_pct": None,
        "oi_6h_pct": None,
        "oi_24h_pct": None,
        "volume_ratio_3h": None,
        "quote_volume_1h_usd": None,
        "quote_volume_3h_usd": None,
        "range_6h_low": None,
        "range_6h_high": None,
        "range_6h_position_pct": None,
        "range_24h_low": None,
        "range_24h_high": None,
        "range_24h_position_pct": None,
        "drawdown_from_24h_high_pct": None,
        "funding_rates_pct": [],
        "errors": [],
    }
    context.update(calculate_4h_market_context([]))

    try:
        klines = get_klines(symbol, interval="1h", limit=30)
        closes = [kline_float(row, 4) for row in klines]
        highs = [kline_float(row, 2) for row in klines]
        lows = [kline_float(row, 3) for row in klines]
        quote_volumes = [kline_float(row, 7) for row in klines]
        context["price_1h_pct"] = pct_from_tail(closes, 1)
        context["price_6h_pct"] = pct_from_tail(closes, 6)
        context["price_24h_pct"] = pct_from_tail(closes, 24)
        last_close = closes[-1] if closes else None
        for label, lookback in (("6h", 6), ("24h", 24)):
            recent_highs = [value for value in highs[-lookback:] if value is not None]
            recent_lows = [value for value in lows[-lookback:] if value is not None]
            if last_close is not None and recent_highs and recent_lows:
                range_low = min(recent_lows)
                range_high = max(recent_highs)
                context[f"range_{label}_low"] = range_low
                context[f"range_{label}_high"] = range_high
                if range_high > range_low:
                    context[f"range_{label}_position_pct"] = (last_close - range_low) / (range_high - range_low) * 100.0
        range_24h_high = to_float(context.get("range_24h_high"))
        if last_close is not None and range_24h_high is not None and range_24h_high > 0:
            context["drawdown_from_24h_high_pct"] = (range_24h_high - last_close) / range_24h_high * 100.0
        recent_volume = avg_clean(quote_volumes[-3:])
        base_volume = list_median([value for value in quote_volumes[-24:-3] if value is not None])
        context["quote_volume_1h_usd"] = quote_volumes[-1] if quote_volumes else None
        recent_quote_volumes = [value for value in quote_volumes[-3:] if value is not None]
        context["quote_volume_3h_usd"] = sum(recent_quote_volumes) if recent_quote_volumes else None
        if recent_volume is not None and base_volume is not None and base_volume > 0:
            context["volume_ratio_3h"] = recent_volume / base_volume
    except Exception as exc:
        context["errors"].append(f"kline: {exc}")

    try:
        four_h_klines = closed_klines(get_klines(symbol, interval="4h", limit=120))
        context.update(calculate_4h_market_context(four_h_klines))
    except Exception as exc:
        context["errors"].append(f"4h: {exc}")

    try:
        oi_rows = get_oi_history(symbol, period="1h", limit=30)
        oi_values = [
            to_float(row.get("sumOpenInterestValue")) or to_float(row.get("sumOpenInterest"))
            for row in oi_rows
        ]
        context["oi_1h_pct"] = pct_from_tail(oi_values, 1)
        context["oi_6h_pct"] = pct_from_tail(oi_values, 6)
        context["oi_24h_pct"] = pct_from_tail(oi_values, 24)
    except Exception as exc:
        context["errors"].append(f"oi: {exc}")

    try:
        context["funding_rates_pct"] = get_funding_history_pct(symbol, limit=8)
    except Exception as exc:
        context["errors"].append(f"funding: {exc}")

    RUNTIME_WGL_CONTEXT_CACHE[symbol] = {"ts": now, "context": context}
    return dict(context)


def ravelab_structure_candidate(row: Any, metrics: dict[str, float | None]) -> dict[str, Any] | None:
    symbol = str(row.symbol).upper()
    if symbol in strategy_blocklist():
        return None
    funding = to_float(row.funding_rate_pct)
    if funding is None or abs(funding) > strategy_max_funding_pct():
        return None
    oi_to_mcap = metrics.get("oi_to_marketcap_pct")

    daily = closed_klines(get_klines(symbol, interval="1d", limit=strategy_daily_lookback_days()))
    if len(daily) < strategy_min_daily_candles():
        return None
    closes = [kline_float(item, 4) for item in daily]
    highs = [kline_float(item, 2) for item in daily]
    lows = [kline_float(item, 3) for item in daily]
    volumes = [kline_float(item, 7) for item in daily]
    if closes[-1] is None:
        return None
    valid_highs = [value for value in highs if value is not None]
    valid_lows = [value for value in lows if value is not None]
    if not valid_highs or not valid_lows:
        return None
    low = min(valid_lows)
    high = max(valid_highs)
    if low <= 0 or high <= low:
        return None
    last = closes[-1]
    range_size = high - low
    range_multiple = high / low
    range_position = (last - low) / range_size * 100.0
    extension_from_low = pct_change(last, low)
    recent_lows_60 = [value for value in lows[-60:] if value is not None]
    recent_low_60 = min(recent_lows_60) if recent_lows_60 else low
    recent_extension_from_low = pct_change(last, recent_low_60)
    drawdown_from_high = (high - last) / high * 100.0
    p1d = pct_change(last, closes[-2] if len(closes) >= 2 else None)
    p3d = pct_change(last, closes[-4] if len(closes) >= 4 else None)
    p7d = pct_change(last, closes[-8] if len(closes) >= 8 else None)
    base_days = sum(
        1
        for close in closes[-30:]
        if close is not None and ((close - low) / range_size * 100.0) <= strategy_max_base_band_from_low_pct()
    )
    daily_volume = avg_clean(volumes[-3:])
    daily_volume_base = list_median([value for value in volumes[-63:-3] if value is not None])
    daily_volume_ratio = daily_volume / daily_volume_base if daily_volume and daily_volume_base and daily_volume_base > 0 else None

    if range_multiple < strategy_min_daily_range_multiple():
        return None
    if drawdown_from_high < strategy_min_drawdown_from_high_pct():
        return None
    if range_position > strategy_max_bottom_range_position_pct():
        return None
    if recent_extension_from_low is None or recent_extension_from_low > strategy_max_recent_low_extension_pct():
        return None
    if p1d is None or p1d > strategy_max_24h_price_pct() or p1d < -18:
        return None
    if p3d is not None and p3d > strategy_max_3d_price_pct():
        return None
    if p7d is not None and p7d > strategy_max_7d_extension_pct():
        return None
    if base_days < strategy_min_base_days():
        return None
    if daily_volume_ratio is not None and daily_volume_ratio > strategy_max_daily_volume_spike_ratio():
        return None

    four_h = closed_klines(get_klines(symbol, interval="4h", limit=180))
    if len(four_h) < 60:
        return None
    h4_closes = [kline_float(item, 4) for item in four_h]
    h4_highs = [kline_float(item, 2) for item in four_h]
    h4_lows = [kline_float(item, 3) for item in four_h]
    h4_volumes = [kline_float(item, 7) for item in four_h]
    if h4_closes[-1] is None:
        return None
    h4_last = h4_closes[-1]
    h4_low_90 = min(value for value in h4_lows[-90:] if value is not None)
    h4_high_90 = max(value for value in h4_highs[-90:] if value is not None)
    h4_low_30 = min(value for value in h4_lows[-30:] if value is not None)
    h4_high_30 = max(value for value in h4_highs[-30:] if value is not None)
    h4_low_prev = min(value for value in h4_lows[-90:-30] if value is not None)
    h4_range_90 = pct_change(h4_high_90, h4_low_90)
    h4_range_30 = pct_change(h4_high_30, h4_low_30)
    h4_position = (h4_last - h4_low_90) / (h4_high_90 - h4_low_90) * 100.0 if h4_high_90 > h4_low_90 else None
    h4_p24 = pct_change(h4_last, h4_closes[-7] if len(h4_closes) >= 7 else None)
    h4_p3d = pct_change(h4_last, h4_closes[-19] if len(h4_closes) >= 19 else None)
    h4_volume = avg_clean(h4_volumes[-6:])
    h4_volume_base = list_median([value for value in h4_volumes[-78:-6] if value is not None])
    h4_volume_ratio = h4_volume / h4_volume_base if h4_volume and h4_volume_base and h4_volume_base > 0 else None
    higher_low = h4_low_30 >= h4_low_prev * 0.92 if h4_low_prev else False
    reclaim = h4_last >= (list_median([value for value in h4_closes[-30:] if value is not None]) or h4_last)
    compressed = h4_range_30 is not None and h4_range_90 is not None and h4_range_30 <= h4_range_90 * 0.75

    if h4_position is None or h4_position > strategy_max_4h_range_position_pct():
        return None
    if h4_p24 is not None and h4_p24 > strategy_max_4h_24h_price_pct():
        return None
    if h4_p24 is not None and h4_p24 < strategy_min_4h_24h_price_pct():
        return None
    if h4_p3d is not None and h4_p3d > strategy_max_4h_3d_price_pct():
        return None
    if h4_p3d is not None and h4_p3d < strategy_min_4h_3d_price_pct():
        return None
    if not (higher_low or reclaim or compressed):
        return None

    score = 0
    score += 18 if range_multiple >= 5 else 12
    score += 20 if range_position <= strategy_max_bottom_range_position_pct() * 0.75 else 12
    score += 12 if drawdown_from_high >= strategy_min_drawdown_from_high_pct() + 15 else 6
    if recent_extension_from_low is not None:
        score += 10 if recent_extension_from_low <= strategy_max_recent_low_extension_pct() * 0.6 else 4
    score += 14 if base_days >= strategy_min_base_days() else 8
    if daily_volume_ratio is not None and daily_volume_ratio >= 1.5:
        score += 10
    elif daily_volume_ratio is not None and daily_volume_ratio >= 1.0:
        score += 5
    if h4_position <= strategy_max_4h_range_position_pct() * 0.75:
        score += 10
    if higher_low:
        score += 10
    if reclaim:
        score += 8
    if compressed:
        score += 8
    if h4_p24 is not None and 0 <= h4_p24 <= strategy_max_4h_24h_price_pct():
        score += 8
    elif h4_p24 is not None and -6 <= h4_p24 < 0:
        score += 3
    if h4_volume_ratio is not None and h4_volume_ratio >= 1.3:
        score += 8
    if abs(funding) <= 0.02:
        score += 4

    if score < ravelab_min_score():
        return None
    reason = (
        f"日線區間{fmt_pct(range_position)}｜高點回撤{fmt_pct(drawdown_from_high)}｜近低+{fmt_pct(recent_extension_from_low)}｜歷史{range_multiple:.1f}x｜"
        f"base {base_days}d｜4H位階{fmt_pct(h4_position)}｜"
        f"4H24 {fmt_pct(h4_p24)}｜4H量x{h4_volume_ratio:.2f}" if h4_volume_ratio else
        f"日線區間{fmt_pct(range_position)}｜高點回撤{fmt_pct(drawdown_from_high)}｜近低+{fmt_pct(recent_extension_from_low)}｜歷史{range_multiple:.1f}x｜"
        f"base {base_days}d｜4H位階{fmt_pct(h4_position)}｜4H24 {fmt_pct(h4_p24)}"
    )
    flags = []
    if higher_low:
        flags.append("4H低點抬高")
    if reclaim:
        flags.append("4H回收均衡")
    if compressed:
        flags.append("4H壓縮")
    return {
        "symbol": symbol,
        "score": score,
        "reason": reason,
        "flags": "、".join(flags) if flags else "4H待確認",
        "oi_to_mcap": oi_to_mcap,
        "funding": funding,
        "market_rank": row.market_rank,
        "marketcap": row.marketcap_usd,
        "extension_from_low": extension_from_low,
        "recent_extension_from_low": recent_extension_from_low,
        "drawdown_from_high": drawdown_from_high,
        "p1d": p1d,
        "p3d": p3d,
        "p7d": p7d,
    }


def launch_structure_candidate(row: Any, metrics: dict[str, float | None]) -> dict[str, Any] | None:
    symbol = str(row.symbol).upper()
    setup = launch_structure_setup(
        symbol,
        row.marketcap_usd,
        to_float(row.funding_rate_pct),
        metrics.get("oi_to_marketcap_pct"),
        row.market_rank,
    )
    return setup


def wgl_stage_candidate(
    row: Any,
    metrics: dict[str, float | None],
    book: Any | None = None,
    bottom: dict[str, Any] | None = None,
    launch: dict[str, Any] | None = None,
) -> dict[str, Any]:
    symbol = str(row.symbol).upper()
    context = wgl_recent_context(symbol)
    funding_rates = list(context.get("funding_rates_pct") or [])
    last_funding = funding_rates[-1] if funding_rates else to_float(row.funding_rate_pct)
    current_funding = to_float(row.funding_rate_pct)
    if current_funding is None:
        current_funding = last_funding

    price_1h = metrics.get("price_1h_pct")
    if price_1h is None:
        price_1h = to_float(context.get("price_1h_pct"))
    oi_1h = metrics.get("contracts_1h_pct")
    if oi_1h is None:
        oi_1h = to_float(context.get("oi_1h_pct"))
    price_6h = to_float(context.get("price_6h_pct"))
    oi_6h = to_float(context.get("oi_6h_pct"))
    price_24h = to_float(context.get("price_24h_pct"))
    oi_24h = to_float(context.get("oi_24h_pct"))
    volume_ratio = to_float(context.get("volume_ratio_3h"))
    oi_to_mcap = metrics.get("oi_to_marketcap_pct")
    current_price = to_float(getattr(row, "mark_price", None)) or to_float(getattr(row, "price", None))
    range_6h_low = to_float(context.get("range_6h_low"))
    range_6h_high = to_float(context.get("range_6h_high"))
    range_24h_low = to_float(context.get("range_24h_low"))
    range_24h_high = to_float(context.get("range_24h_high"))
    range_6h_position = to_float(context.get("range_6h_position_pct"))
    range_24h_position = to_float(context.get("range_24h_position_pct"))
    drawdown_from_24h_high = to_float(context.get("drawdown_from_24h_high_pct"))
    if current_price is not None and range_6h_low is not None and range_6h_high is not None and range_6h_high > range_6h_low:
        range_6h_position = (current_price - range_6h_low) / (range_6h_high - range_6h_low) * 100.0
    if current_price is not None and range_24h_low is not None and range_24h_high is not None and range_24h_high > range_24h_low:
        range_24h_position = (current_price - range_24h_low) / (range_24h_high - range_24h_low) * 100.0
    if current_price is not None and range_24h_high is not None and range_24h_high > 0:
        drawdown_from_24h_high = (range_24h_high - current_price) / range_24h_high * 100.0

    recent_rates = funding_rates[-6:]
    negative_count = sum(1 for value in recent_rates if value < 0)
    deep_negative_count = sum(1 for value in recent_rates if value <= -0.05)
    funding_returning_to_zero = (
        deep_negative_count >= 1
        and last_funding is not None
        and -0.015 <= last_funding <= 0.02
    )
    funding_positive_hot = current_funding is not None and current_funding >= 0.08

    book_ready = bool(book and book.snapshot_count >= orderbook_min_snapshots())
    book_absorption = bool(book_ready and book.score >= orderbook_min_score())
    book_pre_absorption = bool(
        book_ready
        and not book_absorption
        and (
            (book.avg_imbalance_50 is not None and book.avg_imbalance_50 >= 0.25 and (book.positive_imbalance_ratio or 0) >= 0.5)
            or (
                (book.ask_depth_change_pct is not None and book.ask_depth_change_pct <= -10)
                and (
                    (book.bid_depth_change_pct is not None and book.bid_depth_change_pct >= 8)
                    or (book.bid_ask_ratio_change_pct is not None and book.bid_ask_ratio_change_pct >= 5)
                )
            )
        )
    )
    book_distribution = bool(
        book_ready
        and (
            book.verdict == "偏弱/派發"
            or (book.ask_depth_change_pct is not None and book.ask_depth_change_pct >= 25)
            or (book.avg_imbalance_50 is not None and book.avg_imbalance_50 <= -0.25)
        )
    )

    price_oi_sync_1h = (
        price_1h is not None and oi_1h is not None and price_1h >= 2.0 and oi_1h >= 2.0
    )
    price_oi_sync_6h = (
        price_6h is not None and oi_6h is not None and price_6h >= 4.0 and oi_6h >= 4.0
    )
    oi_fading_with_price_hold = oi_1h is not None and oi_1h <= -2.0 and price_1h is not None and price_1h >= 0
    short_side_winning = oi_1h is not None and oi_1h >= 3.0 and price_1h is not None and price_1h <= -1.5
    pullback_price_started = (
        price_24h is not None
        and wgl_pullback_min_24h_price_pct() <= price_24h <= wgl_pullback_max_24h_price_pct()
    )
    pullback_range_ok = (
        range_6h_position is not None
        and wgl_pullback_min_6h_range_position_pct() <= range_6h_position <= wgl_pullback_max_6h_range_position_pct()
    )
    pullback_funding_ok = current_funding is not None and abs(current_funding) <= wgl_pullback_max_funding_pct()
    pullback_oi_ok = oi_1h is None or oi_1h >= wgl_pullback_max_oi_1h_drop_pct()
    pullback_volume_ok = volume_ratio is None or volume_ratio >= wgl_pullback_min_volume_ratio()
    strong_pullback = bool(
        pullback_price_started
        and pullback_range_ok
        and pullback_funding_ok
        and pullback_oi_ok
        and pullback_volume_ok
    )
    price_stretched = (
        (price_24h is not None and price_24h >= 35.0 and not strong_pullback)
        or (price_6h is not None and price_6h >= 25.0 and not strong_pullback)
        or funding_positive_hot
    )

    score = 0.0
    reasons: list[str] = []
    risks: list[str] = []

    if current_funding is not None:
        if abs(current_funding) <= strategy_max_funding_pct():
            score += 8
            reasons.append(f"funding未熱 {fmt_pct(current_funding, 4)}")
        if current_funding < 0:
            score += 5
    if deep_negative_count >= 2:
        score += 16
        reasons.append(f"funding連續深負 {deep_negative_count}/6")
    elif negative_count >= 3:
        score += 8
        reasons.append(f"funding偏空 {negative_count}/6")
    if funding_returning_to_zero:
        risks.append("funding由深負回零，軋空燃料可能減少")
    if funding_positive_hot:
        score -= 20
        risks.append(f"funding偏熱 {fmt_pct(current_funding, 4)}")

    if price_oi_sync_1h:
        score += 18
        reasons.append(f"1H價/OI同步：價 {fmt_pct(price_1h)}、OI {fmt_pct(oi_1h)}")
    elif price_1h is not None and oi_1h is not None and price_1h > 0 and oi_1h > 0:
        score += 8
        reasons.append(f"1H價/OI同向：價 {fmt_pct(price_1h)}、OI {fmt_pct(oi_1h)}")
    if price_oi_sync_6h:
        score += 12
        reasons.append(f"6H價/OI同步：價 {fmt_pct(price_6h)}、OI {fmt_pct(oi_6h)}")
    if volume_ratio is not None and volume_ratio >= 1.5:
        score += 8
        reasons.append(f"3H量能 x{volume_ratio:.2f}")
    if strong_pullback:
        score += 30
        reasons.append(
            f"BLESS型回踩：24h {fmt_pct(price_24h)}、6H位階{fmt_pct(range_6h_position)}、"
            f"funding {fmt_pct(current_funding, 4)}"
        )
    elif pullback_price_started and pullback_range_ok:
        score += 8
        reasons.append(f"強勢回踩觀察：24h {fmt_pct(price_24h)}、6H位階{fmt_pct(range_6h_position)}")

    if book_absorption:
        score += 22
        reasons.append(f"訂單簿吸籌 {book.score}分")
    elif book_pre_absorption:
        score += 14
        reasons.append(
            "訂單簿預吸籌："
            f"Ask {fmt_pct(book.ask_depth_change_pct)}、Bid {fmt_pct(book.bid_depth_change_pct)}、比值 {fmt_pct(book.bid_ask_ratio_change_pct)}"
        )
    elif book and book.snapshot_count < orderbook_min_snapshots():
        risks.append(f"訂單簿快照不足 {book.snapshot_count}/{orderbook_min_snapshots()}")
    if book_distribution:
        score -= 22
        risks.append(f"訂單簿偏派發 {book.score}分")

    if bottom:
        score += min(16, float(bottom.get("score") or 0) * 0.16)
        reasons.append(f"底部結構 {bottom.get('score')}分")
    if launch:
        score += min(14, float(launch.get("score") or 0) * 0.14)
        reasons.append(f"起漲結構 {launch.get('score')}分")

    if oi_fading_with_price_hold:
        score -= 18
        risks.append(f"OI 1H轉負但價格撐住：OI {fmt_pct(oi_1h)}")
    if short_side_winning:
        score -= 16
        risks.append(f"OI增但價格跌，疑似空方佔優：價 {fmt_pct(price_1h)}、OI {fmt_pct(oi_1h)}")
    if price_stretched:
        risks.append("價格或funding已偏伸，不當底部埋伏")

    hard_exit = (
        book_distribution
        or (funding_positive_hot and ((price_6h or 0) > 5 or (price_24h or 0) > 10))
        or oi_fading_with_price_hold
        or short_side_winning
    )

    bottom_like = bool(bottom or book_absorption or book_pre_absorption)
    squeeze_like = bool((price_oi_sync_1h or price_oi_sync_6h) and (deep_negative_count >= 1 or negative_count >= 3))

    if hard_exit:
        stage = "尾端/派發"
        action = "不要進"
    elif strong_pullback:
        stage = "強勢回踩"
        action = "埋伏" if score >= wgl_pullback_min_score() else "再確認偏強"
    elif squeeze_like:
        stage = "軋空啟動"
        action = "再確認偏強" if score >= 55 else "再確認"
    elif bottom_like and not price_stretched:
        stage = "底部吸籌"
        action = "埋伏" if score >= 60 and (book_absorption or book_pre_absorption) else "再確認"
    elif launch and not price_stretched:
        stage = "起漲確認"
        action = "再確認偏強" if score >= 60 else "再確認"
    else:
        stage = "再確認"
        action = "再確認"

    if action == "不要進":
        score = min(score, 49)

    return {
        "symbol": symbol,
        "stage": stage,
        "action": action,
        "score": int(round(max(0, min(100, score)))),
        "reasons": reasons[:5],
        "risks": risks[:4],
        "price_1h_pct": price_1h,
        "oi_1h_pct": oi_1h,
        "price_6h_pct": price_6h,
        "oi_6h_pct": oi_6h,
        "price_24h_pct": price_24h,
        "oi_24h_pct": oi_24h,
        "volume_ratio": volume_ratio,
        "quote_volume_1h_usd": to_float(context.get("quote_volume_1h_usd")),
        "quote_volume_3h_usd": to_float(context.get("quote_volume_3h_usd")),
        "four_h_close": to_float(context.get("four_h_close")),
        "four_h_atr": to_float(context.get("four_h_atr")),
        "four_h_atr_pct": to_float(context.get("four_h_atr_pct")),
        "four_h_ema20": to_float(context.get("four_h_ema20")),
        "four_h_ema50": to_float(context.get("four_h_ema50")),
        "four_h_swing_high_20": to_float(context.get("four_h_swing_high_20")),
        "four_h_swing_low_20": to_float(context.get("four_h_swing_low_20")),
        "four_h_swing_high_60": to_float(context.get("four_h_swing_high_60")),
        "four_h_swing_low_60": to_float(context.get("four_h_swing_low_60")),
        "four_h_range_position_60_pct": to_float(context.get("four_h_range_position_60_pct")),
        "range_6h_position_pct": range_6h_position,
        "range_24h_position_pct": range_24h_position,
        "drawdown_from_24h_high_pct": drawdown_from_24h_high,
        "strong_pullback": strong_pullback,
        "funding_rates_pct": funding_rates,
        "context_errors": context.get("errors") or [],
    }


def legacy_composite_candidate_rows(
    ranked: list[dict[str, Any]],
    onchain_rows: list[dict[str, Any]],
    ravelab_rows: list[dict[str, Any]],
    launch_rows: list[dict[str, Any]],
    orderbook_rows: list[dict[str, Any]],
    wgl_rows: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    onchain_by_symbol = {item["row"].symbol: item["signal"] for item in onchain_rows}
    ravelab_by_symbol = {item["symbol"]: item for item in ravelab_rows}
    launch_by_symbol = {item["symbol"]: item for item in launch_rows}
    orderbook_by_symbol = {item["row"].symbol: item["signal"] for item in orderbook_rows}
    wgl_by_symbol = {item["symbol"]: item for item in wgl_rows}

    output: list[dict[str, Any]] = []
    for item in ranked:
        row = item["row"]
        symbol = row.symbol
        metrics = item["metrics"]
        base_score = int(item.get("score") or 0)
        onchain_signal = onchain_by_symbol.get(symbol)
        bottom = ravelab_by_symbol.get(symbol)
        launch = launch_by_symbol.get(symbol)
        book = orderbook_by_symbol.get(symbol)
        wgl = wgl_by_symbol.get(symbol)

        score = min(18, base_score * 0.18)
        labels: list[str] = []
        reasons: list[str] = []
        risks: list[str] = []

        if wgl:
            wgl_score = int(wgl.get("score") or 0)
            score += min(48, float(wgl_score) * 0.48)
            labels.append(str(wgl.get("stage") or "WGL"))
            if wgl.get("action") == "埋伏":
                score += 10
            elif wgl.get("action") == "再確認偏強":
                score += 7
            elif wgl.get("action") == "不要進":
                score -= 25
            wgl_reason = "；".join((wgl.get("reasons") or [])[:2]) or "WGL條件待補"
            reasons.append(f"WGL {wgl_score}分｜{wgl.get('action')}｜{wgl_reason}")
            if wgl.get("risks"):
                risks.extend([str(item) for item in wgl.get("risks", [])[:2]])

        if bottom:
            score += min(36, float(bottom["score"]) * 0.36)
            labels.append("底部")
            reasons.append(f"底部{bottom['score']}分：{bottom['flags']}")
        if launch:
            score += min(38, float(launch["score"]) * 0.38)
            labels.append("起漲")
            reasons.append(f"起漲{launch['score']}分：{launch['flags']}")

        if book:
            if book.snapshot_count < orderbook_min_snapshots():
                risks.append(f"訂單簿快照不足 {book.snapshot_count}/{orderbook_min_snapshots()}")
            elif book.score >= orderbook_min_score():
                score += min(30, float(book.score) * 0.30)
                labels.append("吸籌")
                first_reason = book.reasons[0] if book.reasons else f"{book.verdict}{book.score}分"
                reasons.append(f"訂單簿{book.score}分：{first_reason}")
            elif book.verdict == "偏弱/派發":
                score -= 18
                risks.append(f"訂單簿偏弱：{book.score}分")
            else:
                score += min(8, float(book.score) * 0.12)

        if onchain_signal:
            score += max(-18, min(18, float(onchain_signal.score) * 4.0))
            if onchain_signal.score >= strategy_onchain_min_score():
                reasons.append(f"鏈上{onchain_signal.verdict} {onchain_signal.score:+d}")
            elif onchain_signal.score < 0:
                risks.append(f"鏈上{onchain_signal.verdict} {onchain_signal.score:+d}")

        funding = to_float(row.funding_rate_pct)
        if funding is not None and abs(funding) <= strategy_max_funding_pct():
            score += 3
        elif funding is not None:
            score -= 10
            risks.append(f"Funding {fmt_pct(funding, 4)}")

        has_real_signal = bool(
            bottom
            or launch
            or (wgl and int(wgl.get("score") or 0) >= wgl_min_score())
            or (book and book.snapshot_count >= orderbook_min_snapshots() and book.score >= orderbook_min_score())
        )
        if not has_real_signal and score < 28:
            continue
        if not labels:
            labels.append("觀察")
        if risks:
            deduped_risks: list[str] = []
            for risk in risks:
                if risk not in deduped_risks:
                    deduped_risks.append(risk)
            reasons.append("風險：" + "；".join(deduped_risks[:2]))
        if not reasons:
            reasons.append("目前只有注意力分數，需再確認")

        deduped_reasons: list[str] = []
        for reason in reasons:
            if reason not in deduped_reasons:
                deduped_reasons.append(reason)

        output.append(
            {
                "symbol": symbol,
                "score": int(round(max(0, min(100, score)))),
                "labels": "、".join(labels),
                "reasons": deduped_reasons[:4],
                "row": row,
                "metrics": metrics,
                "wgl": wgl,
                "book": book,
                "wgl_score": int(wgl.get("score") or 0) if wgl else 0,
                "wgl_action": str(wgl.get("action") or "") if wgl else "",
            }
        )

    output.sort(
        key=lambda data: (
            0 if data.get("wgl_action") == "不要進" else 1,
            data["score"],
            data.get("wgl_score") or 0,
            data["row"].oi_value_usd or 0,
        ),
        reverse=True,
    )
    return output


def composite_candidate_rows(
    ranked: list[dict[str, Any]],
    onchain_rows: list[dict[str, Any]],
    ravelab_rows: list[dict[str, Any]],
    launch_rows: list[dict[str, Any]],
    orderbook_rows: list[dict[str, Any]],
    wgl_rows: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    onchain_by_symbol = {item["row"].symbol: item["signal"] for item in onchain_rows}
    ravelab_by_symbol = {item["symbol"]: item for item in ravelab_rows}
    launch_by_symbol = {item["symbol"]: item for item in launch_rows}
    orderbook_by_symbol = {item["row"].symbol: item["signal"] for item in orderbook_rows}
    wgl_by_symbol = {item["symbol"]: item for item in wgl_rows}

    output: list[dict[str, Any]] = []
    for source in ranked:
        row = source["row"]
        symbol = row.symbol
        metrics = source.get("metrics") or {}
        structure = source.get("structure_screen") or {}
        onchain_signal = onchain_by_symbol.get(symbol)
        bottom = ravelab_by_symbol.get(symbol)
        launch = launch_by_symbol.get(symbol)
        book = orderbook_by_symbol.get(symbol)
        wgl = wgl_by_symbol.get(symbol)

        components = v3_component_scores(
            row=row,
            metrics=metrics,
            structure=structure,
            wgl=wgl,
            book=book,
            bottom=bottom,
            launch=launch,
            onchain=onchain_signal,
            orderbook_min_snapshots=orderbook_min_snapshots(),
            liquidity_reference_notional_usd=liquidity_reference_notional_usd(),
            liquidity_min_quote_volume_24h_usd=liquidity_min_quote_volume_24h_usd(),
            liquidity_min_quote_volume_1h_usd=liquidity_min_quote_volume_1h_usd(),
            liquidity_min_depth_multiple=liquidity_min_depth_multiple(),
            liquidity_max_spread_pct=liquidity_max_spread_pct(),
            liquidity_max_slippage_pct=liquidity_max_slippage_pct(),
        )

        labels = [components["signal_state"]]
        reasons: list[str] = []
        risks: list[str] = []
        if components.get("short_squeeze"):
            labels.append("軋空")
            reasons.append(
                f"負 Funding 軋空：{fmt_pct(to_float(getattr(row, 'funding_rate_pct', None)), 4)}，價/OI 1H 同步"
            )
        if structure.get("reasons"):
            reasons.append("日線：" + "、".join(str(value) for value in structure["reasons"][:3]))
        if wgl:
            labels.append(str(wgl.get("stage") or "資金"))
            wgl_reason = "；".join(str(value) for value in (wgl.get("reasons") or [])[:2])
            if wgl_reason:
                reasons.append(f"資金：{wgl_reason}")
            risks.extend(str(value) for value in (wgl.get("risks") or [])[:2])
        if bottom:
            labels.append("底部")
            reasons.append(f"底部{bottom['score']}分：{bottom['flags']}")
        if launch:
            labels.append("起漲")
            reasons.append(f"起漲{launch['score']}分：{launch['flags']}")
        if book:
            if book.snapshot_count < orderbook_min_snapshots():
                risks.append(f"訂單簿快照不足 {book.snapshot_count}/{orderbook_min_snapshots()}")
            elif book.score >= orderbook_min_score():
                labels.append("吸籌")
                first_reason = book.reasons[0] if book.reasons else f"{book.verdict}{book.score}分"
                reasons.append(f"訂單簿{book.score}分：{first_reason}")
            elif book.verdict == "偏弱/派發":
                risks.append(f"訂單簿偏弱：{book.score}分")
        if onchain_signal:
            if getattr(onchain_signal, "identity_verified", False):
                reasons.append(f"鏈上已驗證：{onchain_signal.verdict} {onchain_signal.score:+d}")
            else:
                reasons.append(
                    f"DEX參考：{onchain_signal.verdict} {onchain_signal.score:+d}（地址未驗證，不計分）"
                )
        if components.get("spot_taker_imbalance") is not None:
            reasons.append(f"現貨主動成交 {fmt_pct(float(components['spot_taker_imbalance']) * 100)}")
        if components.get("basis_pct") is not None:
            reasons.append(f"合約基差 {fmt_pct(components.get('basis_pct'), 4)}")
        liquidity_status = str(components.get("liquidity_status") or "待資料")
        if liquidity_status == "不足":
            detail = "、".join(str(value) for value in components.get("liquidity_reasons") or [])
            risks.append(f"流動性不足：{detail or '未通過成交門檻'}")
        elif liquidity_status == "待資料":
            missing = "、".join(str(value) for value in components.get("liquidity_missing") or [])
            risks.append(f"流動性待確認：{missing or '資料不足'}")

        funding = to_float(getattr(row, "funding_rate_pct", None))
        if funding is not None and abs(funding) > strategy_max_funding_pct():
            risks.append(f"Funding {fmt_pct(funding, 4)}")
        if risks:
            deduped_risks: list[str] = []
            for risk in risks:
                if risk not in deduped_risks:
                    deduped_risks.append(risk)
            reasons.append("風險：" + "；".join(deduped_risks[:2]))
        if not reasons:
            reasons.append("結構已進入候選，但資金與觸發仍不足")

        deduped_reasons: list[str] = []
        for reason in reasons:
            if reason not in deduped_reasons:
                deduped_reasons.append(reason)

        trade_plan = build_trade_plan(
            row=row,
            metrics=metrics,
            structure=structure,
            wgl=wgl,
            book=book,
            components=components,
            min_confidence=trade_plan_min_confidence(),
            min_risk_reward=trade_plan_min_risk_reward(),
        )

        output.append(
            {
                "symbol": symbol,
                "score": components["overall_score"],
                **components,
                **trade_plan,
                "labels": "、".join(dict.fromkeys(labels)),
                "reasons": deduped_reasons[:4],
                "row": row,
                "metrics": metrics,
                "structure_screen": structure,
                "wgl": wgl,
                "book": book,
                "wgl_score": int(wgl.get("score") or 0) if wgl else 0,
                "wgl_action": str(wgl.get("action") or "") if wgl else "",
            }
        )

    output.sort(
        key=lambda data: (
            data.get("plan_priority") or 0,
            data.get("plan_confidence") or 0,
            data.get("state_priority") or 0,
            data["score"],
            data.get("structure_score") or 0,
            data.get("trigger_score") or 0,
            data.get("quality_score") or 0,
            getattr(data["row"], "oi_value_usd", None) or 0,
        ),
        reverse=True,
    )
    return output


def short_text(value: str, limit: int = 70) -> str:
    text = " ".join(str(value).split())
    if len(text) <= limit:
        return text
    return text[: max(0, limit - 1)].rstrip() + "…"


def first_signal_reason(item: dict[str, Any]) -> str:
    for reason in item.get("reasons") or []:
        text = str(reason)
        if text.startswith("風險："):
            continue
        if "｜" in text:
            parts = [part for part in text.split("｜") if part]
            text = parts[-1] if parts else text
        return short_text(text, 64)
    return "等下一根價格/OI確認"


def classify_wgl_trade_item(
    item: dict[str, Any],
    report_rank: int,
    seen: dict[str, dict[str, Any]],
) -> dict[str, Any]:
    symbol = str(item.get("symbol") or "").upper()
    first_seen = seen.get(symbol) or {}
    decision = str(item.get("trade_decision") or "不交易")
    reason = str(item.get("plan_reason") or "未通過交易計畫門檻")
    if decision in ACTIONABLE_DECISIONS:
        return {
            "trade_bucket": "open",
            "trade_decision": decision,
            "trade_setup": f"{decision}計畫",
            "trade_reason": reason,
            "first_seen": first_seen or None,
        }
    return {
        "trade_bucket": "no_trade",
        "trade_decision": "不交易",
        "trade_setup": "不交易",
        "trade_reason": reason,
        "first_seen": first_seen or None,
    }


def format_wgl_compact_line(index: int, item: dict[str, Any]) -> str:
    row = item["row"]
    rank = f"#{row.market_rank}" if row.market_rank else "#n/a"
    oi_to_mcap = item["metrics"].get("oi_to_marketcap_pct")
    return (
        f"{index}. {item['symbol']}｜{item['score']}分｜{item.get('trade_setup')}｜{rank}｜"
        f"MC ${fmt_num(row.marketcap_usd)}｜OI/MC {fmt_pct(oi_to_mcap)}"
    )


def compact_card_time(value: Any | None = None) -> str:
    text = str(value or "").strip()
    if not text:
        return time.strftime("%m/%d %H:%M")
    if "T" in text:
        date_part, time_part = text.split("T", 1)
        text = f"{date_part} {time_part[:5]}"
    if len(text) >= 16 and text[4] == "-" and text[7] == "-":
        return f"{text[5:7]}/{text[8:10]} {text[11:16]}"
    return text


def wgl_card_grade(score: int) -> str:
    if score >= 90:
        return "S級"
    if score >= 75:
        return "A級"
    if score >= 60:
        return "B級"
    return "C級"


def wgl_card_risk(item: dict[str, Any]) -> str:
    risk = int(item.get("risk_score") or 0)
    if risk >= 60:
        return "極高"
    if risk >= 35:
        return "高"
    if risk >= 15:
        return "中"
    return "低"


def wgl_card_mode(item: dict[str, Any]) -> str:
    decision = str(item.get("trade_decision") or "不交易")
    return decision if decision in ACTIONABLE_DECISIONS else "不交易"


def wgl_card_seen_state(item: dict[str, Any], seen: dict[str, dict[str, Any]]) -> dict[str, Any]:
    symbol = str(item.get("symbol") or "").upper()
    row = item.get("row")
    current_price = to_float(getattr(row, "mark_price", None)) or to_float(getattr(row, "price", None))
    first_seen = seen.get(symbol) or item.get("first_seen") or {}
    has_seen = bool(first_seen)
    stored_count = int(
        to_float(first_seen.get("total_push_count")) or to_float(first_seen.get("push_count")) or (1 if has_seen else 0)
    )
    push_count = stored_count + 1 if has_seen else 1
    first_time = first_seen.get("first_seen_local") or first_seen.get("first_seen_utc") or time.strftime("%Y-%m-%d %H:%M")
    first_price = to_float(first_seen.get("first_price")) or current_price
    direction = str(first_seen.get("first_direction") or wgl_card_mode(item))
    return {
        "push_count": push_count,
        "first_time": compact_card_time(first_time),
        "latest_time": compact_card_time(),
        "first_price": first_price,
        "current_price": current_price,
        "direction": direction,
    }


def wgl_card_signal_counts(item: dict[str, Any]) -> tuple[int, int]:
    metrics = item.get("metrics") or {}
    wgl = item.get("wgl") or {}
    book = item.get("book")
    price_1h = to_float(wgl.get("price_1h_pct")) or to_float(metrics.get("price_1h_pct"))
    oi_1h = to_float(wgl.get("oi_1h_pct")) or to_float(metrics.get("contracts_1h_pct"))
    price_180s = to_float(metrics.get("price_180s_pct"))
    oi_180s = to_float(metrics.get("contracts_180s_pct"))
    price_24h = to_float(wgl.get("price_24h_pct"))
    oi_24h = to_float(wgl.get("oi_24h_pct"))

    short_count = sum(
        1
        for ok in [
            price_180s is not None and price_180s > 0,
            oi_180s is not None and oi_180s > 0,
            price_1h is not None and price_1h > 0,
            oi_1h is not None and oi_1h > 0,
            bool(book and getattr(book, "score", 0) >= orderbook_min_score()),
        ]
        if ok
    )
    trend_count = sum(
        1
        for ok in [
            price_24h is not None and price_24h > 0,
            oi_24h is not None and oi_24h > 0,
            bool(wgl.get("strong_pullback")),
            any(key in str(item.get("labels") or "") for key in ["底部", "起漲", "強勢回踩"]),
        ]
        if ok
    )
    return short_count, trend_count


def wgl_card_reasons(item: dict[str, Any]) -> str:
    reasons = [short_text(str(value), 42) for value in (item.get("reasons") or []) if value]
    return " / ".join(reasons[:3]) or "等待更多資料"


def wgl_entry_condition(item: dict[str, Any]) -> str:
    if wgl_card_mode(item) not in ACTIONABLE_DECISIONS:
        return "不建立倉位"
    return f"{fmt_num(item.get('entry_low'), 6)} - {fmt_num(item.get('entry_high'), 6)}"


def format_plan_price(value: Any) -> str:
    number = to_float(value)
    return "-" if number is None else fmt_num(number, 6)


def format_wgl_liquidity(item: dict[str, Any]) -> str:
    status = str(item.get("liquidity_status") or "待資料")
    score = int(to_float(item.get("liquidity_score")) or 0)
    volume_24h = to_float(item.get("quote_volume_24h_usd"))
    bid_depth = to_float(item.get("liquidity_bid_depth_05pct"))
    ask_depth = to_float(item.get("liquidity_ask_depth_05pct"))
    reference = to_float(item.get("liquidity_reference_notional_usd"))
    buy_slippage = to_float(item.get("liquidity_buy_slippage_pct"))
    sell_slippage = to_float(item.get("liquidity_sell_slippage_pct"))

    def money(value: float | None) -> str:
        return "n/a" if value is None else f"${fmt_num(value)}"

    def percent(value: float | None) -> str:
        return "n/a" if value is None else f"{value:.2f}%"

    return (
        f"流動性：{status} {score}/100｜24H {money(volume_24h)}｜"
        f"0.5%深度 B/A {money(bid_depth)}/{money(ask_depth)}｜"
        f"{money(reference)}滑價 買/賣 {percent(buy_slippage)}/{percent(sell_slippage)}"
    )


def format_wgl_funding_card(index: int, item: dict[str, Any], seen: dict[str, dict[str, Any]]) -> str:
    row = item["row"]
    state = wgl_card_seen_state(item, seen)
    score = int(item.get("score") or 0)
    grade = wgl_card_grade(score)
    first_price = to_float(state.get("first_price"))
    current_price = to_float(state.get("current_price"))
    move_pct = pct_change(current_price, first_price)
    up_pct = max(move_pct or 0.0, 0.0)
    down_pct = max(-(move_pct or 0.0), 0.0)
    marketcap = to_float(getattr(row, "marketcap_usd", None))
    decision = wgl_card_mode(item)
    confidence = int(to_float(item.get("plan_confidence")) or 0)
    plan_reason = short_text(str(item.get("plan_reason") or item.get("trade_reason") or "未通過交易門檻"), 80)
    management = str(item.get("plan_management") or "沒有通過條件，不建立倉位")
    invalidation = str(item.get("plan_invalidation") or plan_reason)
    actionable = decision in ACTIONABLE_DECISIONS
    tp1_pct = to_float(item.get("take_profit_1_pct"))
    tp2_pct = to_float(item.get("take_profit_2_pct"))
    stop_pct = to_float(item.get("stop_distance_pct"))
    rr1 = to_float(item.get("risk_reward_1"))
    rr2 = to_float(item.get("risk_reward_2"))

    return "\n".join(
        [
            "🟡 資金異動",
            "",
            f"幣種：{item['symbol']}",
            f"方向：{decision}",
            f"信心：{confidence}/100｜多分 {int(item.get('long_score') or 0)}｜空分 {int(item.get('short_score') or 0)}",
            f"階段：{item.get('signal_state', '-')}｜模型總分 {score}/100｜品質 {grade}｜風險 {wgl_card_risk(item)}",
            (
                f"結構：{item.get('structure_score', 0)}｜資金：{item.get('capital_score', 0)}｜"
                f"觸發：{item.get('trigger_score', 0)}｜資料：{item.get('quality_score', 0)}｜"
                f"風險分：{item.get('risk_score', 0)}"
            ),
            format_wgl_liquidity(item),
            f"進場區($)：{wgl_entry_condition(item)}",
            (
                f"TP1($)：{format_plan_price(item.get('take_profit_1'))}｜報酬 {tp1_pct:.2f}%｜RR {rr1:.2f}"
                if actionable and tp1_pct is not None and rr1 is not None
                else "TP1($)：-"
            ),
            (
                f"TP2($)：{format_plan_price(item.get('take_profit_2'))}｜報酬 {tp2_pct:.2f}%｜RR {rr2:.2f}"
                if actionable and tp2_pct is not None and rr2 is not None
                else "TP2($)：-"
            ),
            (
                f"SL($)：{format_plan_price(item.get('stop_loss'))}｜風險 {stop_pct:.2f}%"
                if actionable and stop_pct is not None
                else "SL($)：-"
            ),
            f"理由：{plan_reason}",
            f"倉位管理：{management}",
            f"失效條件：{invalidation}",
            f"出現次數：第 {state['push_count']} 次",
            f"#：{index}",
            f"首次推送：{state['first_time']}",
            f"首訊方向：{state['direction']}",
            f"最新推送：{state['latest_time']}",
            f"推送價格($)：{fmt_num(first_price, 6)}",
            f"當前幣價($)：{fmt_num(current_price, 6)}",
            f"推送後漲幅：{up_pct:.2f}%",
            f"推送後跌幅：{down_pct:.2f}%",
            "市值條件：無",
            f"市值參考：{'-' if marketcap is None else '$' + fmt_num(marketcap)}（不計分）",
        ]
    )


def wgl_daily_summary_path(day_key: str | None = None) -> Path:
    return WGL_DAILY_SUMMARIES_PATH / f"{compact_day_key(day_key)}.json"


def load_wgl_daily_summary_state() -> dict[str, Any]:
    if not WGL_DAILY_SUMMARY_STATE_PATH.exists():
        return {}
    try:
        data = json.loads(WGL_DAILY_SUMMARY_STATE_PATH.read_text(encoding="utf-8"))
    except Exception:
        return {}
    return data if isinstance(data, dict) else {}


def save_wgl_daily_summary_state(state: dict[str, Any]) -> None:
    WGL_DAILY_SUMMARY_STATE_PATH.parent.mkdir(parents=True, exist_ok=True)
    WGL_DAILY_SUMMARY_STATE_PATH.write_text(json.dumps(state, ensure_ascii=False, indent=2), encoding="utf-8")


def wgl_daily_summary_due(state: dict[str, Any]) -> bool:
    now = time.localtime()
    if now.tm_hour != wgl_daily_summary_hour() or now.tm_min < wgl_daily_summary_minute():
        return False
    return state.get("last_sent_date") != local_day_key()


def wgl_summary_time(value: Any) -> str:
    if value is None or str(value).strip() == "":
        return "-"
    text = compact_card_time(value)
    return text[-5:] if len(text) >= 5 else text


def build_wgl_daily_summary(day_key: str | None = None) -> str:
    day_key = day_key or local_day_key()
    stats = load_wgl_seen_symbols()
    items: list[dict[str, Any]] = []
    for symbol, record in stats.items():
        if not isinstance(record, dict):
            continue
        days = record.get("days") if isinstance(record.get("days"), dict) else {}
        day = days.get(day_key)
        if not isinstance(day, dict):
            continue
        first_price = to_float(day.get("first_price")) or to_float(record.get("first_price"))
        last_price = to_float(day.get("last_price")) or to_float(record.get("last_price"))
        max_price = to_float(day.get("max_price")) or last_price
        min_price = to_float(day.get("min_price")) or last_price
        items.append(
            {
                "symbol": symbol,
                "push_count": int(to_float(day.get("push_count")) or 0),
                "first_seen_local": day.get("first_seen_local"),
                "last_seen_local": day.get("last_seen_local"),
                "first_price": first_price,
                "last_price": last_price,
                "change_pct": pct_change(last_price, first_price),
                "snapshot_mfe_pct": pct_change(max_price, first_price),
                "snapshot_mae_pct": pct_change(min_price, first_price),
                "best_score": int(to_float(day.get("best_score")) or to_float(day.get("last_score")) or 0),
                "worst_score": int(to_float(day.get("worst_score")) or to_float(day.get("last_score")) or 0),
                "last_score": int(to_float(day.get("last_score")) or 0),
                "last_trade_decision": day.get("last_trade_decision") or "-",
                "last_trade_setup": day.get("last_trade_setup") or "-",
                "appearances": day.get("appearances") if isinstance(day.get("appearances"), list) else [],
            }
        )
    items.sort(
        key=lambda item: (
            item["best_score"],
            item["push_count"],
            item["change_pct"] if item["change_pct"] is not None else -9999,
            item["symbol"],
        ),
        reverse=True,
    )

    summary_path = wgl_daily_summary_path(day_key)
    summary_path.parent.mkdir(parents=True, exist_ok=True)
    outcome_payload: dict[str, Any] = {"schema": "wgl-signal-outcomes-v3", "signals": []}
    outcome_signals: list[dict[str, Any]] = []
    if WGL_SIGNAL_OUTCOMES_PATH.exists():
        try:
            outcome_payload = json.loads(WGL_SIGNAL_OUTCOMES_PATH.read_text(encoding="utf-8"))
            outcome_signals = [
                signal
                for signal in outcome_payload.get("signals", [])
                if isinstance(signal, dict) and signal.get("entry_date") == day_key
            ]
        except Exception:
            outcome_signals = []
    try:
        day_end_ts = time.mktime(time.strptime(f"{day_key} 23:59:59", "%Y-%m-%d %H:%M:%S"))
    except ValueError:
        day_end_ts = time.time()
    evaluated_until = min(time.time(), day_end_ts)
    for signal in outcome_signals:
        try:
            exact = evaluate_signal_path_1m(signal, evaluated_until)
        except Exception as exc:
            signal["path_error"] = str(exc)
            continue
        if exact:
            signal.update(exact)
    if outcome_signals:
        outcome_payload["updated_utc"] = time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime())
        WGL_SIGNAL_OUTCOMES_PATH.write_text(
            json.dumps(outcome_payload, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
    tp_count = sum(1 for signal in outcome_signals if signal.get("path_first_hit") == "TP")
    sl_count = sum(1 for signal in outcome_signals if signal.get("path_first_hit") == "SL")
    ambiguous_count = sum(1 for signal in outcome_signals if signal.get("path_first_hit") == "同分鐘不確定")
    open_count = sum(1 for signal in outcome_signals if not signal.get("path_first_hit"))

    payload = {
        "date": day_key,
        "generated_utc": time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime()),
        "generated_local": time.strftime("%Y-%m-%d %H:%M"),
        "symbol_count": len(items),
        "total_appearances": sum(item["push_count"] for item in items),
        "signal_outcomes": {
            "signals": len(outcome_signals),
            "tp_first": tp_count,
            "sl_first": sl_count,
            "ambiguous": ambiguous_count,
            "unresolved": open_count,
            "note": "TP/SL first hit、MFE、MAE 使用 Binance 1m K 線；同分鐘雙觸發不猜順序。",
        },
        "items": items,
    }
    summary_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")

    day_label = day_key.replace("-", "/")
    lines = [
        f"每日資金異動統計｜{day_label}",
        f"出現標的：{len(items)}｜總出現次數：{payload['total_appearances']}",
        f"觸發訊號：{len(outcome_signals)}｜TP先到 {tp_count}｜SL先到 {sl_count}｜"
        f"同分鐘不確定 {ambiguous_count}｜未結束 {open_count}",
        "規則：重複標的保留；TP/SL 與 MFE/MAE 使用 1m K 線，不猜同分鐘內先後。",
    ]
    if not items:
        lines.append("今日尚未記錄到 WGL TOP 標的。")
    else:
        for idx, item in enumerate(items, 1):
            lines.append(
                f"{idx}. {item['symbol']}｜出現 {item['push_count']} 次｜"
                f"{wgl_summary_time(item.get('first_seen_local'))}-{wgl_summary_time(item.get('last_seen_local'))}｜"
                f"最佳{item['best_score']}分｜最後{item['last_trade_decision']}/{item['last_trade_setup']}｜"
                f"漲跌 {fmt_pct(item['change_pct'])}｜MFE {fmt_pct(item['snapshot_mfe_pct'])}｜"
                f"MAE {fmt_pct(item['snapshot_mae_pct'])}"
            )
    lines.append(f"已存檔：{summary_path}")
    return "\n".join(lines)


def orderbook_collection_symbols(history: dict[str, deque[dict[str, Any]]] | None = None) -> list[str]:
    limit = orderbook_watch_candidates()
    history = history or RUNTIME_SPIKE_HISTORY
    ranked_history: list[tuple[float, str]] = []
    for symbol, samples in history.items():
        sample = latest_history_sample(samples)
        if sample is None:
            continue
        oi_value = to_float(sample.get("oi_value_usd")) or 0.0
        ranked_history.append((oi_value, str(symbol).upper()))
    ranked_history.sort(reverse=True)

    selected: list[str] = []
    seen: set[str] = set()
    if WGL_LATEST_REPORT_PATH.exists():
        try:
            latest_report = json.loads(WGL_LATEST_REPORT_PATH.read_text(encoding="utf-8"))
            report_symbols = latest_report.get("symbols") or []
            if isinstance(report_symbols, list):
                for raw_symbol in report_symbols:
                    symbol = normalize_symbol(str(raw_symbol or ""))
                    if not symbol or symbol in seen:
                        continue
                    selected.append(symbol)
                    seen.add(symbol)
                    if len(selected) >= limit:
                        return selected
        except (OSError, ValueError, TypeError) as exc:
            print(f"Orderbook latest report read error: {exc}", file=sys.stderr, flush=True)

    for _, symbol in ranked_history:
        if symbol and symbol not in seen:
            selected.append(symbol)
            seen.add(symbol)
        if len(selected) >= limit:
            return selected

    try:
        for watch in resolve_watch_symbols():
            symbol = watch.symbol.upper()
            if symbol in seen:
                continue
            selected.append(symbol)
            seen.add(symbol)
            if len(selected) >= limit:
                break
    except Exception as exc:
        print(f"Orderbook watch symbol fallback error: {exc}", file=sys.stderr, flush=True)
    return selected


def collect_orderbook_cycle(history: dict[str, deque[dict[str, Any]]] | None = None) -> tuple[int, int]:
    if not orderbook_enabled():
        return 0, 0
    symbols = orderbook_collection_symbols(history)
    if not symbols:
        return 0, 0
    with RUNTIME_ORDERBOOK_LOCK:
        snapshots, errors = collect_orderbook_snapshots(
            symbols,
            orderbook_db_path(),
            max_workers=orderbook_workers(),
            limit=100,
            reference_notional_usd=liquidity_reference_notional_usd(),
        )
        try:
            prune_orderbook_db(orderbook_db_path(), keep_days=orderbook_prune_days())
        except Exception as exc:
            print(f"Orderbook prune error: {exc}", file=sys.stderr, flush=True)
    if errors and len(errors) >= len(symbols):
        first_symbol, first_error = next(iter(errors.items()))
        print(f"Orderbook collect failed for all symbols, first {first_symbol}: {first_error}", file=sys.stderr, flush=True)
    return len(snapshots), len(errors)


def seed_orderbook_symbol(symbol: str) -> None:
    if not orderbook_enabled():
        return
    try:
        collect_orderbook_snapshots(
            [symbol],
            orderbook_db_path(),
            max_workers=1,
            limit=100,
            reference_notional_usd=liquidity_reference_notional_usd(),
        )
    except Exception as exc:
        print(f"Orderbook seed error for {symbol}: {exc}", file=sys.stderr, flush=True)


def legacy_build_onchain_hourly_report(history: dict[str, deque[dict[str, Any]]] | None = None) -> str:
    history = history or RUNTIME_SPIKE_HISTORY
    ranked = ranked_oi_candidates(history, limit=onchain_report_candidate_count())
    if not ranked:
        return "每小時鏈上雷達｜目前沒有可查 OI 標的。"
    if orderbook_enabled():
        try:
            with RUNTIME_ORDERBOOK_LOCK:
                collect_orderbook_snapshots(
                    [item["row"].symbol for item in ranked[:orderbook_report_seed_candidates()]],
                    orderbook_db_path(),
                    max_workers=orderbook_workers(),
                    limit=100,
                    reference_notional_usd=liquidity_reference_notional_usd(),
                )
        except Exception as exc:
            print(f"Hourly orderbook seed error: {exc}", file=sys.stderr, flush=True)

    onchain_rows = []
    ravelab_rows = []
    launch_rows = []
    orderbook_rows = []
    wgl_rows = []
    structure_limit = structure_scan_candidate_count()
    wgl_limit = wgl_scan_candidate_count()
    for index, item in enumerate(ranked):
        row = item["row"]
        watch_symbol = row.market_symbol or row.symbol.removesuffix("USDT")
        setup = None
        launch_setup = None
        book_signal = None
        try:
            signal = analyze_onchain(row.symbol, market_symbol=watch_symbol)
        except Exception as exc:
            print(f"Hourly on-chain scan error for {row.symbol}: {exc}", file=sys.stderr, flush=True)
            signal = None
        if signal is not None:
            onchain_rows.append({"row": row, "metrics": item["metrics"], "signal": signal})

        if index < structure_limit:
            try:
                setup = ravelab_structure_candidate(row, item["metrics"])
            except Exception as exc:
                print(f"RAVE/LAB structure scan error for {row.symbol}: {exc}", file=sys.stderr, flush=True)
                setup = None
            if setup is not None:
                ravelab_rows.append(setup)

            try:
                launch_setup = launch_structure_candidate(row, item["metrics"])
            except Exception as exc:
                print(f"Launch structure scan error for {row.symbol}: {exc}", file=sys.stderr, flush=True)
                launch_setup = None
            if launch_setup is not None:
                launch_rows.append(launch_setup)

        if orderbook_enabled():
            try:
                book_signal = analyze_orderbook_accumulation(
                    row.symbol,
                    orderbook_db_path(),
                    lookback_seconds=orderbook_lookback_seconds(),
                    min_snapshots=orderbook_min_snapshots(),
                )
            except Exception as exc:
                print(f"Orderbook accumulation scan error for {row.symbol}: {exc}", file=sys.stderr, flush=True)
                book_signal = None
            if book_signal is not None:
                orderbook_rows.append({"row": row, "signal": book_signal})

        if index < wgl_limit:
            try:
                wgl_setup = wgl_stage_candidate(row, item["metrics"], book_signal, setup, launch_setup)
            except Exception as exc:
                print(f"WGL stage scan error for {row.symbol}: {exc}", file=sys.stderr, flush=True)
                wgl_setup = None
            if wgl_setup is not None and int(wgl_setup.get("score") or 0) >= wgl_min_score():
                wgl_rows.append(wgl_setup)

    onchain_rows.sort(key=lambda item: item["signal"].score, reverse=True)
    ravelab_rows.sort(key=lambda item: item["score"], reverse=True)
    launch_rows.sort(key=lambda item: item["score"], reverse=True)
    orderbook_rows.sort(key=lambda item: item["signal"].score, reverse=True)
    wgl_rows.sort(key=lambda item: item["score"], reverse=True)
    orderbook_ready = [
        item for item in orderbook_rows if item["signal"].snapshot_count >= orderbook_min_snapshots()
    ]
    composite_rows = composite_candidate_rows(
        ranked,
        onchain_rows,
        ravelab_rows,
        launch_rows,
        orderbook_rows,
        wgl_rows,
    )

    bullish = [item for item in onchain_rows if item["signal"].score >= strategy_onchain_min_score()]
    warning = [
        item
        for item in onchain_rows
        if item["signal"].score < strategy_onchain_min_score() and item["signal"].reasons
    ]
    seen_symbols = load_wgl_seen_symbols()
    report_rows: list[dict[str, Any]] = []
    for idx, raw_item in enumerate(composite_rows[:composite_report_top_n()], 1):
        item = dict(raw_item)
        item["report_rank"] = idx
        item.update(classify_wgl_trade_item(item, idx, seen_symbols))
        report_rows.append(item)

    long_rows = [item for item in report_rows if item.get("trade_decision") == "做多"]
    short_rows = [item for item in report_rows if item.get("trade_decision") == "做空"]
    no_trade_rows = [item for item in report_rows if item.get("trade_decision") == "不交易"]

    lines: list[str] = []
    if report_rows:
        for idx, item in enumerate(report_rows, 1):
            if idx > 1:
                lines.append("")
                lines.append("-----")
                lines.append("")
            lines.append(format_wgl_funding_card(idx, item, seen_symbols))
    else:
        lines.append("🟡 資金異動")
        lines.append("")
        lines.append("目前沒有足夠明確的 TOP 5 候選。")

    strong_books = [item for item in orderbook_ready if item["signal"].score >= orderbook_min_score()]
    lines.append("")
    lines.append(
        "摘要："
        f"掃描 {len(ranked)}｜做多 {len(long_rows)}｜做空 {len(short_rows)}｜不交易 {len(no_trade_rows)}｜"
        f"WGL {len(wgl_rows)}｜"
        f"訂單簿 {len(strong_books)}｜鏈上偏多 {len(bullish)}"
    )
    if orderbook_enabled() and orderbook_rows and not orderbook_ready:
        best = orderbook_rows[0]["signal"]
        lines.append(f"訂單簿資料累積中：最佳 {best.symbol} 快照 {best.snapshot_count}/{orderbook_min_snapshots()}。")

    report_text = "\n".join(lines)
    save_wgl_report_event(report_rows, report_text)
    mark_wgl_report_seen(report_rows, seen_symbols)
    return report_text


def build_onchain_hourly_report(
    history: dict[str, deque[dict[str, Any]]] | None = None,
    *,
    persist: bool = True,
) -> str:
    history = history or RUNTIME_SPIKE_HISTORY
    structured, ranked = structure_first_oi_candidates(history)
    if not structured or not ranked:
        return "每小時資金雷達｜目前沒有可查 OI 標的。"
    selected_symbols = {item["row"].symbol for item in ranked}

    if orderbook_enabled():
        try:
            with RUNTIME_ORDERBOOK_LOCK:
                collect_orderbook_snapshots(
                    [item["row"].symbol for item in ranked[:orderbook_report_seed_candidates()]],
                    orderbook_db_path(),
                    max_workers=orderbook_workers(),
                    limit=100,
                    reference_notional_usd=liquidity_reference_notional_usd(),
                )
        except Exception as exc:
            print(f"Hourly orderbook seed error: {exc}", file=sys.stderr, flush=True)

    onchain_rows: list[dict[str, Any]] = []
    ravelab_rows: list[dict[str, Any]] = []
    launch_rows: list[dict[str, Any]] = []
    orderbook_rows: list[dict[str, Any]] = []
    wgl_rows: list[dict[str, Any]] = []
    structure_limit = min(len(ranked), structure_scan_candidate_count())
    wgl_limit = min(len(ranked), wgl_scan_candidate_count())
    onchain_limit = min(len(ranked), onchain_deep_candidate_count())
    spot_flow_limit = min(len(ranked), spot_flow_candidate_count())

    for index, item in enumerate(ranked):
        row = item["row"]
        watch_symbol = row.market_symbol or row.symbol.removesuffix("USDT")
        setup = None
        launch_setup = None
        book_signal = None

        if index < spot_flow_limit:
            spot_flow = get_spot_taker_flow(watch_symbol)
            if spot_flow:
                item["metrics"]["spot_taker_imbalance"] = spot_flow.get("taker_imbalance")
                item["metrics"]["spot_taker_notional"] = (
                    to_float(spot_flow.get("taker_buy_notional")) or 0
                ) + (to_float(spot_flow.get("taker_sell_notional")) or 0)

        if index < onchain_limit:
            try:
                signal = analyze_onchain(
                    row.symbol,
                    market_symbol=watch_symbol,
                    provider_id=getattr(row, "provider_id", None),
                )
            except Exception as exc:
                print(f"Hourly on-chain scan error for {row.symbol}: {exc}", file=sys.stderr, flush=True)
                signal = None
            if signal is not None:
                onchain_rows.append({"row": row, "metrics": item["metrics"], "signal": signal})

        if index < structure_limit:
            try:
                setup = ravelab_structure_candidate(row, item["metrics"])
            except Exception as exc:
                print(f"RAVE/LAB structure scan error for {row.symbol}: {exc}", file=sys.stderr, flush=True)
            if setup is not None:
                ravelab_rows.append(setup)
            try:
                launch_setup = launch_structure_candidate(row, item["metrics"])
            except Exception as exc:
                print(f"Launch structure scan error for {row.symbol}: {exc}", file=sys.stderr, flush=True)
            if launch_setup is not None:
                launch_rows.append(launch_setup)

        if orderbook_enabled():
            try:
                book_signal = analyze_orderbook_accumulation(
                    row.symbol,
                    orderbook_db_path(),
                    lookback_seconds=orderbook_lookback_seconds(),
                    min_snapshots=orderbook_min_snapshots(),
                )
            except Exception as exc:
                print(f"Orderbook accumulation scan error for {row.symbol}: {exc}", file=sys.stderr, flush=True)
            if book_signal is not None:
                orderbook_rows.append({"row": row, "signal": book_signal})

        if index < wgl_limit:
            try:
                wgl_setup = wgl_stage_candidate(row, item["metrics"], book_signal, setup, launch_setup)
            except Exception as exc:
                print(f"WGL stage scan error for {row.symbol}: {exc}", file=sys.stderr, flush=True)
                wgl_setup = None
            if wgl_setup is not None:
                wgl_rows.append(wgl_setup)

    composite_rows = composite_candidate_rows(
        ranked,
        onchain_rows,
        ravelab_rows,
        launch_rows,
        orderbook_rows,
        wgl_rows,
    )
    if persist:
        save_wgl_full_scan_event(structured, selected_symbols, ranked, composite_rows)
        update_wgl_signal_states(composite_rows)

    seen_symbols = load_wgl_seen_symbols()
    report_rows: list[dict[str, Any]] = []
    for idx, raw_item in enumerate(composite_rows[:composite_report_top_n()], 1):
        item = dict(raw_item)
        item["report_rank"] = idx
        item.update(classify_wgl_trade_item(item, idx, seen_symbols))
        report_rows.append(item)

    long_rows = [item for item in report_rows if item.get("trade_decision") == "做多"]
    short_rows = [item for item in report_rows if item.get("trade_decision") == "做空"]
    no_trade_rows = [item for item in report_rows if item.get("trade_decision") == "不交易"]
    eligible_count = sum(1 for item in structured if (item.get("structure_screen") or {}).get("eligible"))
    structure_covered = sum(
        1 for item in structured if int((item.get("structure_screen") or {}).get("data_points") or 0) > 0
    )

    lines: list[str] = []
    if report_rows:
        for idx, item in enumerate(report_rows, 1):
            if idx > 1:
                lines.extend(["", "-----", ""])
            lines.append(format_wgl_funding_card(idx, item, seen_symbols))
    else:
        lines.extend(["🟡 資金異動", "", "目前沒有足夠明確的 TOP 5 候選。"])

    lines.append("")
    lines.append(
        "摘要："
        f"全市場 {len(structured)}｜結構已掃 {structure_covered}｜待補 {len(structured) - structure_covered}｜"
        f"底部候選 {eligible_count}｜深度分析 {len(ranked)}｜"
        f"做多 {len(long_rows)}｜做空 {len(short_rows)}｜不交易 {len(no_trade_rows)}｜"
        f"鏈上已驗證 {sum(1 for item in report_rows if item.get('onchain_verified'))}"
    )
    lines.append("市值與市值排名不參與准入、評分或排序。")
    report_text = "\n".join(lines)

    if persist:
        save_wgl_report_event(report_rows, report_text)
        mark_wgl_report_seen(report_rows, seen_symbols)
        update_wgl_signal_outcomes(composite_rows, ranked)
        save_latest_wgl_report(report_text, report_rows, len(structured))
    return report_text


def build_research_thesis(raw_symbol: str, history: dict[str, deque[dict[str, Any]]] | None = None) -> str:
    history = history or RUNTIME_SPIKE_HISTORY
    symbol = normalize_symbol(raw_symbol)
    watch = market_context(symbol)
    if watch is None:
        watch = WatchSymbol(
            symbol=symbol,
            market_symbol=symbol.removesuffix("USDT"),
            market_rank=None,
            marketcap_usd=None,
            source="manual",
        )

    rows = get_oi_snapshots([watch], max_workers=1)
    if not rows or rows[0].oi_value_usd is None:
        return f"交易論證卡｜{symbol}\n結論：不要進\n原因：Binance Futures OI 查不到或資料為 n/a。"

    row = rows[0]
    now = time.time()
    remember_report_sample(history, row, now)
    metrics = report_metrics(row, history, now)
    score = attention_score(row, metrics)
    phase_decision = phase_report_decision(row, history)
    if phase_decision:
        decision, decision_reason = phase_decision
    else:
        decision, decision_reason = report_decision(row, metrics, score)

    try:
        onchain_signal = analyze_onchain(
            row.symbol,
            market_symbol=row.market_symbol or symbol.removesuffix("USDT"),
            provider_id=getattr(row, "provider_id", None),
        )
    except Exception as exc:
        onchain_signal = None
        onchain_error = str(exc)
    else:
        onchain_error = ""

    try:
        setup = ravelab_structure_candidate(row, metrics)
    except Exception as exc:
        setup = None
        setup_error = str(exc)
    else:
        setup_error = ""

    try:
        launch_setup = launch_structure_candidate(row, metrics)
    except Exception as exc:
        launch_setup = None
        launch_error = str(exc)
    else:
        launch_error = ""

    try:
        seed_orderbook_symbol(row.symbol)
        orderbook_signal = analyze_orderbook_accumulation(
            row.symbol,
            orderbook_db_path(),
            lookback_seconds=orderbook_lookback_seconds(),
            min_snapshots=orderbook_min_snapshots(),
        ) if orderbook_enabled() else None
    except Exception as exc:
        orderbook_signal = None
        orderbook_error = str(exc)
    else:
        orderbook_error = ""

    try:
        wgl_setup = wgl_stage_candidate(row, metrics, orderbook_signal, setup, launch_setup)
    except Exception as exc:
        wgl_setup = None
        wgl_error = str(exc)
    else:
        wgl_error = ""

    funding = to_float(row.funding_rate_pct)
    oi_to_mcap = metrics.get("oi_to_marketcap_pct")
    positives: list[str] = []
    risks: list[str] = []

    if wgl_setup:
        wgl_line = (
            f"WGL階段：{wgl_setup['stage']}｜{wgl_setup['action']}｜{wgl_setup['score']}分"
        )
        if wgl_setup["action"] == "不要進":
            risks.append(wgl_line)
        else:
            positives.append(wgl_line)
        if wgl_setup.get("reasons"):
            positives.append("；".join(wgl_setup["reasons"][:3]))
        if wgl_setup.get("risks"):
            risks.append("；".join(wgl_setup["risks"][:3]))
    else:
        risks.append(f"WGL階段檢查失敗：{wgl_error}")

    if setup:
        positives.append(f"RAVE/LAB 型態通過：{setup['score']}分，{setup['flags']}。")
        positives.append(str(setup["reason"]))
    else:
        risks.append("RAVE/LAB 日線 + 4H 型態未完整通過。")
        if setup_error:
            risks.append(f"型態檢查錯誤：{setup_error}")

    if launch_setup:
        positives.append(f"起漲確認通過：{launch_setup['score']}分，{launch_setup['flags']}。")
        positives.append(str(launch_setup["reason"]))
    else:
        risks.append("起漲確認未通過。")
        if launch_error:
            risks.append(f"起漲檢查錯誤：{launch_error}")

    if onchain_signal and onchain_signal.identity_verified:
        if onchain_signal.score >= strategy_onchain_min_score():
            positives.append(f"鏈上/DEX 分數達標：{onchain_signal.verdict} {onchain_signal.score:+d}。")
        else:
            risks.append(f"鏈上/DEX 未達標：{onchain_signal.verdict} {onchain_signal.score:+d}。")
        if onchain_signal.reasons:
            negative_markers = ("跌幅偏大", "已大漲", "賣壓", "成交幾乎", "活躍度偏低", "太薄", "轉入")
            positive_reasons = [
                reason for reason in onchain_signal.reasons if not any(marker in reason for marker in negative_markers)
            ]
            negative_reasons = [
                reason for reason in onchain_signal.reasons if any(marker in reason for marker in negative_markers)
            ]
            if positive_reasons and onchain_signal.score >= strategy_onchain_min_score():
                positives.append("；".join(positive_reasons[:2]))
            if negative_reasons:
                risks.append("；".join(negative_reasons[:2]))
            elif onchain_signal.score < strategy_onchain_min_score():
                risks.append("；".join(onchain_signal.reasons[:2]))
    elif onchain_signal:
        risks.append("鏈上合約地址尚未驗證，DEX 資料只作參考且不計分。")
    else:
        risks.append(f"鏈上檢查失敗：{onchain_error}")

    if orderbook_signal:
        if orderbook_signal.score >= orderbook_min_score():
            positives.append(
                f"訂單簿吸籌達標：{orderbook_signal.verdict} {orderbook_signal.score}分，快照 {orderbook_signal.snapshot_count}。"
            )
            if orderbook_signal.reasons:
                positives.append("；".join(orderbook_signal.reasons[:2]))
        elif orderbook_signal.snapshot_count < orderbook_min_snapshots():
            risks.append(
                f"訂單簿資料不足：{orderbook_signal.snapshot_count}/{orderbook_min_snapshots()}，需累積快照。"
            )
        elif orderbook_signal.verdict == "偏弱/派發":
            risks.append(
                f"訂單簿偏弱/派發：{orderbook_signal.score}分，"
                f"Ask {fmt_pct(orderbook_signal.ask_depth_change_pct)}，Bid {fmt_pct(orderbook_signal.bid_depth_change_pct)}。"
            )
        else:
            risks.append(f"訂單簿未達吸籌門檻：{orderbook_signal.verdict} {orderbook_signal.score}分。")
    elif orderbook_enabled():
        risks.append(f"訂單簿檢查失敗：{orderbook_error}")

    if funding is not None:
        if abs(funding) <= strategy_max_funding_pct():
            positives.append(f"Funding 未過熱：{fmt_pct(funding, 4)}。")
        else:
            risks.append(f"Funding 過熱：{fmt_pct(funding, 4)}。")
    if decision == "不要進":
        risks.append(f"OI 階段判定不要進：{decision_reason}")

    hard_block = (
        decision == "不要進"
        or (wgl_setup is not None and wgl_setup.get("action") == "不要進")
        or (funding is not None and abs(funding) > strategy_max_funding_pct())
        or (
            onchain_signal is not None
            and onchain_signal.identity_verified
            and onchain_signal.score < strategy_onchain_min_score()
        )
    )
    orderbook_ok = bool(orderbook_signal and orderbook_signal.score >= orderbook_min_score())
    if hard_block:
        conclusion = "不要進"
    elif wgl_setup and wgl_setup.get("action") == "埋伏":
        conclusion = "埋伏：WGL底部吸籌"
    elif wgl_setup and wgl_setup.get("action") == "再確認偏強":
        conclusion = f"再確認偏強：{wgl_setup.get('stage')}"
    elif (
        setup
        and onchain_signal
        and onchain_signal.identity_verified
        and onchain_signal.score >= strategy_onchain_min_score()
        and orderbook_ok
    ):
        conclusion = "底部吸籌共振候選"
    elif (
        launch_setup
        and onchain_signal
        and onchain_signal.identity_verified
        and onchain_signal.score >= strategy_onchain_min_score()
        and orderbook_ok
    ):
        conclusion = "起漲吸籌共振候選"
    elif setup and onchain_signal and onchain_signal.identity_verified and onchain_signal.score >= strategy_onchain_min_score():
        conclusion = "底部籌碼候選"
    elif launch_setup and onchain_signal and onchain_signal.identity_verified and onchain_signal.score >= strategy_onchain_min_score():
        conclusion = "起漲確認候選"
    else:
        conclusion = "再確認"

    rank = f"#{row.market_rank}" if row.market_rank else "#n/a"
    lines = [
        f"交易論證卡｜{symbol}",
        f"結論：{conclusion}",
        f"市值排名：{rank}｜市值 ${fmt_num(row.marketcap_usd)}｜OI ${fmt_num(row.oi_value_usd)}｜OI/市值 {fmt_pct(oi_to_mcap)}",
        f"價格 {fmt_num(row.mark_price, 6)}｜Funding {fmt_pct(funding, 4)}｜OI階段：{decision}｜{decision_reason}",
        f"WGL：{wgl_setup['stage']}｜{wgl_setup['action']}｜{wgl_setup['score']}分" if wgl_setup else "WGL：檢查失敗",
        "",
        "策略假設：",
        "- 先找日線長底部 + 4H 回收，目標是拿主升前的底部籌碼。",
        "- 若已脫離底部，改用起漲確認：量能、OI、4H 動能同步啟動，但 funding 與位階不能進入尾端。",
        "- 訂單簿吸籌看買賣盤長時間正失衡、賣盤變薄、買盤堆疊與深度異常；它是輔助共振，不單獨開倉。",
        "- 底部必須同時滿足：高點深回撤、目前仍貼近近低、日線低位盤整、4H 有止跌/回收跡象。",
        "- 鏈上/DEX 不偏空，OI 有足夠參與但 funding 不能過熱。",
    ]
    lines.append("")
    lines.append("多頭證據：")
    if positives:
        lines.extend(f"- {item}" for item in positives[:7])
    else:
        lines.append("- 目前沒有足夠多頭證據。")

    lines.append("")
    lines.append("反方證據/風險：")
    if risks:
        lines.extend(f"- {item}" for item in risks[:7])
    else:
        lines.append("- 暫無明確反方訊號，但仍需等價格與 OI 延續。")

    lines.append("")
    lines.append("觸發/失效：")
    lines.append("- 觸發：4H 維持低點抬高或回收均衡，鏈上分數維持 >= 門檻，OI 不轉弱。")
    lines.append("- 失效：鏈上轉偏空、Funding 過熱、4H 跌破底部回收結構，或 OI 1h 轉負。")
    return "\n".join(lines)


def build_realtime_trade_plan(row: Any, metrics: dict[str, Any]) -> dict[str, Any]:
    symbol = str(row.symbol).upper()
    seed_orderbook_symbol(symbol)
    book = analyze_orderbook_accumulation(
        symbol,
        orderbook_db_path(),
        lookback_seconds=orderbook_lookback_seconds(),
        min_snapshots=orderbook_min_snapshots(),
    )
    structure = (RUNTIME_STRUCTURE_CACHE.get(symbol) or {}).get("screen") or {}
    wgl = wgl_stage_candidate(row, metrics, book)
    candidates = composite_candidate_rows(
        [{"row": row, "metrics": metrics, "structure_screen": structure}],
        [],
        [],
        [],
        [{"row": row, "signal": book}],
        [wgl],
    )
    if candidates:
        return candidates[0]
    return {
        "trade_decision": "不交易",
        "plan_reason": "即時資料不足，無法建立交易計畫",
    }


def collect_spike_alerts(
    history: dict[str, deque[dict[str, Any]]],
    last_alert_at: dict[str, float],
) -> list[str]:
    global RUNTIME_SPIKE_CURSOR
    watch_symbols = resolve_watch_symbols()
    if not watch_symbols:
        return []
    batch_size = effective_spike_batch_size(len(watch_symbols))
    start = RUNTIME_SPIKE_CURSOR % len(watch_symbols)
    end = start + batch_size
    if end <= len(watch_symbols):
        scan_symbols_batch = watch_symbols[start:end]
    else:
        scan_symbols_batch = watch_symbols[start:] + watch_symbols[: end - len(watch_symbols)]
    RUNTIME_SPIKE_CURSOR = end % len(watch_symbols)

    now = time.time()
    window = spike_window_seconds()
    check_seconds = spike_check_seconds()
    max_history_age = window + max(120, check_seconds * 4)
    max_history_age = max(max_history_age, report_lookback_seconds() + max(300, check_seconds * 4))
    max_history_age = max(max_history_age, trend_window_seconds() + max(300, check_seconds * 4))
    max_history_age = max(max_history_age, strategy_base_window_seconds() + max(600, check_seconds * 4))
    min_pct = spike_min_change_pct()
    min_usd = spike_min_value_usd()
    min_contracts_pct = spike_min_contracts_pct()
    price_confirm_pct = spike_price_confirm_pct()
    cooldown = spike_cooldown_seconds()

    rows = [
        row
        for row in get_oi_snapshots(scan_symbols_batch, max_workers=oi_snapshot_workers())
        if row.oi_value_usd is not None
    ]
    alerts: list[str] = []

    for row in rows:
        sample = {
            "ts": now,
            "symbol": row.symbol,
            "oi_value_usd": row.oi_value_usd,
            "open_interest": row.open_interest,
            "mark_price": row.mark_price,
            "funding_rate_pct": row.funding_rate_pct,
            "quote_volume_24h_usd": row.quote_volume_24h_usd,
            "market_rank": row.market_rank,
            "marketcap_usd": row.marketcap_usd,
        }
        samples = history.setdefault(row.symbol, deque())
        samples.append(sample)
        while samples and now - float(samples[0]["ts"]) > max_history_age:
            samples.popleft()

        def find_baseline(window_seconds: int) -> dict[str, Any] | None:
            target_ts = now - window_seconds
            found = None
            for old in samples:
                if float(old["ts"]) <= target_ts:
                    found = old
                else:
                    break
            return found

        short_triggered = False
        baseline = find_baseline(window)
        if baseline is not None:
            old_value = to_float(baseline.get("oi_value_usd"))
            new_value = to_float(sample.get("oi_value_usd"))
            old_contracts = to_float(baseline.get("open_interest"))
            new_contracts = to_float(sample.get("open_interest"))
            old_price = to_float(baseline.get("mark_price"))
            new_price = to_float(sample.get("mark_price"))
            change_usd = new_value - old_value if new_value is not None and old_value is not None else None
            change_pct = pct_change(new_value, old_value)
            contracts_pct = pct_change(new_contracts, old_contracts)
            price_pct = pct_change(new_price, old_price)
            spike_key = f"spike:{row.symbol}"
            spike_check_key = f"spike-check:{row.symbol}"
            spike_ready = bool(
                change_pct is not None
                and change_pct >= min_pct
                and change_usd is not None
                and change_usd >= min_usd
                and contracts_pct is not None
                and contracts_pct >= min_contracts_pct
                and quick_liquidity_pass(row)
                and now - last_alert_at.get(spike_key, last_alert_at.get(row.symbol, 0.0)) >= cooldown
                and now - last_alert_at.get(spike_check_key, 0.0) >= window
            )
            if spike_ready:
                last_alert_at[spike_check_key] = now
                regime = classify_spike_regime(price_pct, contracts_pct, price_confirm_pct)
                grade = classify_spike_grade(
                    change_pct,
                    contracts_pct,
                    price_pct,
                    min_pct,
                    min_contracts_pct,
                    price_confirm_pct,
                )
                event = {
                    "event_type": "oi_spike_180s",
                    "timestamp_utc": row.timestamp_utc,
                    "symbol": row.symbol,
                    "regime": regime,
                    "grade": grade,
                    "window_seconds": window,
                    "old_oi_value_usd": old_value,
                    "new_oi_value_usd": new_value,
                    "change_usd": change_usd,
                    "change_pct": change_pct,
                    "contracts_change_pct": contracts_pct,
                    "price_change_pct": price_pct,
                    "old_open_interest": old_contracts,
                    "new_open_interest": new_contracts,
                    "mark_price": row.mark_price,
                    "funding_rate_pct": row.funding_rate_pct,
                    "quote_volume_24h_usd": row.quote_volume_24h_usd,
                    "market_rank": row.market_rank,
                    "marketcap_usd": row.marketcap_usd,
                    "watch_source": watch_source_description(),
                }
                plan_metrics = {
                    "contracts_180s_pct": contracts_pct,
                    "price_180s_pct": price_pct,
                }
                try:
                    trade_plan = build_realtime_trade_plan(row, plan_metrics)
                except Exception as exc:
                    print(f"Realtime trade plan error for {row.symbol}: {exc}", file=sys.stderr, flush=True)
                    trade_plan = {"trade_decision": "不交易", "plan_reason": "即時計畫分析失敗"}
                event["trade_plan"] = compact_trade_plan(trade_plan)
                save_spike_event(event)
                alert = format_trade_plan_alert(
                    row.symbol,
                    trade_plan,
                    source=f"{window}秒 OI爆量 [{grade}] {regime}",
                )
                if alert:
                    last_alert_at[spike_key] = now
                    alerts.append(alert)
                    short_triggered = True

        if short_triggered:
            continue

        trend_baseline = find_baseline(trend_window_seconds())
        if trend_baseline is None:
            continue
        trend_contracts_pct = pct_change(
            to_float(sample.get("open_interest")),
            to_float(trend_baseline.get("open_interest")),
        )
        trend_price_pct = pct_change(
            to_float(sample.get("mark_price")),
            to_float(trend_baseline.get("mark_price")),
        )
        structure = (RUNTIME_STRUCTURE_CACHE.get(row.symbol) or {}).get("screen") or {}
        trend_signal = classify_oi_trend_signal(
            contracts_1h_pct=trend_contracts_pct,
            price_1h_pct=trend_price_pct,
            oi_value_usd=to_float(sample.get("oi_value_usd")),
            funding_rate_pct=to_float(sample.get("funding_rate_pct")),
            structure=structure,
            quote_volume_24h_usd=to_float(sample.get("quote_volume_24h_usd")),
            min_quote_volume_24h_usd=liquidity_min_quote_volume_24h_usd(),
        )
        trend_key = f"trend:{row.symbol}"
        trend_check_key = f"trend-check:{row.symbol}"
        if (
            trend_signal is None
            or now - last_alert_at.get(trend_key, 0.0) < trend_cooldown_seconds()
            or now - last_alert_at.get(trend_check_key, 0.0) < window
        ):
            continue

        last_alert_at[trend_check_key] = now
        structure_score = int(structure.get("score") or 0)
        event = {
            "event_type": "oi_trend_1h",
            "timestamp_utc": row.timestamp_utc,
            "symbol": row.symbol,
            "signal_type": trend_signal["signal_type"],
            "lane": trend_signal["lane"],
            "window_seconds": trend_window_seconds(),
            "contracts_change_pct": trend_contracts_pct,
            "price_change_pct": trend_price_pct,
            "mark_price": row.mark_price,
            "oi_value_usd": row.oi_value_usd,
            "funding_rate_pct": row.funding_rate_pct,
            "quote_volume_24h_usd": row.quote_volume_24h_usd,
            "structure_score": structure_score,
            "structure_eligible": bool(structure.get("eligible")),
            "action": trend_signal["action"],
        }
        plan_metrics = {
            "contracts_1h_pct": trend_contracts_pct,
            "price_1h_pct": trend_price_pct,
        }
        try:
            trade_plan = build_realtime_trade_plan(row, plan_metrics)
        except Exception as exc:
            print(f"Realtime trade plan error for {row.symbol}: {exc}", file=sys.stderr, flush=True)
            trade_plan = {"trade_decision": "不交易", "plan_reason": "即時計畫分析失敗"}
        event["trade_plan"] = compact_trade_plan(trade_plan)
        save_spike_event(event)
        alert = format_trade_plan_alert(
            row.symbol,
            trade_plan,
            source=f"1H OI點火｜{trend_signal['signal_type']}",
        )
        if alert:
            last_alert_at[trend_key] = now
            alerts.append(alert)

    return alerts


def message_chunks(text: str, limit: int = 3800) -> list[str]:
    chunks: list[str] = []
    current = ""
    for line in text.splitlines():
        addition = line + "\n"
        if current and len(current) + len(addition) > limit:
            chunks.append(current.rstrip())
            current = addition
        else:
            current += addition
    if current:
        chunks.append(current.rstrip())
    return chunks


def send_long_message(token: str, chat_id: int, text: str) -> None:
    for chunk in message_chunks(text):
        telegram_call(token, "sendMessage", {"chat_id": chat_id, "text": chunk})
        time.sleep(0.25)


def send_card_message(token: str, chat_id: int, text: str) -> None:
    if not notification_cards_enabled():
        send_long_message(token, chat_id, text)
        return
    try:
        image_bytes = render_notification_card(text)
        telegram_send_photo(
            token,
            chat_id,
            image_bytes,
            caption=notification_caption(text),
        )
    except Exception as exc:
        print(f"Card send fallback for {chat_id}: {exc}", file=sys.stderr, flush=True)
        send_long_message(token, chat_id, text)


def split_command(text: str) -> tuple[str, list[str]]:
    parts = text.strip().split()
    if not parts:
        return "", []
    command = parts[0].split("@", 1)[0].lower()
    return command, parts[1:]


def parse_position_args(args: list[str], default_side: str = "LONG") -> tuple[str, str, float | None]:
    if not args:
        raise ValueError("用法：/entry OPN 0.244 或 /entry OPN short 0.244")
    symbol = normalize_symbol(args[0])
    side = default_side.upper()
    entry_price = None
    for raw in args[1:]:
        item = raw.strip()
        upper = item.upper()
        if upper in {"LONG", "L", "BUY", "多", "做多"}:
            side = "LONG"
            continue
        if upper in {"SHORT", "S", "SELL", "空", "做空"}:
            side = "SHORT"
            continue
        price = to_float(item)
        if price is not None and price > 0:
            entry_price = price
            continue
        raise ValueError(f"無法解析參數：{raw}")
    return symbol, side, entry_price


def register_position(chat_id: int, args: list[str], default_side: str = "LONG") -> str:
    symbol, side, entry_price = parse_position_args(args, default_side)
    sample = current_position_sample(symbol)
    mark_price = to_float(sample.get("mark_price"))
    if entry_price is None:
        if mark_price is None:
            raise ValueError(f"{symbol} 目前查不到 mark price，請手動輸入進場價。")
        entry_price = mark_price

    positions = [
        item
        for item in load_positions()
        if not (
            int(item.get("chat_id", 0)) == int(chat_id)
            and str(item.get("symbol", "")).upper() == symbol
            and str(item.get("side", "LONG")).upper() == side
        )
    ]
    position = {
        "chat_id": int(chat_id),
        "symbol": symbol,
        "side": side,
        "entry_price": entry_price,
        "entry_ts": time.time(),
        "entry_time": time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime()),
        "entry_mark_price": mark_price,
        "initial_open_interest": sample.get("open_interest"),
        "partial_taken": False,
        "break_even_stop": False,
        "last_price": mark_price,
        "last_pnl_pct": position_pnl_pct({"entry_price": entry_price, "side": side}, mark_price),
        "last_funding_rate_pct": sample.get("funding_rate_pct"),
        "last_open_interest": sample.get("open_interest"),
        "last_checked_at": sample.get("ts"),
    }
    positions.append(position)
    save_positions(positions)

    return (
        f"已登記盯盤｜{symbol}\n"
        f"方向：{side_label(side)}｜進場價：{fmt_num(entry_price, 6)}｜目前價：{fmt_num(mark_price, 6)}\n"
        f"監控間隔：{position_check_seconds()} 秒\n"
        f"出場：變成不要進 / funding > {position_funding_exit_pct():.4f}% 且已漲 / OI1h轉負 / -{position_stop_loss_pct():.2f}%停損\n"
        f"獲利 +{position_take_profit_pct():.2f}%：先停利一半，剩餘倉止損拉到開倉價"
    )


def remove_position(chat_id: int, args: list[str]) -> str:
    if not args:
        raise ValueError("用法：/exit OPN")
    symbol = normalize_symbol(args[0])
    side = None
    if len(args) >= 2:
        upper = args[1].upper()
        if upper in {"LONG", "L", "BUY", "多", "做多"}:
            side = "LONG"
        elif upper in {"SHORT", "S", "SELL", "空", "做空"}:
            side = "SHORT"
    positions = load_positions()
    remaining = []
    removed = []
    for item in positions:
        same_chat = int(item.get("chat_id", 0)) == int(chat_id)
        same_symbol = str(item.get("symbol", "")).upper() == symbol
        same_side = side is None or str(item.get("side", "LONG")).upper() == side
        if same_chat and same_symbol and same_side:
            removed.append(item)
        else:
            remaining.append(item)
    save_positions(remaining)
    if not removed:
        return f"{symbol} 沒有登記中的倉位。"
    for item in removed:
        save_position_event(
            {
                "timestamp": time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime()),
                "chat_id": chat_id,
                "symbol": item.get("symbol"),
                "side": item.get("side", "LONG"),
                "entry_price": item.get("entry_price"),
                "mark_price": item.get("last_price"),
                "pnl_pct": item.get("last_pnl_pct"),
                "action": "手動移除",
                "reason": "使用者手動平倉或取消監控",
                "partial_taken": bool(item.get("partial_taken")),
            }
        )
    return "已移除倉位監控：\n" + "\n".join(format_position(item) for item in removed)


def telegram_menu_commands() -> list[dict[str, str]]:
    return [
        {"command": "report", "description": "查看綜合前五與進場判斷"},
        {"command": "wgl", "description": "研究單一幣種完整訊號"},
        {"command": "entry", "description": "登記實盤倉位並啟動盯盤"},
        {"command": "positions", "description": "查看持倉、PnL與出場狀態"},
        {"command": "exit", "description": "停止指定倉位盯盤"},
        {"command": "orderbook", "description": "檢查委託簿與真實成交流"},
        {"command": "onchain", "description": "檢查鏈上與DEX多空證據"},
        {"command": "strategy_report", "description": "查看模擬單與TP/SL績效"},
        {"command": "alerts", "description": "開關每小時報告與即時警報"},
        {"command": "status", "description": "查看系統與資料更新狀態"},
        {"command": "help", "description": "查看核心指令與範例"},
    ]


def configure_alerts(chat_id: int, enabled: bool | None = None) -> str:
    subscribers = load_subscribers()
    if enabled is True:
        subscribers.add(int(chat_id))
        save_subscribers(subscribers)
    elif enabled is False:
        subscribers.discard(int(chat_id))
        save_subscribers(subscribers)

    active = int(chat_id) in subscribers
    state = "已開啟" if active else "已關閉"
    lines = [
        f"通知狀態：{state}",
        "內容：每小時綜合前五、狀態轉強、OI 爆量、持倉出場警報",
    ]
    lines.append("關閉方式：/alerts off" if active else "開啟方式：/alerts on")
    return "\n".join(lines)


def format_system_status(chat_id: int) -> str:
    now = time.time()
    payload: dict[str, Any] = {}
    if WGL_LATEST_REPORT_PATH.exists():
        try:
            loaded = json.loads(WGL_LATEST_REPORT_PATH.read_text(encoding="utf-8"))
            if isinstance(loaded, dict):
                payload = loaded
        except (OSError, ValueError, TypeError):
            payload = {}

    generated_at = to_float(payload.get("generated_at"))
    age_minutes = max(0, int((now - generated_at) / 60)) if generated_at else None
    generated_local = str(payload.get("generated_local") or "尚無報告")
    age_text = f"{age_minutes} 分鐘前" if age_minutes is not None else "尚無資料"
    next_due = initial_report_due_at(now)
    next_report_text = "等待掃描" if next_due <= now + 2 else time.strftime("%m/%d %H:%M", time.localtime(next_due))

    universe_value = to_float(payload.get("universe_size"))
    universe_size = int(universe_value) if universe_value is not None else 0
    raw_symbols = payload.get("symbols") or []
    top_symbols = [str(symbol).upper() for symbol in raw_symbols if symbol] if isinstance(raw_symbols, list) else []
    top_text = "、".join(top_symbols[:5]) or "尚無"

    batch_size = effective_spike_batch_size(universe_size)
    coverage_seconds = (
        math.ceil(universe_size / batch_size) * spike_check_seconds()
        if universe_size > 0 and batch_size > 0
        else None
    )
    coverage_text = f"約 {coverage_seconds} 秒覆蓋全市場" if coverage_seconds is not None else "等待市場清單"
    manual_positions = sum(
        1 for position in load_positions() if int(position.get("chat_id", 0)) == int(chat_id)
    )
    strategy_positions = len(load_strategy_positions())
    alerts_state = "開" if int(chat_id) in load_subscribers() else "關"
    uptime_minutes = max(0, int((now - BOT_STARTED_AT) / 60))

    return "\n".join(
        [
            "系統狀態｜運行中",
            f"通知：{alerts_state}｜本次運行：{uptime_minutes} 分鐘",
            f"最新報告：{generated_local}（{age_text}）",
            f"下次完整掃描：{next_report_text}",
            f"市場：{universe_size or '-'} 合約｜市值條件：無",
            f"目前前五：{top_text}",
            f"OI 雷達：每 {spike_check_seconds()} 秒輪巡｜{coverage_text}｜含 1H 資金點火",
            f"委託簿：{orderbook_watch_candidates()} 檔／{orderbook_collect_interval_seconds()} 秒",
            (
                f"流動性：{fmt_num(liquidity_reference_notional_usd())}U 倉位｜"
                f"24H成交額 >= ${fmt_num(liquidity_min_quote_volume_24h_usd())}｜"
                f"最大滑價 {liquidity_max_slippage_pct():.2f}%"
            ),
            (
                f"主策略：方向型 TP1/TP2/SL｜每 {strategy_scan_interval_seconds() // 60} 分鐘追蹤｜"
                f"手動盯盤 {manual_positions}｜模擬單 {strategy_positions}"
            ),
        ]
    )


def help_text() -> str:
    return (
        "核心操作\n"
        "/report - 綜合前五與進場判斷\n"
        "/wgl COAI - 單幣完整研究卡\n"
        "/orderbook COAI - 委託簿與真實成交流\n"
        "/onchain COAI - 鏈上與 DEX 證據\n\n"
        "實盤盯盤\n"
        "/entry OPN 0.244 - 登記多單\n"
        "/entry OPN short 0.244 - 登記空單\n"
        "/positions - 持倉與 PnL\n"
        "/exit OPN - 停止盯盤\n\n"
        "通知與績效\n"
        "/strategy_report - 模擬單與 TP/SL 績效\n"
        "/alerts on 或 /alerts off - 通知開關\n"
        "/status - 系統與資料更新狀態\n\n"
        "舊指令仍可使用，但不再佔用選單。"
    )


def handle_text(text: str, chat_id: int | None = None) -> str:
    command, args = split_command(text)
    if command in {"/start", "/help"}:
        return help_text()

    if command == "/status":
        if chat_id is None:
            return "此指令只能在 Telegram 使用。"
        return format_system_status(int(chat_id))

    if command == "/alerts":
        if chat_id is None:
            return "此指令只能在 Telegram 使用。"
        if not args or args[0].strip().lower() in {"status", "狀態"}:
            return configure_alerts(int(chat_id))
        action = args[0].strip().lower()
        if action in {"on", "1", "yes", "start", "開", "開啟"}:
            return configure_alerts(int(chat_id), True)
        if action in {"off", "0", "no", "stop", "關", "關閉"}:
            return configure_alerts(int(chat_id), False)
        return "用法：/alerts on、/alerts off 或 /alerts status"

    if command == "/spike_settings":
        return (
            "OI 爆量提醒設定\n"
            f"觀察窗口：{spike_window_seconds()} 秒\n"
            f"檢查間隔：{spike_check_seconds()} 秒\n"
            f"OI價值最小增幅：+{spike_min_change_pct():.2f}%\n"
            f"OI價值最小增加額：+${fmt_num(spike_min_value_usd())}\n"
            f"合約OI最小增幅：+{spike_min_contracts_pct():.2f}%\n"
            f"價格確認門檻：+/-{spike_price_confirm_pct():.2f}%\n"
            f"同幣冷卻時間：{spike_cooldown_seconds()} 秒\n"
            f"1H 底部點火：OI +{trend_min_contracts_pct():.2f}%、價格 +{trend_min_price_pct():.2f}% 至 +{trend_max_bottom_price_pct():.2f}%\n"
            f"1H 強勢延續：OI +{momentum_min_contracts_pct():.2f}%、價格 +{momentum_min_price_pct():.2f}%\n"
            f"流動性快篩：24H成交額 >= ${fmt_num(liquidity_min_quote_volume_24h_usd())}\n"
            f"1H 通知冷卻：{trend_cooldown_seconds()} 秒"
        )

    if command == "/universe":
        symbols = resolve_watch_symbols(force_refresh=True)
        return (
            "動態監控清單\n"
            f"模式：{watch_mode()}\n"
            f"來源：{watch_source_description()}\n"
            f"監控標的數：{len(symbols)}\n"
            f"清單刷新：每 {watchlist_refresh_seconds()} 秒\n"
            f"並行查詢數：{oi_snapshot_workers()}\n"
            f"報表列出：前 {report_top_n()} 名"
        )

    if command in {"/entry", "/in", "entry", "進場", "買入", "做多", "做空"}:
        if chat_id is None:
            return "This command needs a Telegram chat."
        default_side = "SHORT" if command == "做空" else "LONG"
        return register_position(int(chat_id), args, default_side)

    if command in {"/exit", "/out", "exit", "平倉", "出場", "移除"}:
        if chat_id is None:
            return "This command needs a Telegram chat."
        return remove_position(int(chat_id), args)

    if command in {"/positions", "/pos", "倉位", "持倉"}:
        if chat_id is None:
            return "This command needs a Telegram chat."
        return format_positions(int(chat_id))

    if command in {"/onchain", "/chain"}:
        if not args:
            return "用法：/onchain COAI"
        try:
            symbol = normalize_symbol(args[0])
        except Exception:
            symbol = args[0].strip().upper()
        watch = market_context(symbol)
        market_symbol = watch.market_symbol if watch else symbol.removesuffix("USDT")
        try:
            signal = analyze_onchain(
                symbol,
                market_symbol=market_symbol,
                provider_id=watch.provider_id if watch else None,
            )
        except Exception as exc:
            return f"鏈上檢查失敗：{exc}"
        return format_onchain_report(signal)

    if command in {"/onchain_report", "/chain_report"}:
        return load_cached_wgl_report()

    if command in {"/orderbook", "/book", "/ob"}:
        if not args:
            return "用法：/orderbook COAI"
        try:
            symbol = normalize_symbol(args[0])
        except Exception:
            symbol = args[0].strip().upper()
        try:
            seed_orderbook_symbol(symbol)
            signal = analyze_orderbook_accumulation(
                symbol,
                orderbook_db_path(),
                lookback_seconds=orderbook_lookback_seconds(),
                min_snapshots=orderbook_min_snapshots(),
            )
        except Exception as exc:
            return f"訂單簿檢查失敗：{exc}"
        return format_orderbook_signal(signal)

    if command in {"/thesis", "/research", "/wgl"}:
        if not args:
            return "用法：/wgl COAI"
        try:
            return build_research_thesis(args[0], RUNTIME_SPIKE_HISTORY)
        except Exception as exc:
            return f"交易論證卡產生失敗：{exc}"

    if command in {"/strategy_report", "/strategy", "/strat"}:
        return build_strategy_report(RUNTIME_SPIKE_HISTORY)

    if command == "/strategy_settings":
        if strategy_mode() == "trade_plan":
            return (
                "方向型 TP1 / TP2 / SL 模擬策略\n"
                "訊號：只採用每小時報告中的做多或做空完整計畫\n"
                f"每筆：{fmt_num(trade_plan_paper_margin_usd())}U × {fmt_num(trade_plan_paper_leverage())}倍\n"
                "管理：TP1停利一半，剩餘止損移到進場價；TP2出清\n"
                f"訊號有效：{trade_plan_signal_max_age_seconds()} 秒\n"
                f"最多持有：{trade_plan_max_open_positions()} 檔\n"
                f"同幣再進場冷卻：{trade_plan_reentry_cooldown_seconds()} 秒\n"
                "舊策略：只追蹤既有單出場，不再新增"
            )
        return (
            "階段模型 TP/SL 策略設定\n"
            "開倉框架：日線長底部，1m/5m/1h 只作輔助，不再當主因\n"
            f"TP：+{strategy_take_profit_pct():.2f}%\n"
            f"SL：-{strategy_stop_loss_pct():.2f}%\n"
            f"最多持倉：{strategy_max_open_positions()} 檔\n"
            f"每次最多新開：{strategy_max_new_positions_per_scan()} 檔\n"
            f"全域開倉間隔：{strategy_min_global_entry_gap_seconds()} 秒\n"
            f"單幣冷卻：{strategy_signal_cooldown_seconds()} 秒\n"
            f"日線需求：至少 {strategy_min_daily_candles()} 根，回看 {strategy_daily_lookback_days()} 天\n"
            f"歷史日線波動：至少 {strategy_min_daily_range_multiple():.2f}x\n"
            f"日線離低點上限：+{strategy_max_daily_extension_from_low_pct():.2f}%\n"
            f"日線底部區間位置上限：{strategy_max_bottom_range_position_pct():.2f}%\n"
            f"近 60 日低點距離上限：+{strategy_max_recent_low_extension_pct():.2f}%\n"
            f"高點回撤下限：{strategy_min_drawdown_from_high_pct():.2f}%\n"
            f"近 30 日底部天數：至少 {strategy_min_base_days()} 天落在底部 {strategy_max_base_band_from_low_pct():.2f}% 區間\n"
            f"4H底部位階上限：{strategy_max_4h_range_position_pct():.2f}%\n"
            f"起漲日線位階上限：{strategy_max_launch_range_position_pct():.2f}%\n"
            f"起漲近低距離上限：+{strategy_max_launch_recent_low_extension_pct():.2f}%\n"
            f"起漲高點回撤下限：{strategy_min_launch_drawdown_from_high_pct():.2f}%\n"
            f"起漲 24h/3d/7d 漲幅上限：{strategy_max_launch_24h_price_pct():.2f}% / {strategy_max_launch_3d_price_pct():.2f}% / {strategy_max_launch_7d_price_pct():.2f}%\n"
            f"起漲 4H 位階/24h/3d 上限：{strategy_max_launch_4h_range_position_pct():.2f}% / {strategy_max_launch_4h_24h_price_pct():.2f}% / {strategy_max_launch_4h_3d_price_pct():.2f}%\n"
            f"Funding 上限：{strategy_max_funding_pct():.4f}%\n"
            "市值條件：無；市值與排名只顯示參考，不參與准入或評分\n"
            f"鏈上最低開倉分數：{strategy_onchain_min_score():+d}（合約地址驗證後才計分）\n"
            f"訂單簿收集：{'開' if orderbook_enabled() else '關'}｜每 {orderbook_collect_interval_seconds()} 秒｜候選 {orderbook_watch_candidates()} 檔\n"
            f"訂單簿觀察：{orderbook_lookback_seconds()} 秒｜至少 {orderbook_min_snapshots()} 快照｜吸籌門檻 {orderbook_min_score()} 分\n"
            f"策略回報：每 {strategy_report_interval_seconds()} 秒"
        )

    if command == "/position_settings":
        return (
            "倉位盯盤設定\n"
            f"檢查間隔：{position_check_seconds()} 秒\n"
            f"OI 轉弱窗口：{position_oi_window_seconds()} 秒\n"
            f"Funding 出場門檻：>{position_funding_exit_pct():.4f}% 且已獲利\n"
            f"停損：-{position_stop_loss_pct():.2f}%\n"
            f"半停利：+{position_take_profit_pct():.2f}%\n"
            "半停利後：剩餘倉止損拉到開倉價"
        )

    if command == "/subscribe":
        if chat_id is None:
            return "此指令只能在 Telegram 使用。"
        return configure_alerts(int(chat_id), True)

    if command == "/unsubscribe":
        if chat_id is None:
            return "此指令只能在 Telegram 使用。"
        return configure_alerts(int(chat_id), False)

    if command == "/report":
        return load_cached_wgl_report()

    if command in {"/oi_report", "/old_report"}:
        report, _ = build_oi_report()
        return report

    if command == "/scan":
        limit = 15
        if args:
            try:
                limit = min(max(int(args[0]), 1), 30)
            except ValueError:
                return "用法：/scan 15"
        results = scan_symbols(limit)
        for item in results:
            save_analysis(item, DATA_ROOT)
        return format_scan(results)

    if command == "/oi":
        if not args:
            return "用法：/oi BTC"
        analysis = analyze_symbol(args[0])
        save_analysis(analysis, DATA_ROOT)
        return format_analysis(analysis)

    if command.startswith("/"):
        return "未知指令。請用 /help 查看指令。"

    # Convenience mode: sending "BTC" is treated as "/oi BTC".
    if len(text.strip()) <= 20:
        analysis = analyze_symbol(text.strip())
        save_analysis(analysis, DATA_ROOT)
        return format_analysis(analysis)

    return "請使用 /oi BTC 或 /help。"


def clone_history(history: dict[str, deque[dict[str, Any]]]) -> dict[str, deque[dict[str, Any]]]:
    snapshot: dict[str, deque[dict[str, Any]]] = {}
    for symbol, samples in list(history.items()):
        maxlen = samples.maxlen if isinstance(samples, deque) else None
        snapshot[str(symbol)] = deque((dict(item) for item in list(samples)), maxlen=maxlen)
    return snapshot


def run_bot() -> None:
    load_env(ROOT / ".env")
    RUNTIME_STRUCTURE_CACHE.update(load_structure_cache())
    token = os.environ.get("TELEGRAM_BOT_TOKEN", "").strip()
    if not token:
        raise SystemExit("TELEGRAM_BOT_TOKEN is missing. Copy .env.example to .env and fill it.")

    allowed = parse_allowed_chat_ids()
    telegram_call(
        token,
        "setMyCommands",
        {
            "commands": json.dumps(
                telegram_menu_commands(),
                separators=(",", ":"),
            )
        },
    )

    subscribers = load_subscribers()
    offset = initialize_update_offset(token)
    next_report_at = initial_report_due_at()
    next_strategy_report_at = time.time() + strategy_report_interval_seconds()
    next_strategy_scan_at = time.time() + strategy_scan_interval_seconds()
    next_spike_check_at = time.time() + 5
    next_position_check_at = time.time() + 5
    next_orderbook_collect_at = time.time() + 3
    spike_history = RUNTIME_SPIKE_HISTORY
    position_history = RUNTIME_POSITION_HISTORY
    last_spike_alert_at: dict[str, float] = {}
    last_strategy_signal_at: dict[str, float] = {}
    daily_summary_state = load_wgl_daily_summary_state()
    report_executor = ThreadPoolExecutor(max_workers=1, thread_name_prefix="wgl-report")
    maintenance_executor = ThreadPoolExecutor(max_workers=1, thread_name_prefix="orderbook")
    strategy_executor = ThreadPoolExecutor(max_workers=1, thread_name_prefix="strategy")
    report_future: Future[str] | None = None
    orderbook_future: Future[tuple[int, int]] | None = None
    strategy_future: Future[list[str]] | None = None
    print("OI bot polling started. Press Ctrl+C to stop.", flush=True)
    while True:
        try:
            updates = telegram_call(
                token,
                "getUpdates",
                {"offset": offset, "timeout": 15, "allowed_updates": json.dumps(["message"])},
                timeout=25,
            )
            for update in updates:
                offset = max(offset, int(update["update_id"]) + 1)
                message = update.get("message") or {}
                chat = message.get("chat") or {}
                chat_id = chat.get("id")
                text = message.get("text", "")
                if chat_id is None or not text:
                    continue
                if not allowed_chat(int(chat_id), allowed):
                    telegram_call(token, "sendMessage", {"chat_id": chat_id, "text": "Unauthorized chat."})
                    continue
                try:
                    reply = handle_text(text, int(chat_id))
                except (ApiError, ValueError) as exc:
                    reply = f"Data error: {exc}"
                except Exception as exc:
                    reply = f"Bot error: {exc}"
                command, _ = split_command(text)
                if command in {"/report", "/strategy_report", "/positions"}:
                    send_card_message(token, int(chat_id), reply)
                else:
                    send_long_message(token, int(chat_id), reply)

            if orderbook_future is not None and orderbook_future.done():
                try:
                    orderbook_future.result()
                except Exception as exc:
                    print(f"Orderbook collect cycle error: {exc}", file=sys.stderr, flush=True)
                orderbook_future = None
            if (
                time.time() >= next_orderbook_collect_at
                and orderbook_future is None
                and report_future is None
                and time.time() < next_report_at
            ):
                orderbook_future = maintenance_executor.submit(collect_orderbook_cycle, clone_history(spike_history))
                next_orderbook_collect_at = time.time() + orderbook_collect_interval_seconds()

            if report_future is not None and report_future.done():
                try:
                    report = report_future.result()
                except Exception as exc:
                    report = ""
                    next_report_at = min(next_report_at, time.time() + 120)
                    print(f"每小時全市場掃描失敗：{exc}", file=sys.stderr, flush=True)
                subscribers = load_subscribers()
                if report:
                    for chat_id in subscribers:
                        try:
                            send_card_message(token, int(chat_id), report)
                        except Exception as exc:
                            print(f"Report send error for {chat_id}: {exc}", file=sys.stderr, flush=True)
                for alert in consume_wgl_transition_alerts():
                    for chat_id in subscribers:
                        try:
                            send_card_message(token, int(chat_id), alert)
                        except Exception as exc:
                            print(f"Transition send error for {chat_id}: {exc}", file=sys.stderr, flush=True)
                report_future = None
            if time.time() >= next_report_at and report_future is None:
                report_future = report_executor.submit(
                    build_onchain_hourly_report,
                    clone_history(spike_history),
                    persist=True,
                )
                next_report_at = time.time() + report_interval_seconds()

            if time.time() >= next_strategy_report_at:
                subscribers = load_subscribers()
                if subscribers:
                    report = build_strategy_report(spike_history)
                    for chat_id in subscribers:
                        try:
                            send_card_message(token, int(chat_id), report)
                        except Exception as exc:
                            print(f"Strategy report send error for {chat_id}: {exc}", file=sys.stderr, flush=True)
                next_strategy_report_at = time.time() + strategy_report_interval_seconds()

            if wgl_daily_summary_due(daily_summary_state):
                summary_day = local_day_key()
                try:
                    summary = build_wgl_daily_summary(summary_day)
                except Exception as exc:
                    summary = f"每日資金異動統計失敗：{exc}"
                    print(f"WGL daily summary error: {exc}", file=sys.stderr, flush=True)
                subscribers = load_subscribers()
                for chat_id in subscribers:
                    try:
                        send_card_message(token, int(chat_id), summary)
                    except Exception as exc:
                        print(f"WGL daily summary send error for {chat_id}: {exc}", file=sys.stderr, flush=True)
                daily_summary_state["last_sent_date"] = summary_day
                daily_summary_state["last_sent_utc"] = time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime())
                daily_summary_state["last_sent_local"] = time.strftime("%Y-%m-%d %H:%M")
                save_wgl_daily_summary_state(daily_summary_state)

            if time.time() >= next_position_check_at:
                try:
                    position_alerts = collect_position_alerts(position_history)
                except Exception as exc:
                    position_alerts = []
                    print(f"Position monitor error: {exc}", file=sys.stderr, flush=True)
                for chat_id, alert in position_alerts:
                    try:
                        send_card_message(token, int(chat_id), alert)
                    except Exception as exc:
                        print(f"Position send error for {chat_id}: {exc}", file=sys.stderr, flush=True)
                next_position_check_at = time.time() + position_check_seconds()

            if strategy_future is not None and strategy_future.done():
                try:
                    strategy_alerts = strategy_future.result()
                except Exception as exc:
                    strategy_alerts = []
                    print(f"Strategy monitor error: {exc}", file=sys.stderr, flush=True)
                subscribers = load_subscribers()
                for alert in strategy_alerts:
                    for chat_id in subscribers:
                        try:
                            send_card_message(token, int(chat_id), alert)
                        except Exception as exc:
                            print(f"Strategy send error for {chat_id}: {exc}", file=sys.stderr, flush=True)
                strategy_future = None

            if time.time() >= next_spike_check_at and report_future is None and orderbook_future is None:
                subscribers = load_subscribers()
                if subscribers:
                    try:
                        alerts = collect_spike_alerts(spike_history, last_spike_alert_at)
                    except Exception as exc:
                        alerts = []
                        print(f"Spike monitor error: {exc}", file=sys.stderr, flush=True)
                    for alert in alerts:
                        for chat_id in subscribers:
                            try:
                                send_card_message(token, int(chat_id), alert)
                            except Exception as exc:
                                print(f"Spike send error for {chat_id}: {exc}", file=sys.stderr, flush=True)
                    if strategy_future is None and time.time() >= next_strategy_scan_at:
                        strategy_future = strategy_executor.submit(
                            collect_strategy_alerts,
                            clone_history(spike_history),
                            last_strategy_signal_at,
                        )
                        next_strategy_scan_at = time.time() + strategy_scan_interval_seconds()
                next_spike_check_at = time.time() + spike_check_seconds()
        except KeyboardInterrupt:
            report_executor.shutdown(wait=False, cancel_futures=True)
            maintenance_executor.shutdown(wait=False, cancel_futures=True)
            strategy_executor.shutdown(wait=False, cancel_futures=True)
            print("Stopped.", flush=True)
            return
        except Exception as exc:
            print(f"Polling error: {exc}", file=sys.stderr, flush=True)
            time.sleep(5)


def main() -> None:
    load_env(ROOT / ".env")
    parser = argparse.ArgumentParser(description="Binance OI Telegram bot")
    parser.add_argument("--once", metavar="SYMBOL", help="run one local OI analysis without Telegram")
    parser.add_argument("--scan", metavar="N", type=int, help="run one local scan without Telegram")
    parser.add_argument("--report", action="store_true", help="run one local Excel OI ranking report without Telegram")
    parser.add_argument("--universe", action="store_true", help="show current dynamic watch universe size")
    args = parser.parse_args()

    if args.once:
        analysis = analyze_symbol(args.once)
        save_analysis(analysis, DATA_ROOT)
        print(format_analysis(analysis))
        return
    if args.scan:
        results = scan_symbols(args.scan)
        for item in results:
            save_analysis(item, DATA_ROOT)
        print(format_scan(results))
        return
    if args.report:
        report, report_path = build_oi_report()
        print(report)
        print(f"\nReport workbook: {report_path}")
        return
    if args.universe:
        symbols = resolve_watch_symbols(force_refresh=True)
        print(f"mode={watch_mode()}")
        print(f"source={watch_source_description()}")
        print(f"symbols={len(symbols)}")
        for watch in symbols[:30]:
            print(f"{watch.symbol} rank={watch.market_rank} marketcap={watch.marketcap_usd}")
        return
    run_bot()


if __name__ == "__main__":
    main()
